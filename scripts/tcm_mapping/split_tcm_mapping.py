"""证型对应表拆分（10-26 汇报第 12 页“证型对应”一线：拆开“¤”、去重，未填的发回 GD）。

输入：docs/数据对齐/证型对应.xlsx（GD 临床 10-25 交回：355 种证候写法、两院就诊次数、GD 填的方案证型；只含写法与次数）。
150 种写法是“¤”连起来的多个证候（多半是同一证候重复，如“痰热阻肺证¤痰热阻肺证”）。这里先拆开、去重成单个证候：

- 单个证候的证型：取 GD 在“只由这一个证候组成”的写法上填的（如“A”“A¤A”）；这些写法填的不一样记“冲突”；
  都没填、但它在多证候写法（如“A¤B”）里出现、那条写法填了的，记“只在多证候写法里填过”，不自动采用；
- 就诊次数按写法累加（一次就诊写了几种写法会重复计），只作排序参考。

输出一个 Excel：
1 单个证候：证候、两院就诊次数、对应证型、状态（已填 / 冲突 / 只在多证候写法里填过 / 待补）、来自几种写法；
2 待 GD 补填：状态不是“已填”的，空一列给 GD 填（7 个证型之一，或“不对应”）；
3 多证候写法：由 ≥2 种不同证候组成的写法和 GD 填的证型——“一次就诊对到多个证型时怎么取”要另定口径，这张表给讨论用。

用法（BaseModel 根目录）：python scripts/tcm_mapping/split_tcm_mapping.py [--input docs/数据对齐/证型对应.xlsx] [--out 输出.xlsx]
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TYPES = ["风热犯卫", "风寒束表", "表寒里热", "热毒袭肺", "毒热壅盛", "毒热内陷、内闭外脱", "气阴两虚、正气未复", "不对应"]
SEP = "¤"


def atoms(writing: str) -> list[str]:
    """一种写法拆成单个证候（去空白、按出现顺序去重）。"""
    out: list[str] = []
    for x in str(writing).split(SEP):
        x = x.strip()
        if x and x not in out:
            out.append(x)
    return out


def split(rows: list[tuple]) -> tuple[list[dict], list[dict]]:
    """rows：(写法, 已对到方案证型, GD 就诊, GAM 就诊, GD 填的证型)。返回（单个证候, 多证候写法）。"""
    acc: dict[str, dict] = defaultdict(lambda: {"gd": 0, "gam": 0, "pure": set(), "mixed": set(), "n": 0})
    multi = []
    for w, _, gd, gam, ty in rows:
        if not w:
            continue
        ty = (str(ty).strip() or None) if ty is not None else None
        a = atoms(w)
        for x in a:
            d = acc[x]
            d["gd"] += int(gd or 0)
            d["gam"] += int(gam or 0)
            d["n"] += 1
            if ty:
                (d["pure"] if len(a) == 1 else d["mixed"]).add(ty)
        if len(a) >= 2:
            multi.append({"写法": w, "证候": "、".join(a), "GD 就诊": int(gd or 0), "GAM 就诊": int(gam or 0), "GD 填的证型": ty or ""})
    out = []
    for x, d in acc.items():
        if len(d["pure"]) == 1:
            ty, st = next(iter(d["pure"])), "已填"
        elif len(d["pure"]) > 1:
            ty, st = " / ".join(sorted(d["pure"])), "冲突"
        elif d["mixed"]:
            ty, st = " / ".join(sorted(d["mixed"])), "只在多证候写法里填过"
        else:
            ty, st = "", "待补"
        out.append({"证候": x, "GD 就诊": d["gd"], "GAM 就诊": d["gam"], "对应证型": ty, "状态": st, "来自几种写法": d["n"]})
    out.sort(key=lambda r: (r["状态"] == "已填", -(r["GD 就诊"] + r["GAM 就诊"]), r["证候"]))
    multi.sort(key=lambda r: -(r["GD 就诊"] + r["GAM 就诊"]))
    return out, multi


def _sheet(wb: Workbook, title: str, cols: list[str], recs: list[dict], widths: dict[str, int], note: str) -> None:
    ws = wb.create_sheet(title)
    ws.append([note])
    ws["A1"].font = Font(italic=True, color="666666")
    ws.append(cols)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in recs:
        ws.append([r.get(c, "") for c in cols])
    for i, c in enumerate(cols):
        ws.column_dimensions[chr(65 + i)].width = widths.get(c, 12)
    ws.freeze_panes = "A3"
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", type=Path, default=Path("docs/数据对齐/证型对应.xlsx"))
    ap.add_argument("--out", type=Path, default=Path("docs/数据对齐/证型对应_拆分与待补_20261027.xlsx"))
    a = ap.parse_args()
    rows = list(load_workbook(a.input, read_only=True).active.iter_rows(values_only=True))[1:]
    single, multi = split([tuple(r[:5]) for r in rows])
    todo = [dict(r, **{"请填对应证型": ""}) for r in single if r["状态"] != "已填"]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "单个证候", ["证候", "GD 就诊", "GAM 就诊", "对应证型", "状态", "来自几种写法"], single,
           {"证候": 22, "对应证型": 26, "状态": 20},
           f"由 {len(rows)} 种写法拆开“¤”、去重得到 {len(single)} 种证候；就诊次数按写法累加，只作排序参考")
    _sheet(wb, "待 GD 补填", ["证候", "GD 就诊", "GAM 就诊", "状态", "对应证型", "请填对应证型"], todo,
           {"证候": 22, "状态": 20, "对应证型": 26, "请填对应证型": 22},
           "请在最后一列填方案 7 个证型之一，或“不对应”；“对应证型”列是在别的写法里填过的，供参考")
    ws = wb["待 GD 补填"]
    dv = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    if todo:
        dv.add(f"F3:F{len(todo) + 2}")
    _sheet(wb, "多证候写法", ["写法", "证候", "GD 就诊", "GAM 就诊", "GD 填的证型"], multi, {"写法": 40, "证候": 32, "GD 填的证型": 22},
           "由 ≥2 种不同证候组成的写法：一次就诊对到多个证型时怎么取，要另定口径")
    wb.save(a.out)
    st = defaultdict(int)
    for r in single:
        st[r["状态"]] += 1
    print(f"{len(rows)} 种写法 → {len(single)} 种证候（" + "、".join(f"{k} {v}" for k, v in st.items()) + f"）；多证候写法 {len(multi)} 种；写到 {a.out}")


if __name__ == "__main__":
    main()
