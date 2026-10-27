# -*- coding: utf-8 -*-
"""
生成《老年流感转重研究 · 近期工作与论文阅读汇报》（2026-10-26）。

面向导师 / 课题组的汇报：13 页，按演讲者备注约 17 分钟（不含讨论）。

输入：PPT/hospital_优化版.pptx（模板，只读，不会被修改；沿用其主题、封面和版式，同 build_weekly_summary_20261023.py）
      PPT/论文讲解页_卡片_CLIF.png（论文路径讲解页里 CLIF 卡片的截图，不含病例内容）
输出：PPT/近期工作与论文阅读汇报_20261026.pptx

样式与工具函数原样沿用 PPT/build_weekly_summary_20261023.py（标题栏 + 章节标签 + 青色分隔线 + 关键信息条
+ 浅黄结论框 + 页脚页码；字体微软雅黑）。

数字都是聚合量，1–9 的计数写作 "<10"。来源：
- docs/数据对齐/GD_临床负责人问题单_20261025.docx：10-24 的 std_v3 导出（规则集 guideline2025_v0.3）上的计数；
  GD 临床负责人的答复（16 条批注、意见栏的 √ 与文字、附件 2 标黄的 12 个科室，只含名称与次数）；
- docs/数据对齐/证型对应.xlsx：355 种证候写法与 GD 填的对应证型（只含写法与次数）；
- paper/数据集构建调研_20261025/：论文路径讲解页、设计要点笔记（文献数字以笔记为准，含页码 / 章节）、00_论文身份核对.md；
- flucase README“当前状态”（10-17 至 10-24 各行）、docs/claude_project/02_项目背景知识.md §5、§6.3。
这两份 GD 答复文件含真实聚合计数，不入库；本 PPT 与本脚本同样含真实聚合计数，是否入库由研究者定。
没有读任何真实病例。

运行（BaseModel 根目录）：python PPT/build_progress_reading_report_20261026.py
依赖：python-pptx、lxml
"""

import uuid
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "hospital_优化版.pptx"
CARD_CLIF = HERE / "论文讲解页_卡片_CLIF.png"
OUT = HERE / "近期工作与论文阅读汇报_20261026.pptx"

# ---------------------------------------------------------------------------
# 全局样式（与 build_weekly_summary_20261023.py 一致，另加两院与证型分期的颜色）
# ---------------------------------------------------------------------------
FONT = "微软雅黑"
DECK_TITLE = "老年流感转重研究 · 近期工作与论文阅读汇报"

DARK = "1F4E94"     # 深蓝：顶部色条、卡片标题栏、GD
MID = "4075BE"      # 中蓝：标题栏
CYAN = "00A4CB"     # 青色：分隔线、GAM、新材料
LIGHT = "EBF3FB"    # 浅蓝：关键信息条
PALE = "F5F8FC"     # 更浅的蓝灰：行底色
BORDER = "CDDAE8"   # 卡片描边
NAVY = "11376A"     # 深蓝文字
BODY = "273140"     # 正文
MUTED = "646E79"    # 说明文字、页脚
NOTE_BG = "FCEFC6"  # 浅黄：底部结论框
WHITE = "FFFFFF"
RED = "C0392B"      # 重症 / 偏差
AMBER = "B7791F"    # 待定 / 改口径
GREY = "8A94A6"     # 未填 / 等待
AMBER_BG = "FBF1DF"  # 浅琥珀：选项底色
CYAN_BG = "E3F5FA"   # 浅青：10-24 一列的底色
RED_BG = "FBE9E7"    # 浅红：流感高峰

GD_C = DARK         # 图表：GD
GAM_C = CYAN        # 图表：GAM

CL = 0.55           # 版心左边界（英寸）
CW = 12.23          # 版心宽度
SW = 13.333         # 幻灯片宽度

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------------------
# 文字工具（沿用 build_hospital_optimized.py）
# ---------------------------------------------------------------------------
def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_typeface(rPr, font=FONT):
    """西文、东亚、复杂文种字体都设为 font（只写 latin 时中文会回退成宋体）。"""
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
        for attr in ("panose", "charset", "pitchFamily"):
            if attr in el.attrib:
                del el.attrib[attr]


def style_run(run, size, color=BODY, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    set_typeface(rPr)


def set_bullet(paragraph, color=MID, indent=0.2):
    """项目符号（悬挂缩进），不在文字里写“•”。"""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def add_break(paragraph, size):
    br = paragraph._p.add_br()
    rPr = etree.SubElement(br, qn("a:rPr"), lang="zh-CN", sz=str(int(round(size * 100))))
    set_typeface(rPr)


def fill_text(tf, paras, size=14, color=BODY, bold=False, align="l", line_spacing=1.08, space_after=0):
    """paras：字符串；或列表，每段为字符串 / [(文字, {选项})] / {"runs":…, "bullet":…, "size":…}。"""
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = spec if isinstance(spec, dict) else {"runs": spec}
        runs = opts.get("runs", "")
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = ALIGN[opts.get("align", align)]
        p.line_spacing = opts.get("line_spacing", line_spacing)
        p.space_before = Pt(opts.get("space_before", 0))
        p.space_after = Pt(opts.get("space_after", space_after))
        p_size = opts.get("size", size)
        p_color = opts.get("color", color)
        p_bold = opts.get("bold", bold)
        for text, ropt in runs:
            r_size = ropt.get("size", p_size)
            for k, part in enumerate(text.split("\n")):
                if k:
                    add_break(p, r_size)
                if not part:
                    continue
                r = p.add_run()
                r.text = part
                style_run(r, r_size, ropt.get("color", p_color), ropt.get("bold", p_bold), ropt.get("italic", False))
        if opts.get("bullet"):
            set_bullet(p, color=opts.get("bullet_color", MID))


def setup_frame(tf, anchor="t", margins=(0.05, 0.03, 0.05, 0.03), wrap=True):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = ANCHOR[anchor]


def add_text(slide, x, y, w, h, paras, size=14, color=BODY, bold=False, align="l", anchor="t",
             margins=(0.05, 0.03, 0.05, 0.03), line_spacing=1.08, space_after=0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    setup_frame(tb.text_frame, anchor, margins)
    fill_text(tb.text_frame, paras, size, color, bold, align, line_spacing, space_after)
    if name:
        tb.name = name
    return tb


def text_width(s, size_pt):
    """粗略估计单行文字宽度（英寸），用来给徽章、标签定宽。"""
    em = 0.0
    for ch in s:
        if ord(ch) >= 0x2E80:
            em += 1.0
        elif ch == " ":
            em += 0.3
        elif ch in "·•":
            em += 0.5
        else:
            em += 0.58
    return em * size_pt / 72


def bullets(items, gap=4, color=MID):
    """把字符串或 [(文字, 选项)] 列表转成带项目符号的段落。"""
    return [{"runs": t, "bullet": True, "space_after": gap, "bullet_color": color} for t in items]


def lead(label, text, label_color=NAVY):
    """“粗体标签 + 正文”的一段。"""
    return [(label, {"bold": True, "color": label_color}), (text, {})]


# ---------------------------------------------------------------------------
# 形状工具（沿用）
# ---------------------------------------------------------------------------
def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, alpha=16, blur=5, dist=1.5):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)),
                          dist=str(int(dist * 12700)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="1F3B66")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE,
             radius=None, shadow=False, dash=None, name=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    else:
        shp.line.fill.background()
    if radius is not None:
        shp.adjustments[0] = radius
    if shadow:
        add_shadow(shp)
    if name:
        shp.name = name
    return shp


def shape_text(shape, paras, size=14, color=BODY, bold=False, align="c", anchor="m",
               margins=(0.04, 0.02, 0.04, 0.02), wrap=True):
    setup_frame(shape.text_frame, anchor, margins, wrap)
    fill_text(shape.text_frame, paras, size, color, bold, align)
    return shape


def add_oval(slide, cx, cy, d, fill=None, line=None, line_w=1.25, dash=None):
    return add_rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, line_w=line_w,
                    shape=MSO_SHAPE.OVAL, dash=dash)


def add_line(slide, x1, y1, x2, y2, color, width=1.0, arrow_end=False, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow_end:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


# ---------------------------------------------------------------------------
# 表格工具（沿用）
# ---------------------------------------------------------------------------
def set_cell_border(cell, color=None, width=0.75, sides="B"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for i, (tag, side) in enumerate((("a:lnL", "L"), ("a:lnR", "R"), ("a:lnT", "T"), ("a:lnB", "B"))):
        ln = etree.Element(qn(tag))
        if color and side in sides:
            ln.set("w", str(int(width * 12700)))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def add_table(slide, x, y, col_w, row_h, rows, size=13, header_size=13.5, header_fill=DARK,
              line_color="DDE6F0", margins=(0.1, 0.05, 0.1, 0.05)):
    nrows, ncols = len(rows), len(col_w)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tbl.vert_banding = False
    for i, wv in enumerate(col_w):
        tbl.columns[i].width = Inches(wv)
    for i, hv in enumerate(row_h):
        tbl.rows[i].height = Inches(hv)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            spec = val if isinstance(val, dict) else {"text": val}
            cell = tbl.cell(r, c)
            is_head = r == 0
            fill = spec.get("fill", header_fill if is_head else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            set_cell_border(cell, None if is_head else line_color, 0.75, "B")
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [Inches(m) for m in margins]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text(tf, spec.get("text", ""),
                      size=spec.get("size", header_size if is_head else size),
                      color=spec.get("color", WHITE if is_head else BODY),
                      bold=spec.get("bold", is_head),
                      align=spec.get("align", "l"))
    return tbl


def merge_rows(tbl, col, r0, r1, paras, fill="DCE7F5"):
    """合并某一列的若干行，写入新文字（同模板第 13 页的“批次”列）。"""
    cell = tbl.cell(r0, col)
    cell.merge(tbl.cell(r1, col))
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    tf = cell.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    tf.paragraphs[0].text = ""
    fill_text(tf, paras, 13, NAVY, align="c")
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


# ---------------------------------------------------------------------------
# 页面框架（沿用模板：顶部色条、标题栏、章节标签、分隔线、关键信息条、结论框、页脚、页码）
# ---------------------------------------------------------------------------
def add_slide_number(slide, x=11.68, y=7.06, w=1.1, h=0.28):
    number = slide.part.package.presentation_part.presentation.slides.index(slide) + 1
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "页码"
    setup_frame(tb.text_frame, "m", (0, 0, 0, 0))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{%s}" % str(uuid.uuid4()).upper(), type="slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="zh-CN", sz="1050")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=MUTED)
    set_typeface(rPr)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = str(number)


def frame(slide, badge, title, key, ask=None, ask_lead=None, notes=None):
    add_rect(slide, 0, 0, SW, 0.10, fill=DARK, name="顶部色条")
    add_rect(slide, 0, 0.24, 10.30, 0.58, fill=MID, name="标题栏")
    bw = text_width(badge, 14) + 0.34
    b = add_rect(slide, 0.35, 0.33, bw, 0.40, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.25, name="章节标签")
    shape_text(b, badge, 14, DARK, bold=True)
    tx = 0.35 + bw + 0.16
    add_text(slide, tx, 0.24, 10.30 - tx - 0.15, 0.58, title, size=22, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0), name="标题")
    add_rect(slide, 0.50, 0.98, 12.28, 0.04, fill=CYAN, name="分隔线")
    add_rect(slide, CL, 1.12, CW, 0.50, fill=LIGHT, name="关键信息条")
    add_text(slide, CL + 0.2, 1.12, CW - 0.4, 0.50, key, size=16, color=NAVY, bold=True, anchor="m",
             name="关键信息")
    if ask:
        add_rect(slide, CL, 6.42, CW, 0.46, fill=NOTE_BG, name="结论框")
        runs = [(ask_lead, {"bold": True}), (ask, {})] if ask_lead else [(ask, {})]
        add_text(slide, CL + 0.2, 6.42, CW - 0.4, 0.46, [runs], size=15, color=NAVY, anchor="m",
                 name="结论")
    add_text(slide, CL, 7.06, 9.0, 0.28, DECK_TITLE, size=10.5, color=MUTED, anchor="m",
             margins=(0, 0, 0, 0), name="页脚")
    add_slide_number(slide)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def card(slide, x, y, w, h, header, header_h=0.44, header_size=16, header_fill=DARK, shadow=True):
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=shadow)
    add_rect(slide, x, y, w, header_h, fill=header_fill)
    add_text(slide, x + 0.15, y, w - 0.3, header_h, header, size=header_size, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0))


def stat_tile(slide, x, y, w, h, big, label, big_size=20, label_size=12):
    """白底小卡片：上面大数字，下面一行说明。"""
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, x + 0.16, y + 0.06, w - 0.32, h * 0.52, big, size=big_size, color=DARK, bold=True,
             anchor="b", margins=(0, 0, 0, 0))
    add_text(slide, x + 0.16, y + 0.06 + h * 0.52, w - 0.32, h * 0.44 - 0.1, label, size=label_size,
             color=MUTED, anchor="t", margins=(0, 0.02, 0, 0))


def chip(slide, x, y, text, kind, size=12.5, h=0.36):
    """判据小块：kind = guideline（方案原文）/ proxy（代理）/ weak（弱代理）。返回宽度。

    宽度按字数估计后留足余量（≥、%、/ 等符号在雅黑里比估计的宽），并关掉自动换行，避免一个判据折成两行。
    """
    w = text_width(text, size) * 1.06 + 0.42
    if kind == "guideline":
        s = add_rect(slide, x, y, w, h, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        shape_text(s, text, size, WHITE, bold=True, wrap=False)
    elif kind == "proxy":
        s = add_rect(slide, x, y, w, h, fill=WHITE, line=MID, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.3)
        shape_text(s, text, size, NAVY, bold=True, wrap=False)
    else:
        s = add_rect(slide, x, y, w, h, fill=WHITE, line=GREY, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.3, dash=MSO_LINE_DASH_STYLE.DASH)
        shape_text(s, text, size, MUTED, bold=True, wrap=False)
    return w


def chart_fonts(ch):
    """图表文字补上东亚字体（与模板一致）。"""
    for latin in list(ch._chartSpace.iter(qn("a:latin"))):
        parent = latin.getparent()
        if parent.find(qn("a:ea")) is None:
            ea = etree.Element(qn("a:ea"))
            ea.set("typeface", FONT)
            latin.addnext(ea)


def hide_value_axis(va):
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va._element.get_or_add_delete_().set("val", "1")
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE




# ---------------------------------------------------------------------------
# 本次新增的小工具
# ---------------------------------------------------------------------------
def small(n):
    """计数 1–9 写作 "<10"（与 02_项目背景知识.md 的约定一致）。"""
    return "<10" if 1 <= n <= 9 else f"{n:,}"


def bar_chart(slide, x, y, w, h, cats, series, vmax, fmt, name, label_size=11.5, cat_size=12, gap=55,
              overlap=-10, point_colors=None, point_labels=None):
    """横向条形图（全 deck 同一种标签样式）。

    cats 自上而下写；series 为 [(名称, 数值, 颜色)]，同一类别里第一个系列画在最上面。
    point_colors：{(系列序号, 类别序号): 颜色}；point_labels：{(系列序号, 类别序号): 文字}（如 "<10"）。
    """
    cd = CategoryChartData()
    # 条形图自下而上排类别、同一类别内第一个系列画在下方：类别与系列都倒序写入
    cd.categories = list(reversed(cats))
    for sname, vals, _ in reversed(series):
        cd.add_series(sname, list(reversed(vals)))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(cat_size)
    ch.font.name = FONT
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    plot.vary_by_categories = False
    n = len(cats)
    for k, ser in enumerate(plot.series):
        si = len(series) - 1 - k           # 换回 series 里的序号
        color = series[si][2]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(color)
        ser.format.line.fill.background()
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(label_size)
        dl.font.bold = True
        dl.font.color.rgb = rgb(NAVY)
        for ci in range(n):
            pi = n - 1 - ci                # 换回 cats 里的序号
            pc = (point_colors or {}).get((si, pi))
            if pc:
                pt = ser.points[ci]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(pc)
                pt.format.line.fill.background()
            pl = (point_labels or {}).get((si, pi))
            if pl:
                pdl = ser.points[ci].data_label
                pdl.has_text_frame = True
                pdl.text_frame.text = pl
                pdl.position = XL_LABEL_POSITION.OUTSIDE_END
                for p in pdl.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(label_size)
                        r.font.bold = True
                        r.font.color.rgb = rgb(NAVY)
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = vmax
    hide_value_axis(va)
    ca = ch.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.color.rgb = rgb(BODY)
    chart_fonts(ch)
    return ch


def legend(slide, x, y, items, size=11.5, sw=0.18, gap=0.28):
    """自绘图例：[(文字, 颜色)]，从 x 起横排。返回结束位置。"""
    for text, color in items:
        add_rect(slide, x, y + 0.08, sw, 0.14, fill=color)
        tw = text_width(text, size) + 0.18
        add_text(slide, x + sw + 0.05, y, tw, 0.3, text, size=size, color=BODY, anchor="m", margins=(0, 0, 0, 0))
        x += sw + 0.05 + tw + gap
    return x


def pill(slide, x, y, text, fill, color=WHITE, size=11, h=0.28, bold=True, line=None):
    w = text_width(text, size) * 1.06 + 0.3
    s = add_rect(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    shape_text(s, text, size, color, bold=bold, wrap=False)
    return w


def num_badge(slide, cx, cy, text, fill=DARK, d=0.42, size=14):
    o = add_oval(slide, cx, cy, d, fill=fill)
    shape_text(o, text, size, WHITE, bold=True, margins=(0, 0, 0, 0))
    return o


def callout(slide, x, y, w, h, big, color, runs, big_w=1.25, big_size=28, size=12.5):
    """白底卡片：左边大数字，右边一段说明（runs 为 [(文字, 选项)]）。"""
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, x + 0.1, y, big_w, h, big, size=big_size, color=color, bold=True, align="c", anchor="m",
             margins=(0, 0, 0, 0))
    add_line(slide, x + big_w + 0.2, y + 0.2, x + big_w + 0.2, y + h - 0.2, BORDER, 1)
    add_text(slide, x + big_w + 0.35, y, w - big_w - 0.5, h, [runs], size=size, color=BODY, anchor="m",
             margins=(0, 0, 0, 0))


def num_tile(slide, x, y, w, h, big, label, color=DARK, line=BORDER, big_size=24, label_size=11.5):
    """白底小卡片：上面大数字，下面一行说明；color / line 用来点出要强调的那一格。"""
    add_rect(slide, x, y, w, h, fill=WHITE, line=line, line_w=1.5 if line != BORDER else 0.75, shadow=True)
    add_text(slide, x + 0.16, y + 0.06, w - 0.32, h * 0.52, big, size=big_size, color=color, bold=True,
             anchor="b", margins=(0, 0, 0, 0))
    add_text(slide, x + 0.16, y + 0.06 + h * 0.52, w - 0.32, h * 0.44 - 0.1, label, size=label_size,
             color=MUTED, anchor="t", margins=(0, 0.02, 0, 0))


# ---------------------------------------------------------------------------
# 演讲者备注（汇报人口吻；每页标了建议时长，合计约 17 分钟）
# ---------------------------------------------------------------------------
NOTES = {
    "cover": "封面，约 0.3 分钟。全程约 17 分钟（不含讨论），四部分：近期进展与发现、数据集构建论文、GD 临床答复、"
             "下一步。全部数字为聚合统计（1–9 记为 <10），来自 10-24 的 std_v3 导出和 GD 的答复文件。",
    "summary": "约 1 分钟。先给结论：数据管线已经能在真实数据上跑，临床也给了答复；下一步是把答复变成口径，并补足门急诊数据。"
               "四个问题对应后面的页码。",
    "recent": "约 1.5 分钟。按时间讲：10-17 部署到专用服务器；10-18 读 GD 病历系统原始导出；10-19 至 10-21 性能改造和医生端"
              "体征图；10-22、10-23 规则集 v0.3 与统一导出 std_v3；10-24 在服务器导入两院真实数据，当天算出宽表和问题单的计数"
              "（青色一列）；10-25 GD 答复问题单并交回证型对应表；同日完成数据集构建调研。六模型实验 10-21、10-22 完成本人"
              "负责的三个模型。下面四张卡是交付物：6–10% 是删去 40% 观测时、五个模型个体预测的翻转率（5 个 seed 均值）。",
    "data": "约 1.5 分钟。流程：两院原始导出 → flucase 导入 → 规则集 v0.3 分级 → std_v3 统一导出 → 本机宽表脚本只出聚合报告"
            " → 问题单发给 GD 临床负责人。左边是规模：主分析人群（确诊且 ≥60 岁）GD 328、GAM 住院 129、GAM 门急诊 17。"
            "右边两个发现：一是门急诊只有 17 次，而主结局 O3 是门急诊早期进展（D9），样本主要得从新的门急诊数据来；"
            "二是诊断和检测对不上：GD 有 290 次住院诊断写了流感、却没有阳性检测，GAM 门急诊反过来有 219 次阳性检测没写流感"
            "（另有 GAM 住院 44 次），所以病例以检测为准（D16），临床建议再补 NGS 报告。",
    "sitediff": "约 2 分钟。比例的分母是各院全部住院（GD 618、GAM 173）。新开始吸氧 GD 522（84%）对 GAM 13（8%）；"
                "无创 191 对 14；进 ICU 183 对 13；高流量 119 对 40；有创 73 对 16。差异集中在吸氧、无创、ICU 三项，"
                "有创和高流量两院接近。右边三条要连起来看：GD 给氧的 522 次里 428 次在入院 24 小时内，给氧前后能找到低氧证据的 133 次；"
                "但 GD 的 SpO2 只有 189 / 28,655 条推断为吸空气（按吸氧医嘱推断），“吸空气 SpO2 ≤93%”这条几乎用不上，"
                "所以“无低氧”里有一部分可能是测量状态推断造成的，不能直接说成没低氧就给氧。ICU 183 次里同时有器官衰竭的 45 次，"
                "ICU 怎么认见第 9 页。",
    "route": "约 1.5 分钟。七站：看范本、合到一张表、定义转重、查质量、收数据、样本量与验证、接基础模型；箭头上是读完一站带到下一站的问题。"
             "加粗的是必读 13 篇。讲解页 paper/数据集构建调研_20261025/论文路径讲解图_数据集构建.html 每篇都写了设计要点（带页码）"
             "和对本项目的改法；逐篇笔记在设计要点笔记.md。28 号 Riley 2021（外部验证样本量）只读了摘要。",
    "numbers": "约 1.5 分钟。数字都出自论文正文（见设计要点笔记）：Sherman（AMIA 2017；arXiv 1811.12520 §4.1）院内死亡任务，"
               "以结局为锚的方案在同法测试集上 .963，在与结局无关的测试集上 .831，以入院为起点的方案 .882；Lauritsen 2021 p3–4，"
               "AUROC；Agniel 2018（BMJ）233/272 项检验“做没做”与死亡相关，76/210 项的最强预测因子是检验间隔；4CE 2020 p2、p7；"
               "Riley 2020 Example 1（结局 0.05、30 个参数、R²cs 0.05）；Vassar 2013。",
    "apply": "约 1.5 分钟。左上是讲解页里 CLIF 的卡片截图：左边写论文做法，右边写对本项目的改法；左下的呼吸支持长表是虚构示例，"
             "说明结构。右边四类改动都不改转重定义，可以先在虚构数据上做；数据层两项（时间来源、治疗来源类型）本周就能开始。",
    "answers": "约 2 分钟。GD 临床负责人 13 题都答了：第 1–5 题和第 13 题、附件 1、2 用 16 条批注；第 7–12 题在意见栏打 √ 或写明"
               "（第 10 题勾“不排除”）；附件 2 把 12 个科内有监护室的科室标黄；第 6 题另交了证型对应表。同意现行的 6 项："
               "新开始吸氧算重型、进 ICU 即危重、术后转入 ICU 不排除、机械通气含无创与高流量、年龄分层、SpO2 按医嘱推断。"
               "要改的 5 处按影响排在表里：ICU 影响最大（标黄科室 GD 297 次住院现行都不算 ICU）；ICU、再住院两处和文献做法一致。"
               "这些都改变定义（A 档），讨论后登记为新决定再改规则。",
    "tcm": "约 1 分钟。证型按方案分三期：轻症 4 型（风热犯卫、风寒束表、表寒里热、热毒袭肺）、重症 2 型（毒热壅盛、"
           "毒热内陷内闭外脱）、恢复期 1 型（气阴两虚正气未复）。条形是按写法累加的就诊次数（一次就诊可写多个证候）。"
           "现在系统只认方案原名（tcm.py），这张表进来后按对应表认。证候只作研究变量，不进分级（D65）。",
    "decide": "约 2 分钟。前五件改变口径定义，是需要拍板的；第六件是协调数据。①② 影响重型、危重型判定最多；"
              "② 要结合第 5 页的 SpO2 推断问题看；④ 真实影响很小；⑤ 临床已给明确意见，确认即可；"
              "⑥ 门急诊数据决定主结局 O3 能不能做，检索条件用 GD 第 3 题的答复。定下后登记为新决定（D76 起）、出规则集 v0.4。",
    "plan": "约 1.5 分钟。日期是建议。本周：讨论确定口径、出 v0.4 候选；GAM 补充数据先做只出计数的扫描；证型拆分、未填的发回 GD；"
            "按检索条件申请门急诊数据。下周：两院按 v0.3 / v0.4 重算对比，GAM 重新导入。11 月中口径定稿、医生试用、"
            "Healthsheet 初稿与样本量估算；12 月流感高峰前上线前瞻收集。六模型实验等合作同学的 OPL-MT-MNAR 结果。",
    "closing": "结束页。需要讨论的六件事见第 11 页，下一步计划见第 12 页。",
}


# ---------------------------------------------------------------------------
# 各页内容
# ---------------------------------------------------------------------------
def build_summary(slide):
    frame(slide, "要点", "四个问题与一句话回答",
          "一句话：数据管线已能跑真实数据；下一步把临床答复变成口径，并补足门急诊数据。",
          ask_lead="数据说明：", ask="本汇报只含聚合统计（1–9 记为 <10）；真实病例只在服务器和本机处理，不上传在线大模型。",
          notes=NOTES["summary"])
    items = [
        ("01", "近期做成了什么？",
         "两院真实数据第一次按统一口径跑通：GD 住院 618、GAM 住院 173、门急诊 1,166 次。", "第 3–4 页"),
        ("02", "发现了什么问题？",
         "两院判据检出差异集中在吸氧、无创和 ICU；门急诊里符合主分析的只有 17 次。", "第 4–5 页"),
        ("03", "论文读出了什么？",
         "34 篇：t0 与窗口先写死、标明记录来源、ICU 看支持措施、按事件数算样本量。", "第 6–8 页"),
        ("04", "临床答复与下一步？",
         "GD 13 题：6 项同意、5 处改口径；先定口径出 v0.4，并行处理 GAM 新数据与证型。", "第 9–12 页"),
    ]
    h, gap = 0.92, 0.14
    for i, (num, q, a, page) in enumerate(items):
        y = 1.86 + i * (h + gap)
        add_rect(slide, 0.72, y, 11.9, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.12, shadow=True)
        nb = add_rect(slide, 0.95, y + 0.16, 0.72, 0.60, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        shape_text(nb, num, 20, WHITE, bold=True)
        add_text(slide, 1.92, y + 0.07, 9.0, 0.42, q, size=18, color=NAVY, bold=True, anchor="m")
        add_text(slide, 1.92, y + 0.49, 9.15, 0.36, a, size=14, color=BODY, anchor="m")
        add_line(slide, 11.18, y + 0.2, 11.18, y + h - 0.2, BORDER, 1)
        add_text(slide, 11.22, y, 1.32, h, page, size=13, color=MID, bold=True, align="c", anchor="m")


def build_recent(slide):
    frame(slide, "近期", "近两周：系统跑上真实数据，临床与文献同步推进",
          "10-24 第一次在真实数据上出计数；问题单得到 GD 答复；读完 34 篇数据集构建论文。",
          ask_lead="新到、尚未处理：", ask="GAM 补充的原始数据；GD 问题单答复与证型对应表（计划见第 12 页）。",
          notes=NOTES["recent"])
    days = [("10-17", "周六"), ("10-18", "周日"), ("10-19", "周一"), ("10-20", "周二"), ("10-21", "周三"),
            ("10-22", "周四"), ("10-23", "周五"), ("10-24", "周六"), ("10-25", "周日")]
    x0 = CL + 1.5
    dw = (CL + CW - x0) / len(days)
    lanes = [("系统与数据", 2.33), ("临床对接", 2.95), ("文献", 3.57), ("六模型实验", 4.19)]
    lh = 0.56
    hx = x0 + 7 * dw
    add_rect(slide, hx + 0.02, 1.78, dw - 0.04, 4.80 - 1.78, fill=CYAN_BG, name="10-24 底色")
    for i, (d, wd) in enumerate(days):
        add_text(slide, x0 + i * dw, 1.8, dw, 0.44,
                 [{"runs": d, "size": 12.5, "bold": True, "color": NAVY, "align": "c"},
                  {"runs": wd, "size": 10, "color": MUTED, "align": "c"}],
                 anchor="m", margins=(0, 0, 0, 0))
    add_line(slide, x0, 2.27, CL + CW, 2.27, BORDER, 1)
    for name, y in lanes:
        b = add_rect(slide, CL, y, 1.38, lh, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        shape_text(b, name, 12, NAVY, bold=True)

    def ev(lane, d0, span, text, style="plain"):
        y = lanes[lane][1]
        x = x0 + d0 * dw + 0.05
        w = span * dw - 0.10
        if style in ("key", "new"):
            s = add_rect(slide, x, y + 0.04, w, lh - 0.08, fill=DARK if style == "key" else CYAN,
                         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, shadow=True)
            shape_text(s, text, 11, WHITE, bold=True)
        else:
            s = add_rect(slide, x, y + 0.04, w, lh - 0.08, fill=WHITE, line=MID, line_w=1.0,
                         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
            shape_text(s, text, 11, NAVY)

    ev(0, 0, 1, "部署到服务器")
    ev(0, 1, 1, "读 GD\n原始导出")
    ev(0, 2, 3, "性能改造 · 医生端体征图 · 3 次部署")
    ev(0, 5, 2, "规则集 v0.3 · 统一导出 std_v3")
    ev(0, 7, 1, "导入两院\n真实数据", "key")
    ev(1, 7, 1, "宽表计数\n→ 问题单", "key")
    ev(1, 8, 1, "GD 答复\n证型对应表", "new")
    ev(2, 6, 1, "时序数据集\n讲解页")
    ev(2, 8, 1, "数据集构建\n34 篇", "new")
    ev(3, 4, 2, "3 个模型 × 5 seed · 讨论稿")
    add_text(slide, hx - 0.2, 4.8, dw + 0.4, 0.2, "第一次真实计数", size=10, color=CYAN, bold=True, align="c",
             anchor="m", margins=(0, 0, 0, 0))

    cards = [("2 院", "flucase 读通两院真实导出", "10-24 导入 · 207 个测试通过"),
             ("A1–A10", "宽表与聚合报告脚本", "问题单的数据位自动算好"),
             ("6–10%", "删 40% 观测时个体预测翻转", "六模型实验 · 本人 3 模型 × 5 seed"),
             ("34 篇", "数据集构建论文 · 7 站", "13 篇必读 · 讲解页带页码")]
    gap = 0.2
    w = (CW - 3 * gap) / 4
    y, h = 5.05, 1.22
    for i, (big, lab, sub) in enumerate(cards):
        x = CL + i * (w + gap)
        add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
        add_text(slide, x + 0.18, y + 0.08, w - 0.36, 0.46, big, size=22, color=DARK, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.18, y + 0.56, w - 0.36, 0.3, lab, size=12, color=NAVY, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.18, y + 0.86, w - 0.36, 0.28, sub, size=10.5, color=MUTED, anchor="m",
                 margins=(0, 0, 0, 0))


def build_data(slide):
    frame(slide, "成果", "真实数据第一次按统一口径跑通",
          "10-24 在服务器导入两院数据，按规则集 v0.3 分级并统一导出，再在本机算出宽表与计数。",
          ask_lead="含义：", ask="按既定安排，现有住院数据只做住院期重症的描述；预测主结局的样本要从新的门急诊数据来。",
          notes=NOTES["data"])
    steps = [("两院原始导出", "GD 一套 20 个 xlsx\nGAM 住院、门急诊"),
             ("flucase 导入", "服务器 · 10-24\n问题逐条确认"),
             ("规则集 v0.3", "分级时间线\n候选结局"),
             ("std_v3 导出", "两院同表、同编码\n同取值"),
             ("宽表 + 计数", "本机脚本\n只出聚合量"),
             ("问题单 → GD", "13 题 + 3 个附件\n10-25 答复")]
    step = CW / len(steps)
    for i, (t, sub) in enumerate(steps):
        x = CL + i * step
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        s = add_rect(slide, x, 1.84, step + 0.06, 0.5, fill=CYAN if i >= 4 else DARK, shape=shp)
        shape_text(s, t, 12, WHITE, bold=True, margins=(0.1 if i else 0.04, 0, 0.06, 0))
        add_text(slide, x + 0.12, 2.38, step - 0.2, 0.46, sub, size=10.5, color=MUTED, align="c",
                 margins=(0, 0, 0, 0))

    lw = 6.3
    rows = [("导入的就诊", [("618", "GD 住院", DARK, BORDER), ("173", "GAM 住院", DARK, BORDER),
                         ("1,166", "GAM 门急诊", DARK, BORDER)]),
            ("确诊且 ≥60 岁（主分析人群；括号内为 ≥65 岁）",
             [("328", "GD 住院（298）", DARK, BORDER), ("129", "GAM 住院（113）", DARK, BORDER),
              ("17", "GAM 门急诊（17）", AMBER, AMBER)])]
    tw = (lw - 2 * 0.18) / 3
    for r, (label, tiles) in enumerate(rows):
        y = 3.0 + r * 1.55
        add_text(slide, CL, y, lw, 0.3, label, size=13, color=NAVY, bold=True, anchor="m", margins=(0, 0, 0, 0))
        for i, (big, lab, col, ln) in enumerate(tiles):
            num_tile(slide, CL + i * (tw + 0.18), y + 0.34, tw, 1.05, big, lab, color=col, line=ln)

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    callout(slide, rx, 3.0, rw, 1.4, "17", AMBER,
            [("GAM 门急诊里确诊且 ≥60 岁的只有 17 次。", {"bold": True}),
             ("主结局 O3（门急诊早期进展）要靠新的门急诊、发热门诊数据。", {})])
    callout(slide, rx, 4.55, rw, 1.4, "290", RED,
            [("GD 诊断写了流感、却没有阳性检测的住院；", {"bold": True}),
             ("GAM 门急诊反过来有 219 次阳性、诊断没写流感。病例以检测为准，临床建议补上 NGS。", {})])


def build_sitediff(slide):
    frame(slide, "发现", "两院判据：差异集中在吸氧、无创和 ICU",
          "同一套 v0.3 规则：GD 新开始吸氧 84%、GAM 8%；有创通气、高流量两院接近。",
          ask_lead="影响：", ask="主口径 primary 以治疗类判据为主，两院可比性风险最大；病情差异还是记录方式差异，要拆开核实。",
          notes=NOTES["sitediff"])
    lw = 6.95
    add_rect(slide, CL, 1.82, lw, 4.45, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.9, 4.6, 0.36, "各判据占该院住院的比例", size=15, color=NAVY, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    legend(slide, CL + lw - 2.1, 1.93, [("GD", GD_C), ("GAM", GAM_C)])
    cats = ["新开始吸氧", "无创通气", "进入 ICU", "高流量氧疗", "有创通气"]
    gd = [522 / 618 * 100, 191 / 618 * 100, 183 / 618 * 100, 119 / 618 * 100, 73 / 618 * 100]
    gam = [13 / 173 * 100, 14 / 173 * 100, 13 / 173 * 100, 40 / 173 * 100, 16 / 173 * 100]
    bar_chart(slide, CL + 0.1, 2.3, lw - 0.25, 3.55, cats, [("GD", gd, GD_C), ("GAM", gam, GAM_C)],
              vmax=100, fmt='0"%"', name="判据比例条形图", gap=60, overlap=-5)
    add_text(slide, CL + 0.2, 5.86, lw - 0.4, 0.34, "分母：GD 住院 618、GAM 住院 173（10-24 导出，规则集 v0.3）",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    ch, cg = 1.3, 0.14
    cards = [
        ("25%", RED, [("GD 给氧的 522 次住院（428 次在入院 24 小时内），", {}),
                      ("给氧前后 24 小时有低氧*证据的只有 133 次", {"bold": True})]),
        ("<1%", AMBER, [("GD 的 SpO2 只有不到 1% 推断为吸空气：", {"bold": True}),
                        ("“吸空气 SpO2 ≤93%”几乎用不上，上一条要谨慎解读", {})]),
        ("25%", RED, [("GD 进 ICU 的 183 次住院中，", {}), ("同时有器官衰竭的 45 次", {"bold": True}),
                      ("；ICU 怎么认见第 9 页", {})]),
    ]
    for i, (big, col, runs) in enumerate(cards):
        callout(slide, rx, 1.82 + i * (ch + cg), rw, ch, big, col, runs, big_w=1.15, size=12.5)
    add_text(slide, rx, 1.82 + 3 * ch + 2 * cg + 0.06, rw, 0.36,
             "* 低氧：吸空气 SpO2 ≤93%、PaO2/FiO2 ≤300 或呼吸 ≥30 次/分。", size=10.5, color=MUTED,
             anchor="m", margins=(0, 0, 0, 0))


def build_route(slide):
    frame(slide, "文献", "数据集构建：34 篇论文分 7 站读",
          "每一站回答数据集构建的一个问题，读完带着下一个问题进下一站；加粗的 13 篇是必读。",
          ask_lead="可回查：", ask="讲解页与逐篇笔记都标了页码或章节；28 号（外部验证样本量）只读了摘要，已标注。",
          notes=NOTES["route"])
    stations = [
        (1, "看范本", "4 篇 · 必读 2", [("MIMIC-IV", 1), ("eICU", 0), ("PIC", 1), ("MIMIC-IV-ED", 0)]),
        (2, "合到一张表", "6 篇 · 必读 2", [("OMOP", 0), ("N3C", 1), ("4CE", 0), ("CLIF", 1), ("ricu", 0), ("MEDS", 0)]),
        (3, "定义转重", "8 篇 · 必读 4", [("Sherman", 1), ("Lauritsen", 1), ("Agniel", 0), ("Beaulieu-Jones", 1),
                                    ("WHO 量表", 0), ("ISARIC 4C", 1), ("PhysioNet 2019", 0), ("ACES", 0)]),
        (4, "查质量", "4 篇 · 必读 1", [("Weiskopf", 0), ("Kahn", 1), ("OHDSI DQD", 0), ("Groenwold", 0)]),
        (5, "收数据", "4 篇 · 必读 1", [("ISARIC CCP", 1), ("FluSurv-NET", 0), ("GIHSN", 0), ("Vassar", 0)]),
        (6, "样本量与验证", "7 篇 · 必读 3", [("Riley 开发", 1), ("Riley 验证", 0), ("Debray", 1), ("Nestor", 0),
                                       ("TRIPOD+AI", 0), ("PROBAST+AI", 0), ("Healthsheet", 1)]),
        (7, "接基础模型", "1 篇", [("EHRSHOT", 0)]),
    ]
    gx = 0.85
    bw = (CW - 3 * gx) / 4
    bh = 1.86
    ys = (1.86, 4.36)
    # 第一行 1→4 从左到右；第二行 5、6、7 从右到左
    pos = {1: (0, 0), 2: (0, 1), 3: (0, 2), 4: (0, 3), 5: (1, 3), 6: (1, 2), 7: (1, 1)}

    def box_xy(k):
        r, c = pos[k]
        return CL + c * (bw + gx), ys[r]

    for k, name, sub, papers in stations:
        x, y = box_xy(k)
        add_rect(slide, x, y, bw, bh, fill=WHITE, line=MID, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.06, shadow=True)
        num_badge(slide, x + 0.36, y + 0.36, str(k), fill=DARK, d=0.44, size=15)
        add_text(slide, x + 0.68, y + 0.12, bw - 0.8, 0.34, name, size=16, color=NAVY, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.68, y + 0.44, bw - 0.8, 0.26, sub, size=11, color=MUTED, anchor="m",
                 margins=(0, 0, 0, 0))
        runs = []
        for j, (p, must) in enumerate(papers):
            if j:
                runs.append((" · ", {"color": GREY}))
            runs.append((p, {"bold": bool(must), "color": DARK if must else BODY}))
        add_text(slide, x + 0.16, y + 0.8, bw - 0.3, bh - 0.88, [runs], size=11.5, line_spacing=1.15,
                 margins=(0, 0, 0, 0))

    def harrow(k1, k2, label):
        x1, y1 = box_xy(k1)
        x2, _ = box_xy(k2)
        ya = y1 + bh / 2 + 0.12
        if x2 > x1:
            add_line(slide, x1 + bw + 0.06, ya, x2 - 0.06, ya, MID, 1.75, arrow_end=True)
            lx = x1 + bw
        else:
            add_line(slide, x1 - 0.06, ya, x2 + bw + 0.06, ya, MID, 1.75, arrow_end=True)
            lx = x2 + bw
        add_text(slide, lx - 0.05, ya - 0.42, gx + 0.1, 0.34, label, size=11, color=AMBER, bold=True, align="c",
                 anchor="b", margins=(0, 0, 0, 0))

    harrow(1, 2, "多家怎么合")
    harrow(2, 3, "转重怎么定")
    harrow(3, 4, "数据可信吗")
    harrow(5, 6, "要收多少")
    harrow(6, 7, "接基础模型")
    x4, y4 = box_xy(4)
    ax = x4 + bw / 2
    add_line(slide, ax, y4 + bh + 0.05, ax, ys[1] - 0.05, MID, 1.75, arrow_end=True)
    add_text(slide, ax + 0.1, y4 + bh + 0.08, 1.6, 0.44, "新数据怎么收", size=11, color=AMBER, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    x0, y0 = CL, ys[1]
    add_rect(slide, x0, y0, bw, bh, fill=PALE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    add_text(slide, x0 + 0.18, y0 + 0.12, bw - 0.36, bh - 0.24, [
        {"runs": "读法", "bold": True, "color": NAVY, "size": 13, "space_after": 4},
        {"runs": "沿箭头顺序走，先读每站加粗的必读，其余按问题回查。", "space_after": 6},
        {"runs": [("讲解页：", {"bold": True, "color": NAVY}),
                  ("paper/数据集构建调研_20261025/论文路径讲解图_数据集构建.html", {"color": MUTED})]},
    ], size=11.5, color=BODY, margins=(0, 0, 0, 0))


def build_numbers(slide):
    frame(slide, "文献", "论文里最值得记住的六个数字",
          "怎么取样本、怎么定窗口、医生做了什么，都会让模型看起来比实际好。",
          ask_lead="启示：", ask="样本小、只有两个中心时，这些偏差更容易放大——口径和窗口要在看结果之前定下来。",
          notes=NOTES["numbers"])
    items = [
        ("11 · Sherman 2017", ".963 → .831", "以结局为锚：同法测试 .963，换成与结局无关的测试集只剩 .831（以入院为起点：.882）",
         "t0 固定在首次就诊 / 首次阳性，不按“首次重型时间”往回取"),
        ("12 · Lauritsen 2021", ".906 / .756", "同一批数据：“发病前固定 12 小时”框架 AUROC .906，换成全住院滑动窗口只有 .756",
         "O3 写明是“首次就诊触发的单点预测”"),
        ("13 · Agniel 2018", "86%", "272 项检验中 86% 的“做没做”就与死亡相关；36% 的检验里最强因子是检验间隔",
         "测量次数、间隔单列为过程变量，跨院时查它是否只在认医院"),
        ("07 · 4CE 2020", "13%", "14 项检验各定了编码与单位，只有 13% 的站点交齐；ICU 照护级别难识别、不可靠",
         "我们的 ICU 认定也卡在科室名上"),
        ("27 · Riley 2020", "263 个事件", "例：结局 5%、30 个候选参数，至少要 5,249 人、263 个事件；机器学习还要多得多",
         "按我们的结局率用 pmsampsize 算，再定变量个数"),
        ("26 · Vassar 2013", "κ ≥ 0.6", "病历摘录要标准表单与手册，两人独立摘录、报 κ，先做约 10% 的预试",
         "复核与规则抽取补一份一致性报告"),
    ]
    gap = 0.22
    w = (CW - 2 * gap) / 3
    h = (6.27 - 1.82 - gap) / 2
    for i, (tag, big, finding, ours) in enumerate(items):
        x = CL + (i % 3) * (w + gap)
        y = 1.82 + (i // 3) * (h + gap)
        add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
        pill(slide, x + 0.18, y + 0.14, tag, LIGHT, color=NAVY, size=11)
        add_text(slide, x + 0.18, y + 0.44, w - 0.36, 0.48, big, size=26, color=DARK, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.18, y + 0.94, w - 0.36, 0.56, finding, size=11.5, color=BODY, margins=(0, 0, 0, 0))
        add_rect(slide, x + 0.12, y + h - 0.6, w - 0.24, 0.5, fill=PALE)
        add_text(slide, x + 0.2, y + h - 0.6, w - 0.4, 0.5, [[("对我们：", {"bold": True, "color": NAVY}),
                                                            (ours, {})]],
                 size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))


def build_apply(slide):
    frame(slide, "文献", "文献落到项目里：四类改动，大多不改定义",
          "范本的共同做法：保留原名、标明来源、窗口写死、先算样本量——flucase 方向对，还缺几块。",
          ask_lead="说明：", ask="这些都是技术改动，可以先在虚构数据上做；改变定义的放到第 11 页讨论。",
          notes=NOTES["apply"])
    lw = 5.5
    ih = lw / 3.173
    add_text(slide, CL, 1.82, lw, 0.3, [[("讲解页里的一张卡片", {"bold": True, "color": NAVY}),
                                         ("：左边论文做法，右边对本项目", {"color": MUTED})]],
             size=12, anchor="m", margins=(0, 0, 0, 0))
    add_rect(slide, CL - 0.03, 2.16, lw + 0.06, ih + 0.06, fill=WHITE, line=BORDER, shadow=True)
    pic = slide.shapes.add_picture(str(CARD_CLIF), Inches(CL), Inches(2.19), Inches(lw), Inches(ih))
    pic.name = "讲解页卡片截图：CLIF"
    ty = 2.19 + ih + 0.24
    add_text(slide, CL, ty, lw, 0.3, [[("例：“呼吸支持”长表", {"bold": True, "color": NAVY}),
                                       ("（虚构示例，参照 CLIF 08、WHO 量表 15）", {"color": MUTED})]],
             size=12, anchor="m", margins=(0, 0, 0, 0))
    rows = [["设备类别", "开始", "结束", "医院原文"],
            ["鼻导管", "第 1 天 09:10", "第 3 天 16:00", "鼻导管吸氧 2L/min"],
            ["高流量", "第 3 天 16:05", "第 5 天 10:00", "经鼻高流量 FiO2 40%"],
            ["无创通气", "第 5 天 10:20", "—", "BiPAP 辅助通气"]]
    add_table(slide, CL, ty + 0.36, [1.0, 1.4, 1.4, 1.7], [0.34, 0.33, 0.33, 0.33], rows, size=11.5,
              header_size=11.5)

    rx, rw = CL + lw + 0.3, CW - lw - 0.3
    groups = [
        ("数据层", "flucase · 导出", ["事件时间标来源（测量 / 医嘱 / 书写 / 推断）；治疗标医嘱 / 执行 / 提及",
                                  "新增呼吸支持长表；概念字典与原名映射随导出发布"],
         "01 · 02 · 03\n08 · 09", "现在可做"),
        ("任务定义", "t0 · 窗口 · 标签", ["任务写成“触发 + 窗口 + 排除”，结局窗口与中心无关、单独成文件",
                                     "t0 时已达结局的不进主分析；医生行为特征只作对照"],
         "11 · 14 · 16\n17 · 18", "口径定稿时"),
        ("质控与收集", "宽表 · 复核 · 前瞻", ["检查带阈值、出质量成绩单；规则抽取两人独立标注、报 κ",
                                       "前瞻按固定时点采集，加起病前功能状态"],
         "20 · 21 · 23\n25 · 26", "前瞻表单定稿前"),
        ("样本量与验证", "建模 · 文档", ["按事件数算样本量（pmsampsize）；外部验证先比较人群差异",
                                    "新季节按时间切分、只评一次；写中文 Healthsheet"],
         "27 · 29\n30 · 33", "建模前"),
    ]
    gh, gg = 1.03, 0.11
    for i, (name, sub, items, refs, when) in enumerate(groups):
        y = 1.82 + i * (gh + gg)
        add_rect(slide, rx, y, rw, gh, fill=WHITE, line=BORDER, shadow=True)
        lb = add_rect(slide, rx, y, 1.3, gh, fill=LIGHT)
        shape_text(lb, [{"runs": name, "bold": True, "size": 13, "color": NAVY},
                        {"runs": sub, "size": 10, "color": MUTED}], 13, NAVY)
        add_text(slide, rx + 1.45, y, rw - 3.0, gh, bullets(items, gap=3), size=11.5, color=BODY, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, rx + rw - 1.45, y, 1.35, gh,
                 [{"runs": refs, "size": 10.5, "bold": True, "color": MID, "align": "r", "space_after": 4},
                  {"runs": when, "size": 10.5, "bold": True, "color": CYAN if when == "现在可做" else MUTED,
                   "align": "r"}],
                 anchor="m", margins=(0, 0, 0, 0))


ANSWER_KIND = {"同意现行": DARK, "补细则": MID, "改口径": AMBER, "新材料": CYAN}


def build_answers(slide):
    frame(slide, "临床答复", "GD 临床答复：6 项同意，5 处要改口径",
          "13 题都有答复；要改的 5 处里 ICU 影响最大，ICU、再住院两处与文献做法一致。",
          ask_lead="另外：", ask="第 3 题给了门急诊检索条件（诊断 + 症状 + 检验），可用来申请门急诊、发热门诊数据；第 6 题见下一页。",
          notes=NOTES["answers"])
    chips = [(1, "发热", "补细则"), (2, "慢性低氧", "改口径"), (3, "检索", "新材料"), (4, "确诊", "补细则"),
             (5, "再住院", "改口径"), (6, "证型", "新材料"), (7, "新开始吸氧", "同意现行"), (8, "进 ICU", "同意现行"),
             (9, "ICU 科室", "改口径"), (10, "术后 ICU", "同意现行"), (11, "无创高流量", "同意现行"),
             (12, "年龄", "同意现行"), (13, "SpO2", "同意现行")]
    cgap = 0.06
    cw = (CW - 12 * cgap) / 13
    for i, (num, lab, kind) in enumerate(chips):
        x = CL + i * (cw + cgap)
        s = add_rect(slide, x, 1.8, cw, 0.56, fill=ANSWER_KIND[kind], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        shape_text(s, [{"runs": f"第 {num} 题", "size": 9.5}, {"runs": lab, "size": 11, "bold": True}], 11, WHITE,
                   margins=(0.02, 0.02, 0.02, 0.02))
    lx = CL
    for kind, col in ANSWER_KIND.items():
        cnt = sum(1 for c in chips if c[2] == kind)
        lx = legend(slide, lx, 2.42, [(f"{kind} {cnt}", col)], size=11, gap=0.12)

    a = {"text": "—", "align": "c", "color": GREY}
    rows = [["题", "GD 的答复", "拟改口径", "10-24 导出的计数", "文献呼应"],
            ["第 9 题\nICU 科室", "科内有监护室的，按医嘱有无“重症监护”区分；EICU 保留",
             "科室名命中，或标黄科室且有重症监护医嘱", "标黄 12 个科室共 GD 297 次住院，现行都不算 ICU",
             "15 WHO 量表：看支持措施不看地点；07 4CE"],
            ["第 2 题\n慢性低氧", "平时血氧差异大，用血气判断呼吸衰竭；另加间质性肺病、哮喘、心衰",
             "慢性低氧扩到这三类，改用血气（阈值待定）", f"COPD：GD 148、GAM 23；吸空气 SpO2 ≤93% 两院都 {small(1)}", a],
            ["第 5 题\n再住院", "14 天内再住院期间的转重，算进第一次住院",
             "住院任务的标签窗含同一病程的再住院", f"14 天内再住院：GD 0、GAM {small(2)}",
             "23 ISARIC CCP · 18 ACES"],
            ["第 4 题\n确诊", "弱阳性算；加 NGS 报告；只写“甲 / 乙型流感病毒”的认作抗原",
             "检测对应表加 NGS，两项改为抗原", "弱阳性 GD 99 次就诊，其中 82 次只有这一条阳性",
             "06 N3C：检测证据分层"],
            ["第 1 题\n发热", "都按腋温；“昨日到今晨无发热”即退热；不发热者从感冒样症状起算",
             "体温不折算；退热加文字规则；起病按首个感冒样症状", "入院前后 24 小时不发热：GD 24%、GAM 67%", a]]
    for r in rows[1:]:
        r[0] = {"text": r[0], "bold": True, "color": NAVY}
    add_table(slide, CL, 2.82, [1.0, 3.05, 2.85, 3.28, 2.05], [0.36] + [0.616] * 5, rows, size=11.5,
              header_size=12, margins=(0.08, 0.03, 0.08, 0.03))


def build_tcm(slide):
    frame(slide, "临床答复", "中医证型：355 种写法已对到 7 型的 271 种",
          "GD 按 2025 年版诊疗方案的 7 个证型填了对应表；按就诊算，GD 已覆盖 85%，GAM 69%。",
          ask_lead="下一步：", ask="84 种未填的（多为痰、瘀、虚类）发回 GD 补填或标“不对应”；证候只作研究变量，不进分级。",
          notes=NOTES["tcm"])
    lw = 6.9
    add_rect(slide, CL, 1.82, lw, 4.45, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.88, 4.3, 0.34, "对到各证型的就诊次数", size=14, color=NAVY, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    legend(slide, CL + lw - 2.05, 1.92, [("GD", GD_C), ("GAM", GAM_C)], size=11)
    cats = ["风热犯卫（轻）", "风寒束表（轻）", "表寒里热（轻）", "热毒袭肺（轻）", "毒热壅盛（重）",
            "毒热内陷、内闭外脱（重）", "气阴两虚、正气未复（恢复）", "未填"]
    gd = [36, 31, 12, 352, 57, 19, 454, 163]
    gam = [0, 2, 2, 33, 60, 0, 103, 88]
    pl = {(1, k): small(v) for k, v in enumerate(gam) if 1 <= v <= 9}
    pc = {(0, 7): GREY, (1, 7): "B9C1CC"}
    bar_chart(slide, CL + 0.1, 2.26, lw - 0.25, 3.55, cats, [("GD", gd, GD_C), ("GAM", gam, GAM_C)], vmax=540,
              fmt="0", name="证型条形图", label_size=10.5, cat_size=11, gap=45, overlap=-5, point_colors=pc,
              point_labels=pl)
    add_text(slide, CL + 0.2, 5.82, lw - 0.4, 0.4,
             "按写法累加的就诊次数（一次就诊可写多个证候）；轻症 4 型、重症 2 型、恢复期 1 型。",
             size=10.5, color=MUTED, margins=(0, 0, 0, 0))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    tiles = [("355 种", "证候写法（两院去重）"), ("271 种", "GD 已填（76%）"), ("85% · 69%", "已覆盖就诊（GD · GAM）"),
             ("84 种", "待补")]
    tw = (rw - 0.18) / 2
    for i, (big, lab) in enumerate(tiles):
        num_tile(slide, rx + (i % 2) * (tw + 0.18), 1.82 + (i // 2) * 0.98, tw, 0.86, big, lab, big_size=20,
                 label_size=11)
    y = 3.86
    add_rect(slide, rx, y, rw, 6.27 - y, fill=PALE)
    add_text(slide, rx + 0.2, y + 0.1, rw - 0.4, 6.27 - y - 0.2, [
        {"runs": "处理要点", "bold": True, "color": NAVY, "size": 13, "space_after": 4},
    ] + bullets([
        "150 种写法是“¤”连起来的多个证候：先拆开去重，再对应",
        "系统现在只认方案原名：改为读这张对应表，原名保留、随导出发布",
        "待补最多：痰浊阻肺证（GD 60）、痰瘀互结证（GAM 31）",
        "一次就诊对到多个证型时怎么取，需要定口径",
    ], gap=4), size=11.5, color=BODY, margins=(0, 0, 0, 0))


def build_decide(slide):
    frame(slide, "讨论", "需要讨论确定的六件事",
          "前五件改变口径定义，列出选项和证据，标“参考”的是建议；第六件需要协调数据。",
          ask_lead="定下后：", ask="登记为新版本口径（规则集 v0.4），旧版本保留，两院重算后对比。",
          notes=NOTES["decide"])
    items = [
        ("① ICU 怎么认", "第 9 题", "标黄 12 个科室 GD 297 次住院，现行都不算 ICU；EICU 也有普通床位",
         [("A", "维持：只看科室名", False), ("B", "科室名，或标黄科室且有重症监护医嘱", True)],
         "需要：重症监护类医嘱的名称清单", DARK),
        ("② 新开始吸氧", "第 7 题", "临床同意；GD 给氧 522 次中有低氧证据的 25%，但 SpO2 多被推断为吸氧中",
         [("A", "维持代理判据", False), ("B", "要求给氧前后有低氧", False), ("C", "主口径维持，“伴低氧”作敏感性", True)],
         None, DARK),
        ("③ 慢性低氧者", "第 2 题", "临床建议用血气判呼吸衰竭；病种加间质性肺病、哮喘、心衰",
         [("定", "病种清单", False), ("定", "血气阈值，以及没有血气时怎么办", False)], None, DARK),
        ("④ 14 天内再住院", "第 5 题", f"临床：算进第一次住院；影响 GD 0、GAM {small(2)}",
         [("A", "不算（现行）", False), ("B", "算，出新版本标签", True)], None, DARK),
        ("⑤ 细则：确诊与发热", "第 4、1 题", "临床意见明确：弱阳性 GD 99 次；基线不发热 GD 24%、GAM 67%",
         [("√", "弱阳性算；加 NGS；两项认作抗原", False), ("√", "腋温不折算；退热看病程记录", False),
          ("√", "不发热者从感冒样症状起算", False)], None, DARK),
        ("⑥ 门急诊数据", "第 3 题", "门急诊符合主分析的只有 17 次，O3 缺样本",
         [("→", "按临床给的检索条件：诊断 + 症状 + 检验", False), ("→", "申请两院门急诊、发热门诊导出", False)],
         "需要协调：信息科", CYAN),
    ]
    gap = 0.22
    w = (CW - 2 * gap) / 3
    h = (6.27 - 1.82 - gap) / 2
    for i, (title, q, ev, opts, need, hcol) in enumerate(items):
        x = CL + (i % 3) * (w + gap)
        y = 1.82 + (i // 3) * (h + gap)
        add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
        add_rect(slide, x, y, w, 0.42, fill=hcol)
        add_text(slide, x + 0.16, y, w - 0.3, 0.42, [[(title, {"bold": True}), (f"  {q}", {"size": 11.5})]],
                 size=14.5, color=WHITE, anchor="m", margins=(0, 0, 0, 0))
        add_text(slide, x + 0.16, y + 0.48, w - 0.32, 0.5, ev, size=11, color=MUTED, margins=(0, 0, 0, 0))
        oy = y + 1.0
        for tag, text, ref in opts:
            fill, fg = (AMBER, WHITE) if ref else (AMBER_BG, AMBER)
            if hcol == CYAN:
                fill, fg = CYAN_BG, CYAN
            nb = add_rect(slide, x + 0.16, oy + 0.03, 0.3, 0.26, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          radius=0.3)
            shape_text(nb, tag, 10.5, fg, bold=True, margins=(0, 0, 0, 0), wrap=False)
            runs = [(text, {})]
            if ref:
                runs.append(("  参考", {"bold": True, "color": AMBER, "size": 10.5}))
            add_text(slide, x + 0.54, oy, w - 0.7, 0.34, [runs], size=11.5, color=BODY, anchor="m",
                     margins=(0, 0, 0, 0))
            oy += 0.36
        if need:
            add_text(slide, x + 0.16, y + h - 0.36, w - 0.32, 0.3, need, size=11, color=MID, bold=True, anchor="m",
                     margins=(0, 0, 0, 0))


def build_plan(slide):
    frame(slide, "下一步", "下一步：先定口径出 v0.4，赶在流感高峰前定稿",
          "关键路径：定口径 → v0.4 与 GAM 新数据 → 两院重算 → 口径定稿 → 流感高峰前上线前瞻收集。",
          ask_lead="11 月底前要有：", ask="v0.4 在两院的真实计数、定稿的口径与证型对应、Healthsheet 初稿、六模型汇总表。",
          notes=NOTES["plan"])
    cols = [("本周", "10-26 – 11-01"), ("下周", "11-02 – 11-08"), ("11 月中", "11-09 – 11-22"), ("12 月", "流感高峰")]
    lanes = ["口径 · GD 答复", "GAM 补充数据", "证型对应", "门急诊与前瞻", "文献落地", "六模型实验"]
    x0 = CL + 1.62
    cw = (CL + CW - x0) / 4
    for i, (c, d) in enumerate(cols):
        hdr = add_rect(slide, x0 + i * cw + 0.04, 1.82, cw - 0.08, 0.44, fill=RED if i == 3 else DARK)
        shape_text(hdr, [[(c, {"bold": True, "size": 13.5}), (f"  {d}", {"size": 10.5})]], 13.5, WHITE)
    ly0, lh, lg = 2.33, 0.6, 0.05
    for j, name in enumerate(lanes):
        y = ly0 + j * (lh + lg)
        b = add_rect(slide, CL, y, 1.5, lh, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        shape_text(b, name, 11.5, NAVY, bold=True)

    def bar(lane, c0, span, text, style="plain"):
        y = ly0 + lane * (lh + lg) + 0.04
        x = x0 + c0 * cw + 0.06
        w = span * cw - 0.12
        hh = lh - 0.08
        if style == "key":
            s = add_rect(slide, x, y, w, hh, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, shadow=True)
            shape_text(s, text, 11, WHITE, bold=True)
        elif style == "wait":
            s = add_rect(slide, x, y, w, hh, fill=WHITE, line=GREY, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                         radius=0.12, dash=MSO_LINE_DASH_STYLE.DASH)
            shape_text(s, text, 11, MUTED)
        elif style == "hot":
            s = add_rect(slide, x, y, w, hh, fill=RED_BG, line=RED, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                         radius=0.12)
            shape_text(s, text, 11, RED, bold=True)
        else:
            s = add_rect(slide, x, y, w, hh, fill=WHITE, line=MID, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                         radius=0.12)
            shape_text(s, text, 11, NAVY)

    bar(0, 0, 1, "讨论确定 5 处口径\n出规则集 v0.4 候选", "key")
    bar(0, 1, 1, "两院按 v0.3 / v0.4\n重算对比")
    bar(0, 2, 1, "口径定稿", "key")
    bar(1, 0, 1, "只出计数的扫描\n判断读取器要不要升级")
    bar(1, 1, 1, "重新导入、重导 std_v3\n重跑宽表与计数")
    bar(2, 0, 1, "拆开“¤”、去重\n未填的 84 种发回 GD")
    bar(2, 1, 1, "对应表进系统\n导出加证型列")
    bar(3, 0, 2, "按检索条件申请两院门急诊、发热门诊导出")
    bar(3, 2, 1, "医生试用\n前瞻表单定稿")
    bar(3, 3, 1, "上线前瞻收集", "hot")
    bar(4, 0, 1, "时间来源列\n治疗来源类型")
    bar(4, 1, 1, "呼吸支持长表\n质量成绩单")
    bar(4, 2, 1, "Healthsheet 初稿\n样本量估算")
    bar(5, 0, 2, "等合作同学的 OPL-MT-MNAR 结果", "wait")
    bar(5, 2, 1, "六模型汇总\n与配对比较")


def build_closing(slide):
    band = add_rect(slide, 0, 2.55, SW, 1.6, fill="4472C4", name="致谢横带")
    shape_text(band, "谢谢！欢迎讨论与指正", 40, WHITE)
    for r in band.text_frame.paragraphs[0].runs:
        rPr = r._r.get_or_add_rPr()
        eff = etree.Element(qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="50800", dist="38100", dir="5400000", algn="ctr",
                              rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="45000")
        rPr.find(qn("a:solidFill")).addnext(eff)
    add_text(slide, 1.0, 4.45, SW - 2.0, 0.5, "需要讨论的六件事见第 11 页 · 下一步计划见第 12 页",
             size=16, color=NAVY, align="c", anchor="m")
    slide.notes_slide.notes_text_frame.text = NOTES["closing"]


# ---------------------------------------------------------------------------
# 封面与删页（同 build_weekly_summary_20261023.py）
# ---------------------------------------------------------------------------
def set_shape_text(shape, text):
    """改写形状里第一段的第一个文字片段，删掉其余片段，保留原格式。"""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fix_cover(cover):
    for shp in cover.shapes:
        if not shp.has_text_frame:
            continue
        text = shp.text_frame.text.strip()
        if text.startswith("老年流感"):
            set_shape_text(shp, "老年流感转重研究：近期工作与论文阅读汇报")
        elif text.startswith("理想数据"):
            set_shape_text(shp, "近期进展 · 数据集构建文献 · 临床答复 · 下一步")
        elif text.startswith("2026"):
            set_shape_text(shp, "2026年10月26日")
    cover.notes_slide.notes_text_frame.text = NOTES["cover"]


def delete_slides_after_first(prs):
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst)[1:]:
        rid = sld_id.get(qn("r:id"))
        sld_id_lst.remove(sld_id)
        prs.part.drop_rel(rid)


def main():
    prs = Presentation(str(SRC))
    fix_cover(prs.slides[0])
    delete_slides_after_first(prs)
    blank = prs.slide_layouts.get_by_name("Blank")
    builders = [build_summary, build_recent, build_data, build_sitediff,
                build_route, build_numbers, build_apply,
                build_answers, build_tcm,
                build_decide, build_plan, build_closing]
    for build in builders:
        build(prs.slides.add_slide(blank))
    prs.core_properties.title = DECK_TITLE
    prs.core_properties.subject = "近期（10-17 至 10-25）进展与发现、数据集构建论文阅读、GD 临床答复与证型对应、下一步计划"
    prs.save(str(OUT))
    print(f"已生成：{OUT}（共 {len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
