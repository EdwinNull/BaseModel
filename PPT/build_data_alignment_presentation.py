from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "数据对齐五页.pptx"
OUTPUT = ROOT / "数据对齐五页_结构与示例版.pptx"
BLUE = "006FBD"
NAVY = "174876"
PALE = "EAF4FB"
PALE_2 = "F4F8FC"
TEXT = "243447"
MUTED = "52677A"
ORANGE = "D77932"
WHITE = "FFFFFF"


def shape(slide, name, x, y, width, height, *, text=None, fill="none", line="none", size=14,
          color=TEXT, bold=False, align="left", valign="middle", preset="rect", margin=2):
    props = {
        "name": name, "x": f"{x}pt", "y": f"{y}pt", "width": f"{width}pt",
        "height": f"{height}pt", "fill": fill, "line": line, "preset": preset,
    }
    if text is not None:
        props.update({
            "text": text, "font": "微软雅黑", "size": str(size), "color": color,
            "bold": str(bold).lower(), "align": align, "valign": valign,
            "margin": f"{margin}pt", "autoFit": "normal",
        })
    commands[slide].append({"command": "add", "parent": f"/slide[{slide}]", "type": "shape", "props": props})


def label(slide, name, text, x, y, width, height=24, *, size=14, color=NAVY, bold=True,
          align="left"):
    shape(slide, name, x, y, width, height, text=text, size=size, color=color, bold=bold, align=align)


def card(slide, name, x, y, width, height, title, detail, *, fill=PALE_2, title_size=15, body_size=12.5):
    shape(slide, f"{name}-bg", x, y, width, height, fill=fill, preset="roundRect")
    label(slide, f"{name}-title", title, x + 10, y + 7, width - 20, 24, size=title_size)
    shape(slide, f"{name}-detail", x + 10, y + 32, width - 20, height - 36,
          text=detail, size=body_size, color=TEXT, valign="top", margin=0)


def footer(slide, source, takeaway):
    label(slide, f"source-{slide}", source, 18, 522, 875, 16, size=8.3, color=MUTED, bold=False)
    label(slide, f"takeaway-{slide}", takeaway, 128, 474, 720, 38,
          size=16, color=BLUE, align="center")


def note(slide, text):
    commands[slide].append({"command": "add", "parent": f"/slide[{slide}]", "type": "notes",
                            "props": {"text": text}})


def set_header(slide, title, subtitle):
    commands[slide].extend([
        {"command": "set", "path": f"/slide[{slide}]/shape[@name=矩形 19]",
         "props": {"text": f" {title}", "size": "32", "width": "420pt"}},
        {"command": "set", "path": f"/slide[{slide}]/shape[@name=矩形: 圆角 13-text]",
         "props": {"text": subtitle, "size": "17"}},
    ])


presentation = Presentation(SOURCE)
for slide in presentation.slides:
    for old_shape in list(slide.shapes)[25:]:
        if old_shape.has_chart:
            chart_element = old_shape._element.xpath(".//c:chart")[0]
            slide.part.drop_rel(chart_element.get(qn("r:id")))
        old_shape._element.getparent().remove(old_shape._element)
for container in list(presentation.slide_masters) + list(presentation.slide_layouts):
    for field in container._element.xpath(".//a:fld"):
        if not field.xpath("./a:t/text()"):
            field.getparent().remove(field)
presentation.save(OUTPUT)
commands = {page: [] for page in range(1, 6)}


set_header(1, "对齐后的数据结构", "来源不同，使用相同的表粒度、变量编码与就诊关联键")
label(1, "key", "关联键：source + encounter_id（同一来源内的一次就诊）", 284, 156, 576, 32,
      size=15, color=BLUE)
label(1, "origin-label", "来源数据", 104, 160, 150, 25, size=15)
for index, (name, caption) in enumerate([
    ("GD_CRF", "297 次住院"), ("GAM_IP", "173 次住院"),
    ("MIMIC-IV", "659 次可联结住院"), ("eICU", "15,929 ARI 候选"),
    ("MC-MED", "118,385 次急诊"),
]):
    position = 193 + index * 49
    shape(1, f"origin-{index}", 102, position, 153, 41, fill=PALE,
          preset="roundRect", text=f"{name}  {caption}", size=12.6, color=NAVY,
          bold=True, align="center", margin=3)
shape(1, "origin-arrow", 260, 297, 22, 27, fill=BLUE, preset="rightArrow")
card(1, "encounters", 296, 195, 564, 77, "encounters   1 行 / 就诊",
     "source · encounter_id · patient_id · anchor_time / anchor_rule · age / sex · flu_dx",
     fill="DDEFFA", title_size=17, body_size=13.4)
for index, (name, title, detail) in enumerate([
    ("events", "events   N 行 / 观测", "var · value · unit_cdm\ntime_precision · day_rel"),
    ("treatments", "treatments   N 行 / 治疗", "var · status · start / end\ntime_source"),
    ("complications", "complications   N 行 / 并发症", "var · POA · 判定依据\n时间来源"),
]):
    card(1, name, 296 + index * 191, 303, 181, 86, title, detail,
         title_size=14, body_size=12)
    shape(1, f"down-{index}", 373 + index * 191, 278, 26, 19,
          fill=BLUE, preset="downArrow")
shape(1, "derived-arrow", 566, 392, 23, 14, fill=BLUE, preset="downArrow")
shape(1, "wide-bg", 296, 409, 564, 42, fill=PALE, preset="roundRect")
label(1, "wide-label", "crf_view   1 行 / 住院：adm_* · idx_* · lab_* · OUT00x", 308, 413, 540, 34,
      size=15, color=NAVY)
footer(1, "资料：docs/数据对齐.md；实验执行计划 §4.3。来源规模为原稿状态（2026-09-28）。",
       "一次就诊连接多条事件和治疗，最后派生一行可追溯的建模宽表。")
note(1, "先解释表粒度：encounters 一次就诊一行，events 是长表，治疗和并发症独立保存，crf_view 由这些记录派生。五个来源并不意味着患者可合并，只复用编码、单位和表结构。eICU 与 MC-MED 还未完成正式桥接验收。")


set_header(2, "医院数据落表示例", "用同一例 GD_CRF 记录展示：一次住院如何关联多条体征观测")
for index, (headline, subline) in enumerate([
    ("GD_CRF  297", "住院，日期级"),
    ("GAM_IP  173", "住院，部分时刻可得"),
    ("GAM_OP  1,166", "门诊，不并入住院任务"),
]):
    x_pos = 100 + index * 258
    shape(2, f"hospital-chip-{index}", x_pos, 156, 248, 44, fill=PALE, preset="roundRect")
    label(2, f"hospital-count-{index}", headline, x_pos + 8, 159, 232, 20,
          size=14, color=BLUE, align="center")
    label(2, f"hospital-desc-{index}", subline, x_pos + 8, 178, 232, 19,
          size=11.5, color=MUTED, bold=False, align="center")
shape(2, "patient-bg", 100, 211, 297, 239, fill=PALE_2, preset="roundRect")
label(2, "encounter-head", "encounters  /  1 行", 112, 218, 270, 29, size=16)
for index, (key, value) in enumerate([
    ("source", "GD_CRF"), ("encounter_id", "EXAMPLE_01*"),
    ("age / sex", "75 / M"), ("flu_dx", "1（诊断有流感）"),
    ("anchor_rule", "首个体征"), ("anchor_precision", "日期级"),
]):
    y_pos = 252 + index * 27
    label(2, f"patient-key-{index}", key, 113, y_pos, 126, 23,
          size=12.3, color=MUTED, bold=False)
    label(2, f"patient-value-{index}", value, 242, y_pos, 139, 23,
          size=12.8, color=TEXT, bold=False)
label(2, "related", "另有关联记录：TRT001 奥司他韦 / CLP001 病毒性肺炎", 112, 421,
      273, 26, size=10.6, color=MUTED, bold=False)
shape(2, "join-arrow", 399, 302, 23, 24, fill=BLUE, preset="rightArrow")
shape(2, "event-bg", 427, 211, 438, 239, fill=WHITE, preset="roundRect")
label(2, "event-title", "events  /  6 行观测（节选）", 439, 218, 414, 29, size=16)
columns = [("编码 / 变量", 440, 141), ("value / unit_cdm", 584, 159),
           ("day_rel", 746, 58), ("精度", 806, 49)]
shape(2, "event-table-header", 438, 253, 417, 31, fill="DDEFFA")
for index, (caption, x_pos, width) in enumerate(columns):
    label(2, f"event-col-{index}", caption, x_pos, 255, width, 27,
          size=11.7, color=NAVY, align="center")
for index, (code, value) in enumerate([
    ("VS001  体温", "37.8 ℃"), ("VS002  心率", "68 次/分"),
    ("VS003  呼吸", "17 次/分"), ("VS004  收缩压", "130 mmHg"),
    ("VS005  舒张压", "83 mmHg"), ("VS006  SpO₂", "99 %"),
]):
    y_pos = 286 + index * 26
    if index % 2 == 0:
        shape(2, f"row-bg-{index}", 438, y_pos, 417, 26, fill=PALE_2)
    for suffix, content, x_pos, width in [
        ("var", code, 443, 137), ("value", value, 587, 152),
        ("day", "第 0 天", 748, 57), ("precision", "日期", 808, 44),
    ]:
        label(2, f"row-{index}-{suffix}", content, x_pos, y_pos + 2, width, 22,
              size=11.6, color=TEXT, bold=False,
              align="center" if suffix in ("day", "precision") else "left")
footer(2, "字段值据 docs/数据对齐.md 中的病例示例；*EXAMPLE_01 为展示用关联键，非真实 ID。",
       "同一套编码能落入相同列；GD 的“第 0 天”不等于精确 0 小时。")
note(2, "这页是本地文档中的字段值示例，EXAMPLE_01 是排版用键。GD 缺少真实入院时刻，锚点用首个体征；为跨源比较，这里只展示 day_rel，而不是伪造小时轨迹。门诊用于补充流感证据，不直接并入住院预测队列。")


set_header(3, "公开数据映射边界", "对照源字段、单位与覆盖率；候选映射不等于已通过全量验收")
for index, text in enumerate([
    "MIMIC-IV  659 次可联结住院", "eICU  15,929 ARI 候选 / 严格流感 0",
    "MC-MED  118,385 次急诊 / 字段 QC",
]):
    shape(3, f"public-chip-{index}", 100 + index * 258, 157, 248, 42,
          fill=PALE, preset="roundRect", text=text, size=12.8,
          color=NAVY, bold=True, align="center", margin=4)
shape(3, "mapping-bg", 100, 210, 764, 235, fill=WHITE, preset="roundRect")
shape(3, "mapping-header", 101, 211, 762, 35, fill="DDEFFA")
for index, (caption, x_pos, width) in enumerate([
    ("共同变量", 106, 168), ("医院数据", 274, 162),
    ("MIMIC-IV 候选字段", 437, 289), ("处理", 731, 125),
]):
    label(3, f"mapping-col-{index}", caption, x_pos, 214, width, 27,
          size=13, color=NAVY, align="center")
mapping_rows = [
    ("VS001  体温 ℃", "GD 37.8 ℃ 等", "223762（℃）/ 223761（℉→℃）", "核对时间后用"),
    ("VS004  收缩压", "文本 SBP/DBP", "220179 无创 / 220050 有创", "仅同测法合并"),
    ("LAB001  白细胞", "统一 10⁹/L", "51301/51300；K/uL ×1", "核对单位后用"),
    ("LAB006  CRP", "医院覆盖较高", "50889；窗口内仅 7/659（1.1%）", "暂不入共同核心"),
]
for index, row in enumerate(mapping_rows):
    y_pos = 248 + index * 48
    if index % 2 == 0:
        shape(3, f"mapping-bg-{index}", 102, y_pos, 760, 47, fill=PALE_2)
    for column, (content, x_pos, width) in enumerate(zip(row, [107, 276, 440, 733],
                                                         [165, 158, 285, 122])):
        label(3, f"mapping-{index}-{column}", content, x_pos, y_pos + 7, width, 33,
              size=12.4 if column != 2 else 11.9,
              color=ORANGE if index == 3 and column == 3 else TEXT,
              bold=column == 0, align="center" if column == 3 else "left")
footer(3, "资料：实验执行计划 §4.2–4.3；原稿的 MIMIC 覆盖率。itemid 均需由字典与单位复核。",
       "共同核心选定义、测法和覆盖都站得住的变量；其余保留来源扩展。")
note(3, "这些公开来源目前状态不同：MIMIC 的 659 次住院可联结，但 itemid 是待核对的候选映射；eICU 当前严格流感规则结果是 0，不能据此说数据库没有流感；MC-MED 尚未通过正式 CDM 验收。白细胞较有希望共用，CRP 在 MIMIC 窗口只有 7/659。")


set_header(4, "对齐表进入模型", "长表提取索引日前可用信息，形成一行输入，再做跨医院验证")
for index, text in enumerate([
    "GAM 训练  163 次 / 10 事件", "GD 验证  292 次 / 38 事件",
    "GD 当天转重  26 / 38",
]):
    shape(4, f"trial-chip-{index}", 100 + index * 258, 157, 248, 42,
          fill=PALE, preset="roundRect", text=text, size=13.5,
          color=NAVY, bold=True, align="center", margin=4)
shape(4, "feature-bg", 100, 210, 343, 238, fill=PALE_2, preset="roundRect")
label(4, "feature-head", "从 events 到一行输入（字段示意）", 113, 218, 316, 29, size=15)
for index, line in enumerate([
    "VS001   37.8 ℃    →  体温首值", "VS002   68 次/分 →  心率首值",
    "VS004   130 mmHg →  收缩压首值",
]):
    shape(4, f"feature-row-{index}", 114, 257 + index * 31, 313, 27,
          fill=WHITE if index % 2 == 0 else PALE,
          text=line, size=12.8, color=TEXT, margin=7)
shape(4, "feature-down", 256, 353, 24, 22, fill=BLUE, preset="downArrow")
shape(4, "wide-row", 113, 384, 317, 43, fill="DDEFFA", preset="roundRect",
      text="crf_view：年龄 75 / 男 + 15 项临床测量", size=13,
      color=NAVY, bold=True, align="center", margin=3)
label(4, "feature-caution", "左侧数值仅说明字段形态，不是试验个体结果。", 112, 428,
      320, 20, size=9.8, color=MUTED, bold=False)
shape(4, "result-bg", 455, 210, 408, 238, fill=WHITE, preset="roundRect")
label(4, "result-head", "GD 外部验证  ·  AUROC", 468, 218, 380, 30, size=16)
for index, (model, value, dot_color) in enumerate([
    ("LR", 0.733, NAVY), ("LightGBM", 0.737, "7FB2D8"),
    ("TabPFN", 0.739, BLUE), ("MIRA R0b", 0.609, ORANGE),
]):
    y_pos = 263 + index * 38
    label(4, f"result-name-{index}", model, 468, y_pos - 4, 101, 27,
          size=12.8, color=TEXT, bold=False)
    shape(4, f"result-track-{index}", 580, y_pos + 8, 224, 2, fill="CCDCEB")
    x_pos = 580 + (value - 0.5) / 0.3 * 224
    shape(4, f"result-dot-{index}", round(x_pos - 7, 1), y_pos + 2, 14, 14,
          fill=dot_color, preset="ellipse")
    label(4, f"result-value-{index}", f"{value:.3f}", 808, y_pos - 4, 44, 27,
          size=12.6, color=NAVY, align="right")
for index, tick in enumerate(["0.5", "0.6", "0.7", "0.8"]):
    label(4, f"tick-{index}", tick, 567 + index * 74.7, 416, 42, 20,
          size=10.5, color=MUTED, bold=False, align="center")
footer(4, "资料：原稿 GD 外部验证结果；docs/数据对齐.md 字段示例。O/E = 1.18–2.81。",
       "输入已可用于外部验证；TabPFN 与 LR 接近，当天事件使时点解释受限。")
note(4, "左侧是结构示意：医院 v1.0 取索引日附近各变量第一个值，加年龄、性别；字段数值来自文档示例，不能与右侧验证结果解释为同一患者。右侧点图刻度从 0.5 到 0.8，并明确标出每个 AUROC；TabPFN 0.739 与 LR 0.733 的差异不稳健。GD 38 个事件中 26 个发生在索引当天。")


set_header(5, "对齐边界与下一步", "对齐的不是患者，而是变量含义、可用时间与可比较的结局")
shape(5, "time-bg", 101, 155, 762, 60, fill=PALE, preset="roundRect")
label(5, "time-head", "统一输入按天对齐", 114, 160, 150, 27, size=15)
label(5, "hospital-time", "医院：第 0 天  ●── 第 1 天 ── 第 2 天", 277, 159, 560, 24,
      size=12.8, color=NAVY, bold=False)
label(5, "public-time", "公开： 0 h   ·   6 h   ·   12 h   ·   18 h   ·   24 h", 277, 183, 560, 24,
      size=12.8, color=MUTED, bold=False)
for index, (title, example, action) in enumerate([
    ("01  单位与测法", "D-二聚体 FEU / DDU 未区分；\n血压要分无创与有创。", "先冻结字典、单位和测量方式"),
    ("02  时间可用性", "GD 多为日期级；MIMIC 检验\n预测时需核对 storetime。", "保留 event_time / available_time"),
    ("03  标签与缺失", "eICU 全员在 ICU；当前严格流感 0；\n未测 CRP 不等于阴性。", "先做标签与泄漏 QC，再验证"),
]):
    x_pos = 101 + index * 260
    shape(5, f"boundary-bg-{index}", x_pos, 227, 242, 213, fill=PALE_2,
          preset="roundRect")
    shape(5, f"boundary-number-bg-{index}", x_pos + 12, 239, 34, 34,
          fill=ORANGE if index == 2 else BLUE, preset="ellipse")
    label(5, f"boundary-number-{index}", str(index + 1), x_pos + 13, 241,
          32, 30, size=15, color=WHITE, align="center")
    label(5, f"boundary-title-{index}", title, x_pos + 51, 241, 176, 28, size=15)
    shape(5, f"boundary-detail-{index}", x_pos + 13, 283, 216, 77,
          text=example, size=12.8, color=TEXT, valign="top", margin=1)
    shape(5, f"boundary-action-bg-{index}", x_pos + 13, 378, 216, 43,
          fill="DDEFFA", preset="roundRect")
    label(5, f"boundary-action-{index}", action, x_pos + 18, 384, 206, 30,
          size=11.5, color=NAVY, align="center")
footer(5, "资料：docs/数据对齐.md；实验执行计划 §4.2–4.3；原稿中的质量问题与限制（2026-09-28）。",
       "先冻结字段、时间和结局规则，再开展正式训练与跨来源验证。")
note(5, "三类边界要分别处理：单位和测法核对后才能比较；GD 日期精度不支持补造小时级输入；T-hour 任务必须看数据何时可用。eICU 全员已在 ICU，不能定义新发 ICU 终点。当前规则下严格流感为零只是规则结果。CRP 未测应保留缺失而不是写成阴性。")


with tempfile.TemporaryDirectory() as temporary_dir:
    for slide_number, slide_commands in commands.items():
        command_path = Path(temporary_dir) / f"slide_{slide_number}.json"
        command_path.write_text(json.dumps(slide_commands, ensure_ascii=False), encoding="utf-8")
        subprocess.run(["officecli", "batch", str(OUTPUT), "--input", str(command_path),
                        "--stop-on-error"], check=True, stdout=subprocess.DEVNULL)
        print(f"slide {slide_number}: {len(slide_commands)} edits")
subprocess.run(["officecli", "save", str(OUTPUT)], check=True, stdout=subprocess.DEVNULL)
print(OUTPUT)
