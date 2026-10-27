"""虚构的 GD 式原始导出（结构仿照第一轮扫描的列名；姓名、编号、内容全部虚构）。"""
import datetime as dt, random, re, shutil, zipfile
from pathlib import Path
from openpyxl import Workbook
random.seed(7)
R = Path("fake_gd"); shutil.rmtree(R, ignore_errors=True); (R / "病历/入院记录").mkdir(parents=True); (R / "病历/首次病程").mkdir(parents=True)
NAMES = ["张三", "李四", "王五", "赵六", "欧阳娜娜", "陈小明", "刘伟", "周芳"] * 4
pts = []
for k, nm in enumerate(NAMES):
    mrn = 1200000 + k if k != 5 else 1200001          # 两个住院属于同一病人（再入院）
    enct = 20240000000 + k * 7
    adm = dt.datetime(2024, 1, 2) + dt.timedelta(days=k * 3)
    pts.append(dict(nm=nm, mrn=mrn, enct=enct, uuid=f"U{enct}-{k:04d}-ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrst", adm=adm,
                    los=random.randint(4, 15), jzh=str(enct)[3:]))
def save(wb, p, bad_dim=False):
    wb.save(p)
    if bad_dim:                                   # 把表尺寸写成 A1（仿照病历导出）
        tmp = p.with_suffix(".tmp")
        with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                data = zi.read(it.filename)
                if it.filename.startswith("xl/worksheets/sheet"):
                    data = re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1"/>', data)
                zo.writestr(it, data)
        tmp.replace(p)
# 名单：Sheet1 / Sheet2 / Sheet3（按行对应）
wb = Workbook(); a = wb.active; a.title = "Sheet1"
a.append(["病案号", "姓名", "年龄", "入院日期", "入院科别", "天数", "出院科别", "出院日期", "ICD码", "诊断类型", "疾病名称", "主要诊断疾病名称",
          "主要诊断ICD码", "其他诊断ICD码", "其他诊断疾病名称", "主病中医码", "主病中医疾病名称", "主证中医码", "主证中医疾病名称"])
for p in pts:
    a.append([p["mrn"], p["nm"], random.randint(60, 95), p["adm"].replace(hour=0), "呼吸与危重症医学科", p["los"], "呼吸与危重症医学科",
              (p["adm"] + dt.timedelta(days=p["los"])).replace(hour=0), "J11.100", random.choice([1, 2]), "流行性感冒", "社区获得性肺炎",
              "J18.900", random.choice(["J96.000,I10.x00,E11.900", "J44.100,I25.103", "N17.900,E87.100,I48.000"]),
              "呼吸衰竭,高血压,糖尿病", "A01.02.03", "风温肺热病", "B02.01.02.03.04", "痰热壅肺证"])
b = wb.create_sheet("Sheet2"); b.append(["病案号", "姓名", "出院日期", None])
for p in pts: b.append([p["mrn"], p["nm"], p["adm"], "WY241016"])
c = wb.create_sheet("Sheet3"); c.append(["ENCT_ID"])
for p in pts: c.append([p["enct"]])
save(wb, R / "原始的流感患者名单.xlsx")
# 医嘱
wb = Workbook(); a = wb.active; a.title = "HJJ_WY241016_ORDERS"
a.append(["ENCT_ID", "ENCT_UUID", "住院号", "姓名", "科室", "医嘱类型", "组号", "医嘱名称", "通用名", "药品规格", "总量", "总量单位", "剂量",
          "剂量单位_CD", "频次名称", "天数", "用法", "开始时间", "首日次数", "停止时间", "末日次数", "普通或出院", "医嘱备注"])
ORD = [("磷酸奥司他韦胶囊", "奥司他韦"), ("无创呼吸机辅助通气", None), ("鼻导管吸氧", None), ("重酒石酸去甲肾上腺素注射液", "去甲肾上腺素"),
       ("痰热清注射液", "痰热清"), ("病危", None), ("转重症医学科", None), ("头孢曲松钠", "头孢曲松")]
for p in pts:
    for k in range(12):
        n, g = random.choice(ORD)
        t = p["adm"] + dt.timedelta(hours=6 * k + 2)
        a.append([str(p["enct"]), p["uuid"], str(p["mrn"]), p["nm"], random.choice(["呼吸与危重症医学科", "重症医学科ICU"]),
                  random.choice(["长期医嘱", "临时医嘱"]), f"A{k}", n, g, "75mg", 1, "盒", 75, "mg", "每日两次", 5, "口服", t, 1,
                  t + dt.timedelta(days=3) if k % 2 else None, 1, random.choice(["普通", "出院带药"]), None])
save(wb, R / "HJJ_WY241016_长期临时医嘱.xlsx")
# 生命体征
wb = Workbook(); a = wb.active; a.title = "TEMP_JJ_WY241016"
a.append(["住院号", "ENCT_ID", "姓名", "项目编码", "项目", "值1", "值2", "值3", "值4", "值5", "单位", "备注", "记录时间"])
for p in pts:
    for k in range(40):
        t = p["adm"] + dt.timedelta(hours=4 * k)
        it = random.choice([("TW_T", "体温", f"{random.uniform(36, 39.5):.1f}", None, "℃"), ("MB_P", "脉搏", str(random.randint(60, 120)), None, "次/分"),
                            ("XY_BP", "血压", str(random.randint(100, 160)), str(random.randint(60, 90)), "mmHg"),
                            ("SPO2", "血氧饱和度", str(random.randint(85, 99)), None, "%"), ("HX_R", "呼吸", str(random.randint(14, 30)), None, "次/分")])
        a.append([str(p["mrn"]) if k % 7 else "440101194501011234", str(p["enct"]), p["nm"], it[0], it[1], it[2], it[3], "1" if it[3] else None,
                  None, None, it[4], None, t])
save(wb, R / "生命体征_WY241016.xlsx")
# 检验（只有 JZH）
wb = Workbook(); a = wb.active; a.title = "TEMP_JJ_LIS_WY241016"
a.append(["名称", "结果", "结果单位", "条码号", "参考范围", "参考范围低值", "参考范围高值", "危急参考范围低值", "危急参考范围高值", "是否危急", "采集时间",
          "是否异常", "报告日期", "JZH", "标本类型"])
LAB = [("白细胞计数(WBC)", "8.2", "10^9/L", "全血"), ("C反应蛋白(CRP)", "35.1", "mg/L", "血清"), ("甲型流感病毒抗原", "阳性(+)", None, "咽拭子"),
       ("肌酐(Cr)", "88", "umol/L", "血清"), ("尿白细胞(WBC)", "3", "/uL", "尿液")]
for p in pts:
    for k in range(10):
        n, v, u, sp = random.choice(LAB)
        t = p["adm"] + dt.timedelta(hours=8 * k)
        a.append([n, v, u, f"{random.randint(10**11, 10**12 - 1)}", "3.5-9.5", "3.5", "9.5", "NULL", "NULL", None, t, None, t, p["jzh"], sp])
save(wb, R / "检验结果_WY241016.xlsx")
# 检查
wb = Workbook(); a = wb.active
a.append(["ENCT_ID", "姓名", "住院号", "影像所见", "影像诊断", "报告时间", "类别", "项目"])
for p in pts:
    a.append([str(p["enct"]), p["nm"], str(p["mrn"]), f"{p['nm']}双肺见斑片状磨玻璃影，\n右侧少量胸腔积液。", "双肺炎症；右侧胸腔积液",
              (p["adm"] + dt.timedelta(hours=5)).strftime("%Y-%m-%d %H:%M:%S"), "CT(平扫)", "胸部CT平扫"])
save(wb, R / "检查结果.xlsx")
# 病历（表尺寸写坏）
for sub, fn, head, title in [("入院记录", "入院记录查询_第1页.xlsx", "编号", "入院记录"), ("首次病程", "首次病程记录查询_第1页.xlsx", "住院号", "首次病程记录")]:
    wb = Workbook(); a = wb.active; a.title = "第1页"
    a.append([head, "ENCT_UUID", "姓名", "病历名称", "书写时间", "病历内容"])
    for p in pts:
        txt = (f"姓名：{p['nm']}  性别：男  年龄：78岁\n入院时间：{p['adm']:%Y-%m-%d %H:%M}  记录时间：{p['adm']:%Y-%m-%d %H:%M}\n"
               f"主诉：发热、咳嗽3天。\n现病史：患者3天前受凉后出现发热，最高体温39.2℃，伴咳嗽咳痰。\n既往史：高血压病史10年。\n体格检查：T 38.5℃ P 98次/分 R 22次/分 BP 135/80mmHg")
        a.append([str(p["mrn"]) if head == "住院号" else f"{random.randint(100000, 999999)}", p["uuid"], p["nm"], title, p["adm"] + dt.timedelta(hours=2), txt])
    save(wb, R / "病历" / sub / fn, bad_dim=True)
print("ok", len(pts))
