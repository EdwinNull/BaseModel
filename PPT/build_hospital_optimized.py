# -*- coding: utf-8 -*-
"""
生成《老年流感多中心数据对齐与研究协作汇报》的优化版 PPT。

输入：PPT/hospital.pptx（原稿，只读，不会被修改）
输出：PPT/hospital_优化版.pptx

为什么要重建而不是逐页微调：
1. 字体：原稿每段只写了西文字体（latin=微软雅黑），没写东亚字体（ea），
   主题里的东亚字体又是空的，于是中文全部回退成宋体，形成"中文宋体、英文雅黑"的混排。
   这里先把主题的东亚字体改成微软雅黑，新写的每段文字也显式写入东亚字体。
2. 结构：原稿问题四占 7 页，"起点已重症 / 住院后转重 / 未转重"重复出现 4 次。
   本版按"四个问题"重排为 14 页，每页只讲一件事；问题四每页底部写明"请院方确认"的事项，
   最后一页汇总成分批核实清单。
3. 事实：GD 样本统一为住院次数（292 次 / 38 个事件；原稿 290 是患者数）；
   无创通气的编码改为 treatments 表的 TRT006（原稿的 OUT_NIV 在 CDM 中不存在）。
4. 版式：保留封面；内容页沿用原稿的标题栏、关键信息条、结论框和配色；
   校徽改用封面上的矢量校徽（原稿内容页的位图校徽只有 151×113 像素，且被横向拉伸约 2.2 倍）。

数字均为聚合量，来源：
- docs/两院数据对齐工作与示例数据格式_20261006.docx（第二、五、七、八、九节）
- docs/数据对齐.md、docs/实验进度.md
- docs/数据理想形态与两中心对齐_交互演示_20261005.html（时间跨度、年龄、观测密度）
示例记录（体温、PCT、病程文字）全部为虚构，取自上述 docx 的"示例数据格式"一节。

运行（仓库根目录）：python PPT/build_hospital_optimized.py
依赖：python-pptx、lxml、xlsxwriter（python-pptx 生成原生图表时需要）
"""

import copy
import uuid
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "hospital.pptx"
OUT = HERE / "hospital_优化版.pptx"

# ---------------------------------------------------------------------------
# 全局样式
# ---------------------------------------------------------------------------
FONT = "微软雅黑"
DECK_TITLE = "老年流感多中心数据对齐与研究协作汇报"  # 与封面标题一致（原稿页脚写的是"两中心…院方协作"）

# 配色沿用原稿（取自原稿 slide XML 中的实际色值）
DARK = "1F4E94"     # 深蓝：顶部色条、卡片标题栏
MID = "4075BE"      # 中蓝：标题栏
CYAN = "00A4CB"     # 青色：标题下分隔线、点缀
LIGHT = "EBF3FB"    # 浅蓝：关键信息条、浅色底
PALE = "F5F8FC"     # 更浅的蓝灰：表格隔行
BORDER = "CDDAE8"   # 卡片描边
NAVY = "11376A"     # 深蓝文字：小标题、关键信息
BODY = "273140"     # 正文
MUTED = "646E79"    # 说明文字、页脚
NOTE_BG = "FCEFC6"  # 浅黄：底部结论框（"请院方确认"）
ASK_BG = "FEF6E0"   # 更浅的黄：表格中"需院方确认"一列
WHITE = "FFFFFF"
RED = "C0392B"      # 重症 / 转重
SOFT = "9DBBE3"     # 非重症病程
GREEN = "2E7D5B"    # 文本三态：明确存在
GREY = "8A94A6"     # 文本三态：未提及；参考值

# 版心：左边界 0.55 英寸、宽 12.23 英寸（右边界 12.78），与原稿关键信息条对齐
CL = 0.55
CW = 12.23
SW = 13.333  # 幻灯片宽度（16:9，13.333 × 7.5 英寸）

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------------------
# 文字工具
# ---------------------------------------------------------------------------
def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_typeface(rPr, font=FONT):
    """把一个 <a:rPr>（或 <a:endParaRPr>/<a:defRPr>）的西文、东亚、复杂文种字体都设为 font。

    python-pptx 的 font.name 只写 <a:latin>，中文会去找 <a:ea>，找不到就用主题字体。
    这里把三个都写上；已有的元素只改 typeface 并去掉与新字体不符的 panose/charset。
    """
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))  # 依次追加，顺序正好是 latin → ea → cs
        el.set("typeface", font)
        for attr in ("panose", "charset", "pitchFamily"):
            if attr in el.attrib:
                del el.attrib[attr]


def style_run(run, size, color=BODY, bold=False, italic=False):
    """设置一个文字片段的字号、颜色、粗细和字体（含东亚字体）。"""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    set_typeface(rPr)


def set_bullet(paragraph, color=MID, indent=0.2):
    """给段落加项目符号（悬挂缩进），不在文字里写"•"字符。"""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    # 子元素顺序须为 lnSpc/spcBef/spcAft → buClr → buFont → buChar（此前已由 python-pptx 写好间距）
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def add_break(paragraph, size):
    """在段落末尾加一个软换行 <a:br/>：同一段内换行，不会新起一个项目符号。

    换行符也要写字号和字体，否则这一行的行高会按母版默认字号（18 磅）计算，行距忽大忽小。
    """
    br = paragraph._p.add_br()
    rPr = etree.SubElement(br, qn("a:rPr"), lang="zh-CN", sz=str(int(round(size * 100))))
    set_typeface(rPr)


def fill_text(tf, paras, size=14, color=BODY, bold=False, align="l", line_spacing=1.08, space_after=0):
    """往文本框（text_frame）里写若干段文字。

    paras 的写法：
    - 字符串：一段，统一样式；
    - 列表：每个元素是一段；一段可以是
        * 字符串；
        * [(文字, {选项}), ...] 富文本片段，选项可含 size/color/bold/italic；
        * {"runs": 上面两种之一, "bullet": True, "align": "c", "size": 13, "color": ..., "bold": ...,
           "space_after": 磅, "space_before": 磅}
    文字里的 "\\n" 表示段内换行（软换行）。中文在任意字之间都可能被自动折行，
    较长的句子在这里手动指定断点，避免出现"……精 / 度。"这类孤字。
    """
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = spec if isinstance(spec, dict) else {"runs": spec}
        runs = opts.get("runs", "")
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = ALIGN[opts.get("align", align)]
        p.line_spacing = opts.get("line_spacing", line_spacing)
        p.space_before = Pt(opts.get("space_before", 0))
        p.space_after = Pt(opts.get("space_after", space_after))
        p_size = opts.get("size", size)
        p_color = opts.get("color", color)
        p_bold = opts.get("bold", bold)
        for text, ropt in runs:
            r_size = ropt.get("size", p_size)
            for k, part in enumerate(text.split("\n")):
                if k:
                    add_break(p, r_size)
                if not part:
                    continue
                r = p.add_run()
                r.text = part
                style_run(r, r_size, ropt.get("color", p_color), ropt.get("bold", p_bold), ropt.get("italic", False))
        if opts.get("bullet"):
            set_bullet(p, color=opts.get("bullet_color", MID))


def setup_frame(tf, anchor="t", margins=(0.05, 0.03, 0.05, 0.03), wrap=True):
    """文本框通用设置：自动换行、不自动缩放、内边距、垂直对齐。"""
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = ANCHOR[anchor]


def add_text(slide, x, y, w, h, paras, size=14, color=BODY, bold=False, align="l", anchor="t",
             margins=(0.05, 0.03, 0.05, 0.03), line_spacing=1.08, space_after=0, name=None):
    """新建一个文本框并写入文字，返回该形状。坐标单位均为英寸。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    setup_frame(tb.text_frame, anchor, margins)
    fill_text(tb.text_frame, paras, size, color, bold, align, line_spacing, space_after)
    if name:
        tb.name = name
    return tb


def text_width(s, size_pt):
    """粗略估计单行文字宽度（英寸），只用来给徽章、标签定宽。"""
    em = 0.0
    for ch in s:
        if ord(ch) >= 0x2E80:
            em += 1.0       # 汉字与全角标点约 1 个字宽
        elif ch == " ":
            em += 0.3
        elif ch in "·•":
            em += 0.5
        else:
            em += 0.58      # 西文字母、数字
    return em * size_pt / 72


# ---------------------------------------------------------------------------
# 形状工具
# ---------------------------------------------------------------------------
def strip_style(shape):
    """删除形状自带的 <p:style>。

    python-pptx 新建的形状会引用主题的线条、填充和效果；本主题的效果样式带阴影，
    不删掉的话每个色块都会带一层阴影、一圈主题色描边。删掉后只保留我们显式设置的样式。
    """
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, alpha=16, blur=5, dist=1.5):
    """给卡片加一层很淡的下投影（只用于白色卡片，让它与白底分开）。"""
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)),
                          dist=str(int(dist * 12700)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="1F3B66")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE,
             radius=None, shadow=False, dash=None, name=None):
    """新建一个色块 / 圆角矩形 / 其他预设形状。fill、line 传十六进制颜色，None 表示无。"""
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    else:
        shp.line.fill.background()
    if radius is not None:
        shp.adjustments[0] = radius
    if shadow:
        add_shadow(shp)
    if name:
        shp.name = name
    return shp


def shape_text(shape, paras, size=14, color=BODY, bold=False, align="c", anchor="m",
               margins=(0.04, 0.02, 0.04, 0.02)):
    """直接在形状内部写字（徽章、箭头块等）。"""
    setup_frame(shape.text_frame, anchor, margins)
    fill_text(shape.text_frame, paras, size, color, bold, align)
    return shape


def add_oval(slide, cx, cy, d, fill=None, line=None, line_w=1.25, dash=None):
    """以圆心 (cx, cy)、直径 d 画圆点。"""
    return add_rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, line_w=line_w,
                    shape=MSO_SHAPE.OVAL, dash=dash)


def add_line(slide, x1, y1, x2, y2, color, width=1.0, arrow_end=False, dash=None):
    """画直线；arrow_end=True 时在终点加三角箭头。"""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow_end:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


# ---------------------------------------------------------------------------
# 表格工具
# ---------------------------------------------------------------------------
def set_cell_border(cell, color=None, width=0.75, sides="B"):
    """设置单元格边框：sides 中列出的边画 color 色细线，其余边不画。

    <a:tcPr> 的子元素顺序必须是 lnL → lnR → lnT → lnB → 填充，
    所以这里把边框插在最前面（填充已经由 python-pptx 写在后面）。
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for i, (tag, side) in enumerate((("a:lnL", "L"), ("a:lnR", "R"), ("a:lnT", "T"), ("a:lnB", "B"))):
        ln = etree.Element(qn(tag))
        if color and side in sides:
            ln.set("w", str(int(width * 12700)))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def add_table(slide, x, y, col_w, row_h, rows, size=13, header_size=13.5, header_fill=DARK,
              line_color="DDE6F0", margins=(0.1, 0.05, 0.1, 0.05)):
    """新建原生表格。

    rows：二维列表，第 0 行为表头。每个单元格可以是字符串，也可以是字典：
    {"text": 文字或 fill_text 支持的段落写法, "fill": 底色, "color": 字色, "bold": 粗体,
     "align": "l"/"c", "size": 字号}
    返回 python-pptx 的 table 对象，便于之后合并单元格。
    """
    nrows, ncols = len(rows), len(col_w)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tbl.vert_banding = False
    for i, wv in enumerate(col_w):
        tbl.columns[i].width = Inches(wv)
    for i, hv in enumerate(row_h):
        tbl.rows[i].height = Inches(hv)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            spec = val if isinstance(val, dict) else {"text": val}
            cell = tbl.cell(r, c)
            is_head = r == 0
            fill = spec.get("fill", header_fill if is_head else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            set_cell_border(cell, None if is_head else line_color, 0.75, "B")
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [Inches(m) for m in margins]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text(tf, spec.get("text", ""),
                      size=spec.get("size", header_size if is_head else size),
                      color=spec.get("color", WHITE if is_head else BODY),
                      bold=spec.get("bold", is_head),
                      align=spec.get("align", "l"))
    return tbl


# ---------------------------------------------------------------------------
# 页面框架：标题栏 + 校徽 + 关键信息条 + 底部结论框 + 页脚
# ---------------------------------------------------------------------------
class LogoStamp:
    """把封面右上角的华中科技大学校徽（矢量图形 + 位图）原样复制到其他页。"""

    def __init__(self, cover):
        self.elements = []
        self.image_parts = {}
        for shp in cover.shapes:
            if shp.left is None:
                continue
            # 封面校徽在右上角：left ≥ 10.6 英寸、top < 0.75 英寸
            if shp.left >= Inches(10.6) and shp.top < Inches(0.75):
                el = copy.deepcopy(shp._element)
                self.elements.append(el)
                for blip in el.iter(qn("a:blip")):
                    rid = blip.get(qn("r:embed"))
                    self.image_parts[rid] = cover.part.related_part(rid)

    def stamp(self, slide):
        sp_tree = slide.shapes._spTree
        next_id = max([int(v) for v in sp_tree.xpath(".//p:cNvPr/@id")] + [1]) + 1
        for el in self.elements:
            new = copy.deepcopy(el)
            # 位图需要在新页面上重新建立到图片的关系
            for blip in new.iter(qn("a:blip")):
                old = blip.get(qn("r:embed"))
                blip.set(qn("r:embed"), slide.part.relate_to(self.image_parts[old], RT.IMAGE))
            # 形状 id 在同一页内必须唯一
            for c_nv_pr in new.iter(qn("p:cNvPr")):
                c_nv_pr.set("id", str(next_id))
                next_id += 1
            sp_tree.insert_element_before(new, "p:extLst")


def add_slide_number(slide, x=11.68, y=7.06, w=1.1, h=0.28):
    """页码：使用 PowerPoint 的页码域（slidenum），调整页序后会自动更新。"""
    # 域里的文字只是缓存值，PowerPoint 打开时会重算；这里先写入当前页序号
    number = slide.part.package.presentation_part.presentation.slides.index(slide) + 1
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "页码"
    setup_frame(tb.text_frame, "m", (0, 0, 0, 0))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{%s}" % str(uuid.uuid4()).upper(), type="slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="zh-CN", sz="1050")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=MUTED)
    set_typeface(rPr)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = str(number)


def frame(slide, logo, badge, title, key, ask=None, ask_lead=None, notes=None):
    """内容页的统一框架。

    badge：标题栏左侧的章节标签（如"问题二"）；title：本页标题；key：关键信息条（本页结论）；
    ask_lead + ask：底部浅黄结论框（问题四各页写"请院方确认"的事项）；notes：演讲者备注。
    """
    add_rect(slide, 0, 0, SW, 0.10, fill=DARK, name="顶部色条")
    add_rect(slide, 0, 0.24, 10.30, 0.58, fill=MID, name="标题栏")
    bw = text_width(badge, 14) + 0.34
    b = add_rect(slide, 0.35, 0.33, bw, 0.40, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.25, name="章节标签")
    shape_text(b, badge, 14, DARK, bold=True)
    tx = 0.35 + bw + 0.16
    add_text(slide, tx, 0.24, 10.30 - tx - 0.15, 0.58, title, size=22, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0), name="标题")
    logo.stamp(slide)
    add_rect(slide, 0.50, 0.98, 12.28, 0.04, fill=CYAN, name="分隔线")
    add_rect(slide, CL, 1.12, CW, 0.50, fill=LIGHT, name="关键信息条")
    add_text(slide, CL + 0.2, 1.12, CW - 0.4, 0.50, key, size=16, color=NAVY, bold=True, anchor="m",
             name="关键信息")
    if ask:
        add_rect(slide, CL, 6.42, CW, 0.46, fill=NOTE_BG, name="结论框")
        runs = [(ask_lead, {"bold": True}), (ask, {})] if ask_lead else [(ask, {})]
        add_text(slide, CL + 0.2, 6.42, CW - 0.4, 0.46, [runs], size=15, color=NAVY, anchor="m",
                 name="结论")
    add_text(slide, CL, 7.06, 9.0, 0.28, DECK_TITLE, size=10.5, color=MUTED, anchor="m",
             margins=(0, 0, 0, 0), name="页脚")
    add_slide_number(slide)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def card(slide, x, y, w, h, header, header_h=0.44, header_size=16, header_fill=DARK, shadow=True):
    """白底卡片 + 深蓝标题栏。"""
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=shadow)
    add_rect(slide, x, y, w, header_h, fill=header_fill)
    add_text(slide, x + 0.15, y, w - 0.3, header_h, header, size=header_size, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0))


# ---------------------------------------------------------------------------
# 演讲者备注（每页预计用时、讲解要点、数字来源）
# ---------------------------------------------------------------------------
NOTES = {
    "cover": (
        "预计用时：30 秒。\n"
        "今天汇报两中心数据对齐的进展，以及需要院方协助确认的事项。按四个问题展开：理想数据、两院差异、"
        "已完成的对齐工作、需要共同核实的问题；第 13 页是分批核实清单。\n"
        "（原稿此处的备注是另一份“时序大模型后训练”汇报的内容，已替换。）"
    ),
    "summary": (
        "预计用时：1 分钟。\n"
        "先给出四个问题的一句话回答。核心诉求在第 4 条：转重的定义与发生时间、真实入院时间、"
        "检验单位与字段含义，需要院方帮助确认。右侧是各部分对应的页码。"
    ),
    "ideal": (
        "预计用时：1.5 分钟。\n"
        "用一条病程时间轴说明理想数据：入院、流感检测阳性（索引时点）、住院期间多次测量、"
        "呼吸支持或入 ICU、首次转重、出院或死亡。每个环节都要带时间。\n"
        "两个条件：时间可追溯、病程可重建；同时要纳入未转重病例，否则没有对照。"
    ),
    "sources": (
        "预计用时：1.5 分钟。\n"
        "两院数据的来源完全不同：GAM 是 EHR 多表导出，GD 是人工整理的 CRF 宽表。两院没有共同患者"
        "（同日三项血常规完全相同的记录为 0 组），只能在变量与定义层面对齐。\n"
        "GAM 是全院提取，109/173（63%）次住院在入院 48 小时后才首次流感阳性，入院值不能直接与 GD 比较。\n"
        "来源：docs/数据对齐.md；docs/两院数据对齐工作与示例数据格式_20261006.docx；"
        "时间跨度与年龄见 docs/数据理想形态与两中心对齐_交互演示_20261005.html。"
    ),
    "compare": (
        "预计用时：2 分钟。\n"
        "逐行对照理想记录与两院现状。重点讲三行：入院时间（GD 没有）、体征（GD 每项 1–2 次；GAM 用的是"
        "病历书写时间）、呼吸支持（GAM 需从文字推断）。\n"
        "强调差异来自采集方式，不是数据质量高低。来源：docx 第九节“同一件事，理想的记录与现在的记录”。"
    ),
    "work": (
        "预计用时：1.5 分钟。\n"
        "四类工作：结构、语义与时间、标签、质量检查。本页数字是 CDM 对齐口径（全部记录）。\n"
        "SEV_B 是重症候选规则：ICU、有创或无创通气、ECMO、死亡、I 型或 II 型呼衰任一即判为重症；"
        "在 GD 293 次住院上与人工判定的一致性 κ = 0.917（漏判 0、多判 10）。\n"
        "流水线共 9 步，原始表格只在服务器上只读使用。来源：docx 第二、四、七节。"
    ),
    "records": (
        "预计用时：1.5 分钟。\n"
        "用虚构示例说明统一格式：GD 一行里的体温和 PCT 拆成带时间精度的观测记录；GAM 病程文字里的 SpO₂ "
        "进 events 表，无创通气进 treatments 表（TRT006）；“必要时气管插管”是假设语境，不记为有创通气。\n"
        "每条记录都保留原值、原时间、时间精度和来源，修正只打标记。\n"
        "（原稿写的 OUT_NIV 在 CDM 中不存在，已改为 TRT006。）"
    ),
    "events": (
        "预计用时：1.5 分钟。\n"
        "口径为 v1.0：转重 = 有创通气、ICU、ECMO 或死亡（E1），住院期间任意时间发生。GAM 163 次住院（161 人）、"
        "10 个事件，全部用于训练；GD 292 次住院（290 人）、38 个事件，全部用于外部验证。\n"
        "外部验证的样本量研究建议至少 100 个事件（Collins 等，Stat Med 2016）。约 2/3 的事件发生在流感索引当天。"
        "GAM 经单人逐例裁定后为 14 个事件，说明规则识别有遗漏，需要联合复核。\n"
        "注意：原稿 GD 写“290 例 / 13.1%”，290 是患者数；按住院计为 292 次、13.0%。"
        "GAM 的 10 / 14 两个版本需在正式材料中统一。来源：docs/实验进度.md；docs/数据对齐.md。"
    ),
    "course": (
        "预计用时：1.5 分钟。\n"
        "四条泳道对应四种病程。重点是第三种：转重后好转出院，只看出院状态会误判为“未转重”。"
        "第一种“起点已重症”需要单列，如何纳入需共同确认。\n"
        "GAM 的相关信息分散在病历叙述中，候选事件需要结合临床记录核验。"
    ),
    "time": (
        "预计用时：1.5 分钟。\n"
        "示意图比较三种数据形态：理想的逐次测量；类似 GAM 的每天 1–3 次（病历书写时间）；类似 GD 的"
        "每次住院 1–2 次（部分只有日期）。\n"
        "GD 用首个有效体征时间作时间零点（292/297 次），这不等于真实入院时间；只有日期的数据保留精度标志，"
        "不推断为分钟。来源：docx 第八节（时间精度真实计数）；体征时间点中位数见交互演示页 1.2 节。"
    ),
    "text": (
        "预计用时：1 分钟。\n"
        "文本三态：明确存在、明确不存在、未提及。未提及不等于没有发生。还要识别假设、计划、拒绝、"
        "围术期和伪否定等语境。希望院方协助确认术语词表与否定规则。"
    ),
    "fields": (
        "预计用时：1 分钟。\n"
        "左列是已完成的技术处理，右列是需院方确认的口径。重点：D-二聚体 FEU / DDU 相差约 2 倍；"
        "CRF 表头 TRT009 重复；“判定依据”1–4 的含义；急性肾损伤判定标准（CRF 与 KDIGO 一致性 κ = 0.347）。"
    ),
    "actions": (
        "预计用时：2 分钟。\n"
        "按对研究的影响分三批：第一批决定主要结局（转重定义、时间锚点）；第二批决定两院可比性"
        "（单位与方法、字段字典、文本规则）；持续推进补充病例与导出规范。\n"
        "核实完成后形成最终字段字典、冻结版 CDM、统一结局标签和可复现的实验数据。"
    ),
    "closing": "致谢与讨论。",
}


# ---------------------------------------------------------------------------
# 各页内容
# ---------------------------------------------------------------------------
def build_summary(slide, logo):
    """第 2 页：四个问题与一句话回答（目录 + 结论前置）。"""
    frame(slide, logo, "要点", "四个问题与一句话回答",
          "核心诉求：与院方共同确认“谁发生转重、何时转重”，并补齐时间锚点与字段含义。",
          ask_lead="汇报目的：", ask="说明研究需求与数据现状，形成可执行的分批核实清单（第 13 页）。",
          notes=NOTES["summary"])
    items = [
        ("01", "我们理想的数据是怎样的？",
         "以病程为主线：观测带真实测量时间、住院期间多次测量、结局带发生时间。", "第 3 页"),
        ("02", "GAM 和 GD 的数据差异在哪里？",
         "采集方式不同（EHR 导出 vs 人工整理 CRF），差距集中在时间精度与观测密度。", "第 4–5 页"),
        ("03", "我们做了哪些对齐工作？",
         "建成 8 张表的统一数据模型，逐一处理 22 个问题；重症规则与人工判定 κ = 0.917。", "第 6–7 页"),
        ("04", "哪些问题需要共同解决？",
         "转重定义与首次时间、真实入院时间、检验单位与字段含义、文本“未提及”的编码。", "第 8–13 页"),
    ]
    h, gap = 0.92, 0.14
    for i, (num, q, a, page) in enumerate(items):
        y = 1.86 + i * (h + gap)
        add_rect(slide, 0.72, y, 11.9, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.12, shadow=True)
        nb = add_rect(slide, 0.95, y + 0.16, 0.72, 0.60, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                      radius=0.15)
        shape_text(nb, num, 20, WHITE, bold=True)
        add_text(slide, 1.92, y + 0.07, 9.0, 0.42, q, size=18, color=NAVY, bold=True, anchor="m")
        add_text(slide, 1.92, y + 0.49, 9.1, 0.36, a, size=14, color=BODY, anchor="m")
        add_line(slide, 11.18, y + 0.2, 11.18, y + h - 0.2, BORDER, 1)
        add_text(slide, 11.22, y, 1.32, h, page, size=13, color=MID, bold=True, align="c", anchor="m")


def build_ideal(slide, logo):
    """第 3 页（问题一）：一条病程时间轴 + 五类要素 + 两个条件。"""
    frame(slide, logo, "问题一", "理想的数据：沿病程记录“发生了什么、何时发生”",
          "目标：以患者病程为主线，支持“转重是否发生、何时发生”的动态研究。",
          ask_lead="研究重点：", ask="不仅记录最终出院状态，更关注住院期间病情变化的时间顺序。",
          notes=NOTES["ideal"])
    ly = 2.62  # 时间轴所在高度
    add_line(slide, 0.85, ly, 12.45, ly, DARK, 2.5, arrow_end=True)
    markers = [  # (x, 圆点填充, 圆点描边, 上方文字, 下方文字, 上方文字颜色)
        (1.30, WHITE, DARK, "入院 / 门急诊", "观察起点", NAVY),
        (3.05, CYAN, WHITE, "流感检测阳性", "索引时点", NAVY),
        (8.00, MID, WHITE, "呼吸支持 / 入 ICU", "开始与结束时间", NAVY),
        (9.75, RED, WHITE, "首次转重", "发生时间", RED),
        (11.75, DARK, WHITE, "出院 / 死亡", "结局与去向", NAVY),
    ]
    # 住院期间的多次测量：沿时间轴等距的小圆点（避开大圆点）
    x = 3.55
    while x < 11.35:
        if all(abs(x - m[0]) > 0.22 for m in markers):
            add_oval(slide, x, ly, 0.10, fill=CYAN)
        x += 0.33
    for mx, fill, line, top, bottom, tc in markers:
        add_oval(slide, mx, ly, 0.24, fill=fill, line=line, line_w=2)
        add_text(slide, mx - 0.85, 2.0, 1.7, 0.44, top, size=13.5, color=tc, bold=True, align="c", anchor="b")
        add_text(slide, mx - 0.85, 2.8, 1.7, 0.36, bottom, size=12, color=MUTED, align="c")
    add_text(slide, 4.55, 2.0, 2.6, 0.44, "体征、检验多次测量", size=13.5, color=MID, bold=True, align="c",
             anchor="b")
    add_text(slide, 4.55, 2.8, 2.6, 0.36, "每次带真实测量时间", size=12, color=MUTED, align="c")

    cards = [
        ("患者与就诊", ["脱敏患者 ID、门急诊", "入院、出院、ICU 时间"]),
        ("连续观测", ["多次生命体征与检验", "每次真实测量时间"]),
        ("临床过程", ["病历、病原、影像", "用药与呼吸支持"]),
        ("临床结局", ["ICU、通气、ECMO", "死亡、出院方式"]),
        ("预测标签", ["是否发生转重", "首次转重时间"]),
    ]
    cw, gap = 2.3, (CW - 5 * 2.3) / 4
    for i, (head, lines) in enumerate(cards):
        x = CL + i * (cw + gap)
        card(slide, x, 3.38, cw, 1.45, head, header_h=0.42, header_size=15)
        add_text(slide, x + 0.12, 3.86, cw - 0.24, 0.9, lines, size=13.5, color=BODY, space_after=3)

    conds = [("时间可追溯", "保留原始测量、治疗与事件时间\n并注明时间精度"),
             ("病程可重建", "保留多次观测、明确结局与观察起点\n同时纳入未转重病例作为对照")]
    tw = (CW - 0.2) / 2
    for i, (k, v) in enumerate(conds):
        x = CL + i * (tw + 0.2)
        add_rect(slide, x, 5.1, 1.55, 0.8, fill=DARK)
        add_text(slide, x, 5.1, 1.55, 0.8, k, size=15, color=WHITE, bold=True, align="c", anchor="m")
        add_rect(slide, x + 1.55, 5.1, tw - 1.55, 0.8, fill=LIGHT)
        add_text(slide, x + 1.7, 5.1, tw - 1.85, 0.8, v, size=14, color=BODY, anchor="m")


def build_sources(slide, logo):
    """第 4 页（问题二）：两中心的数据来源与研究对象。"""
    frame(slide, logo, "问题二", "两中心的数据来源与研究对象",
          "两院采集方式不同、患者不重叠，只能在“变量与定义”层面对齐。",
          ask_lead="互补价值：", ask="GAM 保留较丰富的诊疗过程，GD 提供专病结构化字段；需统一变量含义与纳排标准。",
          notes=NOTES["sources"])
    cw = (CW - 0.23) / 2
    hospitals = [
        dict(head="GAM · 广安门医院",
             stats=[("173", "次住院（168 人）"), ("1,166", "次门诊（1,136 人）")],
             span="时间跨度：2024-11 至 2026-03",
             rows=[("数据形态", "多科室临床 EHR 导出：住院 6 张原始表 + 门诊记录"),
                   ("信息组织", "病历、检验、影像、医嘱分散于不同记录"),
                   ("人群特点", "全院提取，含 <60 岁 44 次\n63% 的住院在入院 48 小时后才首次流感阳性"),
                   ("研究用途", "先筛选老年流感病例，再提取临床变量")]),
        dict(head="GD · 广东省中医院",
             stats=[("297", "次住院（295 人）"), ("≥60 岁", "全部为老年流感住院")],
             span="时间跨度：2021-11 至 2026-06",
             rows=[("数据形态", "已整理的老年流感研究队列 CRF（人工整理宽表）"),
                   ("信息组织", "症状、基础病、治疗及结局字段较集中"),
                   ("人群特点", "年龄 60–97 岁，中位数 68 岁"),
                   ("研究用途", "适合按统一字段开展队列分析与验证")]),
    ]
    row_h = [0.5, 0.5, 0.62, 0.5]
    for i, hsp in enumerate(hospitals):
        x = CL + i * (cw + 0.23)
        card(slide, x, 1.8, cw, 4.05, hsp["head"], header_h=0.46, header_size=17)
        for j, (num, label) in enumerate(hsp["stats"]):
            sx = x + 0.25 + j * 2.85
            add_text(slide, sx, 2.36, 2.7, 0.52, num, size=28, color=DARK, bold=True, anchor="b",
                     margins=(0, 0, 0, 0))
            add_text(slide, sx, 2.88, 2.7, 0.3, label, size=12.5, color=MUTED, margins=(0, 0, 0, 0))
        add_text(slide, x + 0.25, 3.2, cw - 0.5, 0.3, hsp["span"], size=12.5, color=MUTED, margins=(0, 0, 0, 0))
        add_line(slide, x + 0.25, 3.58, x + cw - 0.25, 3.58, BORDER, 1)
        y = 3.66
        for (label, value), rh in zip(hsp["rows"], row_h):
            add_text(slide, x + 0.25, y, 1.05, rh, label, size=13, color=NAVY, bold=True, anchor="m",
                     margins=(0, 0, 0, 0))
            add_text(slide, x + 1.35, y, cw - 1.6, rh, value, size=13.5, color=BODY, anchor="m",
                     margins=(0, 0, 0, 0))
            y += rh
    add_text(slide, CL, 5.92, CW, 0.32,
             "注：数字为 CDM 对齐口径（全部记录，截至 2026-10-05）；两院同日三项血常规完全相同的记录为 0 组，未发现重复患者。",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))


def build_compare(slide, logo):
    """第 5 页（问题二）：同一件事，理想的记录与两院现状（原生表格）。"""
    frame(slide, logo, "问题二", "同一件事：理想的记录与两院现状",
          "差异集中在“时间”与“密度”：GD 缺入院时间、检验只有日期；GAM 体征与呼吸支持需从文字推断。",
          ask_lead="说明：", ask="差异主要来自采集方式，而非数据质量高低；对齐目标是“可追溯、可解释、知道哪里不可比”。",
          notes=NOTES["compare"])
    first = lambda t: {"text": t, "fill": LIGHT, "color": NAVY, "bold": True}  # 第一列：浅蓝底、深蓝粗体
    ideal = lambda t: {"text": t, "fill": PALE, "color": NAVY}                  # 理想列：淡底，作为参照
    rows = [
        ["信息", "理想的记录", "GD 现状", "GAM 现状"],
        [first("入院时间"), ideal("入院登记时间"), "没有，以首个有效体征时间代替", "入院记录中的“入院时间”"],
        [first("流感检测"), ideal("采样时间、方法、结果"), "没有检测时间", "检验表有时间"],
        [first("体温、心率、SpO₂"), ideal("逐次测量 + 测量时间 + 吸氧方式与流量"), "CRF 摘录，每项每次住院 1–2 次",
         "从病历文字抽取，时间为书写时间"],
        [first("检验"), ideal("采样 / 报告时间 + 单位 + 方法"), "只有日期；只有表头单位", "到分钟，有单位、参考范围、异常标记"],
        [first("呼吸支持"), ideal("开始 / 结束时间、方式、参数"), "CRF 勾选 + 日期",
         "从文字推断\n需排除“必要时”“拒绝”“术后”"],
        [first("用药"), ideal("医嘱开始 / 停止时间"), "CRF 开始 / 结束日期", "医嘱无时间，以病程首次提及近似"],
        [first("症状"), ideal("结构化“有 / 无 / 未问”"), "CRF 0/1 + 记录日期", "文字三态：肯定 / 否定 / 未提及"],
    ]
    # 列宽按最长一行的实测宽度分配；表格底边（6.14）与底部结论框（6.42）之间留出空隙
    add_table(slide, CL, 1.8, [1.9, 3.5, 3.0, 3.83], [0.42, 0.55, 0.55, 0.55, 0.55, 0.62, 0.55, 0.55], rows,
              size=13, header_size=14)


def build_work(slide, logo):
    """第 6 页（问题三）：已完成的四类对齐工作 + 9 步流水线。"""
    frame(slide, logo, "问题三", "已完成的四类对齐工作",
          "对齐时始终保留原始数值和来源；修正另加标记，具体临床口径仍以共同确认结果为准。",
          ask_lead="口径说明：", ask="本页数量为 CDM 对齐数据口径（全部住院 / 门诊记录），非后续转重分析的入模样本数。",
          notes=NOTES["work"])
    # 每条要点控制在一行内（卡片内可用宽度约 2.46 英寸，13 磅约 13 个汉字），避免折出孤字
    works = [
        ("① 结构对齐", "8 张表", ["统一数据模型 CDM v3.1", "GAM 住院 173、门诊 1,166", "GD 住院 297"]),
        ("② 语义与时间对齐", "22 个问题", ["统一编码、单位与时间零点", "确定流感索引时点", "症状三态、保留时间精度"]),
        ("③ 标签对齐", "κ = 0.917", ["重症规则 SEV_B 与人工判定", "GD 293 次住院中比较", "漏判 0、多判 10"]),
        ("④ 质量检查", "0 命中", ["隐私审计 0 命中", "两院重复患者 0 组", "两次构建 28 个产物一致"]),
    ]
    cw, gap = 2.92, (CW - 4 * 2.92) / 3
    for i, (head, big, lines) in enumerate(works):
        x = CL + i * (cw + gap)
        card(slide, x, 1.8, cw, 2.42, head, header_h=0.44, header_size=15.5)
        add_text(slide, x + 0.13, 2.3, cw - 0.26, 0.55, big, size=26, color=DARK, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.13, 2.92, cw - 0.26, 1.22,
                 [{"runs": t, "bullet": True, "space_after": 4} for t in lines],
                 size=13, color=BODY, margins=(0, 0, 0, 0))
    # 罗马数字"Ⅰ/Ⅱ"在雅黑里会带出多余间距，这里改用西文字母 I / II
    add_text(slide, CL, 4.28, CW, 0.3,
             "SEV_B（重症候选规则）：ICU、有创或无创通气、ECMO、死亡、I 型或 II 型呼吸衰竭，任一发生即判为重症。",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))

    add_text(slide, CL, 4.72, 4.0, 0.36, "对齐流水线（9 步）", size=15, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    steps = ["只读读入", "识别就诊", "化名", "变量映射", "时间处理", "单位与范围", "衍生候选", "写盘审计", "质量汇总"]
    sw = (CW + 8 * 0.08) / 9  # 相邻箭头块略微重叠 0.08 英寸，形成连贯的流程条
    for i, s in enumerate(steps):
        x = CL + i * (sw - 0.08)
        shp_type = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        fill = DARK if i % 2 == 0 else MID
        blk = add_rect(slide, x, 5.15, sw, 0.62, fill=fill, shape=shp_type)
        blk.adjustments[0] = 0.28
        # 箭头形状的文字区本身已让出两端的尖角（约 0.17 英寸），内边距只需很小，五字标签才放得下
        shape_text(blk, s, 12.5, WHITE, bold=True, margins=(0.08, 0, 0.04, 0))
    add_text(slide, CL, 5.85, CW, 0.32,
             "原始表格只在服务器上只读使用；每一步的修正都另存新值并打标记，可回到原表核对。",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))


def build_records(slide, logo):
    """第 7 页（问题三）：原始记录 → 统一 CDM 记录（虚构示例）。"""
    frame(slide, logo, "问题三", "对齐后的记录如何保留原始信息",
          "以下为虚构记录示例，仅用于说明实际采用的数据结构与可追溯原则。",
          ask_lead="原则：", ask="统一数据不覆盖原始数据；文本提及的时间也不直接等同于事件真实发生时间。",
          notes=NOTES["records"])
    lw = 4.85
    # 左上：GD 原始 CRF
    card(slide, CL, 1.8, lw, 1.72, "GD 原始 CRF（一次住院占多行）", header_h=0.42, header_size=14.5)
    add_table(slide, CL + 0.15, 2.32, [1.62, 0.66, 1.3, 0.97], [0.34, 0.38],
              [["体征记录时间", "体温", "检验记录时间", "PCT"],
               ["2024-01-03 10:30", "38.6", "2024-01-08", "<0.05"]],
              size=12, header_size=11.5, header_fill="5A7FB8", margins=(0.06, 0.02, 0.06, 0.02))
    add_text(slide, CL + 0.15, 3.07, lw - 0.3, 0.4, "体征、检验、症状各有自己的“记录时间”，分别展开。",
             size=11.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))
    # 左下：GAM 病历文本
    card(slide, CL, 3.68, lw, 1.72, "GAM 病历文本（日常病程）", header_h=0.42, header_size=14.5)
    add_text(slide, CL + 0.15, 4.16, lw - 0.3, 1.2, [
        [("2024-01-05 08:30　", {"color": MUTED, "size": 12})],
        [("患者喘憋加重，", {}), ("SpO₂ 88%", {"color": MID, "bold": True}),
         ("（鼻导管吸氧 3 L/min），予", {}), ("无创呼吸机辅助通气", {"color": GREEN, "bold": True}),
         ("。……病情加重", {}), ("必要时气管插管", {"color": GREY, "bold": True}), ("。", {})],
    ], size=13, color=BODY, margins=(0, 0, 0, 0), line_spacing=1.1)
    add_text(slide, CL + 0.15, 4.96, lw - 0.3, 0.36, "同一段文字拆出 2 条记录，并排除 1 处假设语境。",
             size=11.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))
    # 中间箭头
    arr = add_rect(slide, CL + lw + 0.12, 3.25, 0.5, 0.7, fill=CYAN, shape=MSO_SHAPE.RIGHT_ARROW)
    arr.name = "转换箭头"
    # 右侧：统一 CDM 记录
    rx = CL + lw + 0.74
    rw = CL + CW - rx
    add_text(slide, rx, 1.8, rw, 0.36, "统一 CDM 记录（events / treatments 表）", size=15, color=NAVY, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    code = lambda t: {"text": t, "bold": True, "color": NAVY}
    # 时间单元格分两行：第一行是采用的时间，第二行（灰色小字）是时间精度
    when = lambda t, prec: {"text": [t, {"runs": prec, "size": 11, "color": MUTED}]}
    rows = [
        ["表", "编码与变量", "数值", "时间与精度", "来源 / 标记"],
        ["events", code("VS001 体温"), "38.6", when("2024-01-03 10:30", "到分钟"), "GD CRF"],
        ["events", code("LAB007 PCT"), "0.05", when("2024-01-08", "只有日期"), "删失标记“<”"],
        ["events", code("VS006 SpO₂"), "88", when("2024-01-05 08:30", "病历书写时间"), "GAM 病历"],
        ["treatments", code("TRT006 无创通气"), "1", when("2024-01-05 08:30", "首次肯定提及"), "GAM 病历"],
        ["treatments", code("TRT007 有创通气"), "不记", "—",
         {"text": ["假设语境", {"runs": "“必要时”插管", "size": 11, "color": MUTED}]}],
    ]
    add_table(slide, rx, 2.2, [1.15, 1.7, 0.62, 1.7, rw - 5.17], [0.38] + [0.56] * 5, rows,
              size=12, header_size=12.5)
    # 底部：每条记录同步保留的字段
    add_rect(slide, CL, 5.58, CW, 0.62, fill=LIGHT)
    add_text(slide, CL + 0.2, 5.58, 1.85, 0.62, "每条记录同步保留", size=14, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    chips = ["原值 raw_value", "原时间 time_raw", "时间精度 time_precision", "来源 origin", "修正标记（如 year+1）"]
    x = CL + 2.1  # 五个标签总宽约 9.8 英寸，从这里起排，右端不超出浅蓝底条
    for c in chips:
        w = text_width(c, 12.5) + 0.32
        chip = add_rect(slide, x, 5.71, w, 0.36, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        radius=0.5)
        shape_text(chip, c, 12.5, NAVY)
        x += w + 0.12


def build_events(slide, logo):
    """第 8 页（问题四 A）：研究样本与可识别的转重事件数（原生条形图）。"""
    frame(slide, logo, "问题四", "A. 研究样本与可识别的转重事件数",
          "当前口径下可识别的转重事件有限（GAM 10 例、GD 38 例），直接制约模型训练与评估的稳定性。",
          ask_lead="请院方协助：", ask="复核候选转重事件；评估可补充的“有明确事件时间的转重病例”与“明确未转重病例”。",
          notes=NOTES["events"])
    add_rect(slide, CL, 1.8, 6.35, 4.42, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.88, 6.0, 0.4, "转重事件数（按住院计）", size=15, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    cd = CategoryChartData()
    # 横向条形图自下而上排列类别，所以把参考值放在第一个，GAM 放在最后（显示在最上面）
    cd.categories = ["参考：外部验证建议", "GD（外部验证）", "GAM（训练）"]
    cd.add_series("转重事件数", (100, 38, 10))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(CL + 0.1), Inches(2.3), Inches(6.15),
                                Inches(3.2), cd)
    gf.name = "事件数条形图"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(13)
    ch.font.name = FONT
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.vary_by_categories = False
    ser = plot.series[0]
    for i, col in enumerate(["D9DFE8", DARK, MID]):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
        if i == 0:  # 参考值用浅灰 + 虚线框，与两家医院的真实数据区分
            pt.format.line.color.rgb = rgb(GREY)
            pt.format.line.width = Pt(1)
            pt.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    for i, t in enumerate(["≥100 例", "38 例（13.0%）", "10 例（6.1%）"]):
        dl = ser.points[i].data_label
        dl.has_text_frame = True
        dl.text_frame.text = t
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        # 数据标签默认会按图表宽度自动折行（"38 例 / (13.0%)"），这里关掉折行
        dl.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")
        for p in dl.text_frame.paragraphs:
            for r in p.runs:
                style_run(r, 13, GREY if i == 0 else NAVY, bold=True)
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = 135
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    # 隐藏数值轴：python-pptx 写出的 <c:delete/> 不带 val，PowerPoint 会当作"不删除"，
    # 所以显式写 val="1"；再把刻度文字也关掉，双保险
    va._element.get_or_add_delete_().set("val", "1")
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    ca = ch.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(13)
    ca.tick_labels.font.color.rgb = rgb(BODY)
    # 图表文字同样补上东亚字体
    for latin in list(ch._chartSpace.iter(qn("a:latin"))):
        parent = latin.getparent()
        if parent.find(qn("a:ea")) is None:
            ea = etree.Element(qn("a:ea"))
            ea.set("typeface", FONT)
            latin.addnext(ea)
    add_text(slide, CL + 0.2, 5.55, 6.0, 0.6,
             ["括号内为事件占住院次数的比例。参考值：外部验证样本量研究建议至少 100 个事件（Collins 等，Stat Med 2016）。"],
             size=10.5, color=MUTED, margins=(0, 0, 0, 0))

    # 右侧：口径与注意事项
    rx, rw = CL + 6.6, CW - 6.6
    card(slide, rx, 1.8, rw, 2.15, "当前口径（v1.0 研究口径）", header_h=0.42, header_size=14.5)
    add_text(slide, rx + 0.18, 2.32, rw - 0.36, 1.58, [
        {"runs": [("转重：", {"bold": True, "color": NAVY}), ("有创通气、ICU、ECMO 或死亡", {})],
         "bullet": True, "space_after": 5},
        {"runs": [("时间窗：", {"bold": True, "color": NAVY}), ("住院期间任意时间", {})],
         "bullet": True, "space_after": 5},
        {"runs": [("GAM：", {"bold": True, "color": NAVY}), ("163 次住院（161 人），全部用于训练", {})],
         "bullet": True, "space_after": 5},
        {"runs": [("GD：", {"bold": True, "color": NAVY}), ("292 次住院（290 人），全部用于外部验证", {})],
         "bullet": True},
    ], size=13.5, color=BODY, margins=(0, 0, 0, 0))
    card(slide, rx, 4.08, rw, 2.14, "需要注意", header_h=0.42, header_size=14.5, header_fill="B7791F")
    add_text(slide, rx + 0.18, 4.6, rw - 0.36, 1.56, [
        {"runs": "约 2/3 的事件发生在流感索引当天，\n可用于“提前预警”的事件更少", "bullet": True,
         "bullet_color": "B7791F", "space_after": 8},
        {"runs": "GAM 经单人逐例裁定后为 14 例：\n规则识别存在遗漏，需联合复核", "bullet": True,
         "bullet_color": "B7791F"},
    ], size=13.5, color=BODY, margins=(0, 0, 0, 0))


def build_course(slide, logo):
    """第 9 页（问题四 B）：四种病程泳道图。"""
    frame(slide, logo, "问题四", "B. 转重标签应反映病程，而不仅是出院结局",
          "建议明确区分“起点已处于重症”“观察期间转重”“转重后好转”“始终未转重”四种情况。",
          ask_lead="请院方确认：", ask="观察起点的重症状态、首次转重时间，以及转重后恢复出院的病例。",
          notes=NOTES["course"])
    tx0, tx1 = 3.05, 8.95       # 病程轨道的起止横坐标
    track_w = tx1 - tx0
    # 轨道上方的坐标说明与图例
    # 两端说明分别以起点、终点圆点为中心
    add_text(slide, tx0 - 0.6, 1.78, 1.2, 0.3, "观察起点 t0", size=11.5, color=MUTED, align="c",
             margins=(0, 0, 0, 0))
    add_text(slide, tx1 - 0.6, 1.78, 1.2, 0.3, "出院 / 死亡", size=11.5, color=MUTED, align="c",
             margins=(0, 0, 0, 0))
    add_text(slide, tx0 + 1.6, 1.78, 2.7, 0.3, "住院时间 →", size=11.5, color=MUTED, align="c", margins=(0, 0, 0, 0))
    # 图例放在右侧说明列的上方，右端与版心对齐，避免与"出院 / 死亡"重叠
    lx = 9.62
    add_rect(slide, lx, 1.85, 0.32, 0.16, fill=SOFT)
    add_text(slide, lx + 0.38, 1.76, 0.7, 0.32, "非重症", size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))
    add_rect(slide, lx + 1.18, 1.85, 0.32, 0.16, fill=RED)
    add_text(slide, lx + 1.56, 1.76, 1.62, 0.32, "重症（ICU、通气等）", size=11.5, color=BODY, anchor="m",
             margins=(0, 0, 0, 0))

    lanes = [
        # (名称, 说明, [(起点比例, 终点比例, 颜色), ...], 首次转重位置比例或 None, 右侧说明, 说明颜色)
        ("起点已处于重症", "观察起点时已在 ICU 或通气", [(0, 0.45, RED), (0.45, 1, SOFT)], None,
         "需单列；纳入方式待共同确认", NAVY),
        ("观察期间转重", "起点非重症，住院中首次转重", [(0, 0.55, SOFT), (0.55, 1, RED)], 0.55,
         "预测目标：需记录首次转重时间", NAVY),
        ("转重后好转出院", "转重经治疗恢复后出院", [(0, 0.35, SOFT), (0.35, 0.68, RED), (0.68, 1, SOFT)], 0.35,
         "只看出院状态会误判为“未转重”", RED),
        ("始终未转重", "全程未出现重症事件", [(0, 1, SOFT)], None,
         "需与“未记录 / 未提及”区分", NAVY),
    ]
    lane_h, lane_gap, y0 = 0.86, 0.1, 2.16
    for i, (name, sub, segs, sev_at, note, note_color) in enumerate(lanes):
        y = y0 + i * (lane_h + lane_gap)
        cy = y + lane_h / 2
        add_rect(slide, CL, y, CW, lane_h, fill=PALE if i % 2 == 0 else WHITE, line=None)
        add_text(slide, CL + 0.15, y + 0.06, 2.3, 0.4, name, size=15, color=NAVY, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, CL + 0.15, y + 0.44, 2.35, 0.34, sub, size=11.5, color=MUTED, anchor="m",
                 margins=(0, 0, 0, 0))
        for a, b, col in segs:
            add_rect(slide, tx0 + a * track_w, cy - 0.11, (b - a) * track_w, 0.22, fill=col)
        add_oval(slide, tx0, cy, 0.2, fill=WHITE, line=DARK, line_w=1.75)   # 观察起点
        add_oval(slide, tx1, cy, 0.2, fill=DARK, line=WHITE, line_w=1.25)   # 出院 / 死亡
        if sev_at is not None:  # 首次转重：轨道上方的倒三角
            mx = tx0 + sev_at * track_w
            add_rect(slide, mx - 0.12, cy - 0.38, 0.24, 0.2, fill=RED, shape=MSO_SHAPE.FLOWCHART_MERGE)
            add_text(slide, mx + 0.14, cy - 0.44, 1.2, 0.3, "首次转重", size=11, color=RED, bold=True,
                     anchor="m", margins=(0, 0, 0, 0))
        add_text(slide, 9.3, y, CL + CW - 9.3 - 0.1, lane_h, note, size=13.5, color=note_color, bold=True,
                 anchor="m", margins=(0, 0, 0, 0))
    add_text(slide, CL, 5.98, CW, 0.36,
             [[("补充：", {"bold": True, "color": NAVY}),
               ("GAM 的相关信息分散在病历叙述中，可能存在不同表述、否定语境或记录不完整，候选事件需结合临床记录核验。", {})]],
             size=12.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))


def build_time(slide, logo):
    """第 10 页（问题四 C）：三种数据形态的观测密度示意 + 两院时间现状。"""
    frame(slide, logo, "问题四", "C. 时间信息与连续观测是动态预测的关键",
          "现有数据适用于描述性和初步验证；动态病程建模需要更明确的时间信息。",
          ask_lead="请院方补充：", ask="真实入出院时间、首次转重时间，以及体征 / 护理的多次观测与时间戳。",
          notes=NOTES["time"])
    pw = 7.3
    add_rect(slide, CL, 1.8, pw, 4.45, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.86, pw - 0.4, 0.4, "同一位病人住院前 5 天的心率记录（示意）", size=15, color=NAVY,
             bold=True, anchor="m", margins=(0, 0, 0, 0))
    x0, x1 = 2.85, 7.6
    day_w = (x1 - x0) / 5
    rows = [
        ("理想", "逐次测量，真实测量时间"),
        ("类似 GAM", "每天 1–3 次，病历书写时间"),
        ("类似 GD", "每次住院 1–2 次，部分仅日期"),
    ]
    centers = [2.82, 3.74, 4.66]
    for (name, sub), cy in zip(rows, centers):
        add_text(slide, CL + 0.2, cy - 0.36, 2.1, 0.36, name, size=14, color=NAVY, bold=True, anchor="b",
                 margins=(0, 0, 0, 0))
        add_text(slide, CL + 0.2, cy, 2.15, 0.34, sub, size=11, color=MUTED, margins=(0, 0, 0, 0))
        add_line(slide, x0, cy, x1, cy, BORDER, 1)
        for k in range(6):
            add_line(slide, x0 + k * day_w, cy - 0.08, x0 + k * day_w, cy + 0.08, BORDER, 1)
    hx = lambda hour: x0 + hour / 120 * (x1 - x0)  # 把"第几小时"换算成横坐标
    for hour in range(2, 120, 4):                  # 理想：每 4 小时一次
        add_oval(slide, hx(hour), centers[0], 0.09, fill=MID)
    for hour in [10, 15, 33, 40, 58, 64, 82, 88, 106]:  # 类似 GAM：每天 1–3 次，空心表示"书写时间"
        add_oval(slide, hx(hour), centers[1], 0.13, fill=WHITE, line=MID, line_w=1.5)
    add_oval(slide, hx(10.5), centers[2], 0.13, fill=DARK)  # 类似 GD：一次到分钟的测量
    band = add_rect(slide, hx(48), centers[2] - 0.15, day_w, 0.3, fill="D6E4F5", line=MID, line_w=1,
                    dash=MSO_LINE_DASH_STYLE.DASH)          # 另一次只有日期：整天都可能
    shape_text(band, "只有日期", 10.5, NAVY, bold=True)
    for k in range(5):
        add_text(slide, x0 + k * day_w, 5.0, day_w, 0.28, f"第 {k + 1} 天", size=10.5, color=MUTED, align="c",
                 margins=(0, 0, 0, 0))
    # 图例
    ly = 5.48
    add_oval(slide, CL + 0.3, ly, 0.11, fill=MID)
    add_text(slide, CL + 0.42, ly - 0.15, 1.6, 0.3, "实际测量时间", size=11, color=BODY, anchor="m",
             margins=(0, 0, 0, 0))
    add_oval(slide, CL + 2.15, ly, 0.13, fill=WHITE, line=MID, line_w=1.5)
    add_text(slide, CL + 2.28, ly - 0.15, 1.6, 0.3, "病历书写时间", size=11, color=BODY, anchor="m",
             margins=(0, 0, 0, 0))
    add_rect(slide, CL + 3.95, ly - 0.09, 0.3, 0.18, fill="D6E4F5", line=MID, line_w=1, dash=MSO_LINE_DASH_STYLE.DASH)
    add_text(slide, CL + 4.32, ly - 0.15, 2.5, 0.3, "只有日期（当天任意时刻）", size=11, color=BODY, anchor="m",
             margins=(0, 0, 0, 0))
    add_text(slide, CL + 0.2, 5.72, pw - 0.4, 0.46,
             "示意图，数值虚构。每次住院体征时间点数（中位数）：GD 2 个，GAM 7–9 个。",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))

    rx, rw = CL + pw + 0.25, CW - pw - 0.25
    card(slide, rx, 1.8, rw, 2.12, "GD · 时间与观测点", header_h=0.42, header_size=14.5)
    add_text(slide, rx + 0.16, 2.32, rw - 0.32, 1.56, [
        {"runs": "尚无入出院时间：以首个有效体征时间\n作时间零点（292/297 次），\n不等同于真实入院时间",
         "bullet": True, "space_after": 8},
        {"runs": "检验 6,170 条、体征 484 条只有日期；\n保留精度标志，不推断为分钟", "bullet": True},
    ], size=14, color=BODY, margins=(0, 0, 0, 0))
    card(slide, rx, 4.08, rw, 2.17, "GAM · 文本与测量时间", header_h=0.42, header_size=14.5)
    add_text(slide, rx + 0.16, 4.6, rw - 0.32, 1.58, [
        {"runs": "体征从病历文字抽取，时间为书写时间\n（8,112 条），不是测量时间", "bullet": True, "space_after": 8},
        {"runs": "写法不规范会漏抽；\n如能提供结构化护理记录，可还原变化轨迹", "bullet": True},
    ], size=14, color=BODY, margins=(0, 0, 0, 0))


def build_text(slide, logo):
    """第 11 页（问题四 D）：文本三态编码 + 需识别的语境。"""
    frame(slide, logo, "问题四", "D. 文本记录的“未提及”需要独立编码",
          "文本提取结果是研究候选变量，需考虑临床语境与记录方式，并以核验后规则为准。",
          ask_lead="请院方确认：", ask="症状、治疗和重症状态的术语词表、否定规则及事件时间。",
          notes=NOTES["text"])
    states = [
        ("明确存在", GREEN, "“已插管”“已行无创通气”", "可作为候选发生线索（记 1）"),
        ("明确不存在", RED, "“未行插管”“无胸闷气促”", "识别否定表达（记 0）"),
        ("未提及 / 待核实", GREY, "病历中未检出相关词语", "不等同于临床上确定未发生（单独标记）"),
    ]
    add_text(slide, 2.95, 1.76, 5.0, 0.3, "病历中的写法", size=12, color=MUTED, margins=(0, 0, 0, 0))
    add_text(slide, 8.3, 1.76, 4.4, 0.3, "含义与编码", size=12, color=MUTED, margins=(0, 0, 0, 0))
    for i, (name, col, example, meaning) in enumerate(states):
        y = 2.08 + i * 0.9
        chip = add_rect(slide, CL, y, 2.2, 0.74, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        shape_text(chip, name, 16, WHITE, bold=True)
        add_rect(slide, 2.9, y, 5.2, 0.74, fill=PALE, line=BORDER)
        add_text(slide, 3.1, y, 4.9, 0.74, example, size=15, color=BODY, anchor="m")
        add_rect(slide, 8.25, y, CL + CW - 8.25, 0.74, fill=WHITE, line=BORDER)
        add_text(slide, 8.42, y, CL + CW - 8.6, 0.74, meaning, size=14, color=col if i < 2 else NAVY,
                 bold=True, anchor="m")
    card(slide, CL, 4.88, CW, 1.38, "还需识别的语境：以下写法不代表该项治疗已经实施", header_h=0.42, header_size=14.5)
    chips = [("“必要时插管”", "假设"), ("“拟转 ICU”", "计划"), ("“家属拒绝插管”", "拒绝"), ("“全麻下 / 术后插管”", "围术期")]
    x = CL + 0.2
    for text, tag in chips:
        w = text_width(text, 13.5) + text_width(tag, 11) + 0.55
        c = add_rect(slide, x, 5.43, w, 0.38, fill=LIGHT, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.5)
        shape_text(c, [[(text, {}), ("  " + tag, {"size": 11, "color": MUTED})]], 13.5, NAVY, bold=True)
        x += w + 0.15
    add_text(slide, CL + 0.2, 5.86, CW - 0.4, 0.36,
             [[("伪否定：", {"bold": True, "color": NAVY}),
               ("“无明显诱因出现发热”中的“无”并不否定发热，需先屏蔽这类短语再判断否定。", {})]],
             size=12.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))


def build_fields(slide, logo):
    """第 12 页（问题四 E）：其他字段的处理进展与待确认口径（原生表格）。"""
    frame(slide, logo, "问题四", "E. 其他字段：已完成的处理与待确认口径",
          "格式与字段的预处理已经开展；临床定义、检验方法和数据来源需进一步与院方确认。",
          ask_lead="依据：", ask="临床字段的最终解释以双方确认的数据字典及原始病历记录为依据。",
          notes=NOTES["fields"])
    first = lambda t: {"text": t, "fill": LIGHT, "color": NAVY, "bold": True}
    ask = lambda t: {"text": t, "fill": ASK_BG}
    rows = [
        ["字段类别", "已完成的技术处理", "需院方确认"],
        # 每格两行，断点放在分号处（"\n" 为段内换行）
        [first("标识与日期"), "兼容表头、跨行与混合日期格式\n年份笔误在住院时间窗内修正并标记（105 条）",
         ask("重复病案号、个别年份\nCRF 表头 TRT009 重复（ECMO 与“其他治疗”）")],
        [first("检验与单位"), "按标本区分同名检验（如尿 / 血肌酐）\n统一单位，并保留原单位",
         ask("IL-6、铁蛋白的单位与检测方法\nD-二聚体为 FEU 还是 DDU（相差约 2 倍）")],
        [first("治疗与并发症"), "建立治疗 / 诊断候选提取规则\n区分入院时已存在与住院后新发",
         ask("医嘱执行时间；急性肾损伤的判定标准\n并发症“判定依据”编码 1–4 的含义")],
        [first("缺失与隐私"), "人口学关联与隐私清洗\n逐文件隐私审计 0 命中",
         ask("空白、未记录与重复记录的含义\n（如空白是否代表“未评估”）")],
    ]
    add_table(slide, CL, 1.82, [2.0, 4.95, 5.28], [0.46] + [0.96] * 4, rows, size=14, header_size=14.5)


def build_actions(slide, logo):
    """第 13 页：分批核实清单（原生表格，第一列按批次合并）。"""
    frame(slide, logo, "下一步", "建议共同推进的核实事项",
          "感谢两家医院前期在数据整理与提供方面的支持；建议按对研究的影响分批确认。",
          ask_lead="核实完成后共同形成：", ask="最终字段字典 · 冻结版 CDM · 统一结局标签 · 可复现的实验数据。",
          notes=NOTES["actions"])
    item = lambda t: {"text": t, "bold": True, "color": NAVY}
    who = lambda t: {"text": t, "align": "c", "bold": True, "color": MID}
    rows = [
        ["批次", "事项", "具体内容", "主要涉及"],
        ["", item("① 转重定义"), "观察起点是否已重症、首次转重时间、转重后恢复的病例、未转重组", who("两院")],
        ["", item("② 时间锚点"), "真实入出院时间（最好到分钟）、流感检测时间、关键治疗开始时间", who("两院（GD 为主）")],
        ["", item("③ 单位与方法"), "D-二聚体 FEU / DDU；IL-6、铁蛋白的单位与检测方法", who("GD")],
        ["", item("④ 字段字典"),
         "TRT009 表头；判定依据 1–4；急性肾损伤判定标准\n呼吸机（OUT003）与无创 / 有创通气（TRT006 / 007）的关系",
         who("GD")],
        ["", item("⑤ 文本规则"), "症状 / 治疗的“有、明确无、未提及”判定；术语词表与否定规则", who("GAM")],
        ["", item("⑥ 补充病例"), "有明确事件时间的转重病例、明确未转重病例", who("两院")],
        ["", item("⑦ 导出规范"), "每行填写病案号、按文字格式保留前导零；说明空白格是“阴性”还是“未评估”", who("GD")],
    ]
    tbl = add_table(slide, CL, 1.8, [1.85, 1.75, 6.98, 1.65], [0.42, 0.55, 0.55, 0.5, 0.66, 0.55, 0.5, 0.55],
                    rows, size=13, header_size=14)
    groups = [(1, 2, "第一批", "决定主要结局"), (3, 5, "第二批", "决定两院可比性"), (6, 7, "持续推进", "样本与规范")]
    for r0, r1, name, sub in groups:
        cell = tbl.cell(r0, 0)
        cell.merge(tbl.cell(r1, 0))
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb("DCE7F5")
        tf = cell.text_frame
        for p in list(tf.paragraphs)[1:]:  # 合并后只保留一段，再重新写入
            p._p.getparent().remove(p._p)
        tf.paragraphs[0].text = ""
        fill_text(tf, [{"runs": name, "bold": True, "size": 15, "color": DARK, "align": "c"},
                       {"runs": sub, "size": 11.5, "color": MUTED, "align": "c"}], 13, NAVY, align="c")
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def build_closing(slide, logo):
    """第 14 页：致谢页（呼应封面的蓝色横带）。"""
    logo.stamp(slide)
    band = add_rect(slide, 0, 2.55, SW, 1.6, fill="4472C4", name="致谢横带")
    shape_text(band, "谢谢！欢迎讨论与指正", 40, WHITE)
    for r in band.text_frame.paragraphs[0].runs:  # 与封面标题一致：白字 + 下投影
        rPr = r._r.get_or_add_rPr()
        eff = etree.Element(qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="50800", dist="38100", dir="5400000", algn="ctr",
                              rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="45000")
        rPr.find(qn("a:solidFill")).addnext(eff)
    add_text(slide, 1.0, 4.45, SW - 2.0, 0.5, "感谢广安门医院、广东省中医院在数据整理与提供方面的支持",
             size=20, color=NAVY, align="c", anchor="m")
    add_text(slide, 1.0, 5.0, SW - 2.0, 0.4, "核实清单见第 13 页", size=14, color=MUTED, align="c", anchor="m")
    slide.notes_slide.notes_text_frame.text = NOTES["closing"]


# ---------------------------------------------------------------------------
# 主题、封面与删页
# ---------------------------------------------------------------------------
def fix_theme_fonts(prs):
    """把主题（幻灯片母版和备注母版）里空着的东亚字体改为微软雅黑。

    原主题 <a:ea typeface=""/>，而 Hans（简体中文）脚本字体是宋体，
    所以凡是没写东亚字体的中文都显示成宋体。改主题后，原封面中引用 +mn-ea 的文字也随之统一。
    """
    themes = [m.part.part_related_by(RT.THEME) for m in prs.slide_masters]
    themes.append(prs.notes_master.part.part_related_by(RT.THEME))
    for theme in themes:
        if hasattr(theme, "_element"):
            for ea in theme._element.iter(qn("a:ea")):
                if not ea.get("typeface"):
                    ea.set("typeface", FONT)
        else:
            theme._blob = theme.blob.replace(b'<a:ea typeface=""/>', f'<a:ea typeface="{FONT}"/>'.encode("utf-8"))


def fix_cover(cover):
    """封面只做小修：标题与副标题由 Times New Roman（中文回退宋体）改为微软雅黑；日期居中；替换备注。"""
    for shp in cover.shapes:
        if not shp.has_text_frame:
            continue
        text = shp.text_frame.text.strip()
        if text.startswith("老年流感") or text.startswith("理想数据"):
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    set_typeface(r._r.get_or_add_rPr())
                end = p._p.find(qn("a:endParaRPr"))
                if end is not None:
                    set_typeface(end)
        elif text.startswith("2026"):
            p = shp.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].text = p.runs[0].text.lstrip()  # 原稿用一串空格把日期"推"到中间
    cover.notes_slide.notes_text_frame.text = NOTES["cover"]


def delete_slides_after_first(prs):
    """删除第 2 页及以后的幻灯片（保留封面）。

    做法：从 presentation.xml 的幻灯片列表里移除条目，并断开到该页的关系；
    断开后这些页面不再被引用，保存时不会写入文件（其备注页也一并丢弃）。
    """
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst)[1:]:
        rid = sld_id.get(qn("r:id"))
        sld_id_lst.remove(sld_id)
        prs.part.drop_rel(rid)


def main():
    prs = Presentation(str(SRC))
    fix_theme_fonts(prs)
    cover = prs.slides[0]
    logo = LogoStamp(cover)  # 先从封面取校徽，再改封面
    fix_cover(cover)
    delete_slides_after_first(prs)
    blank = prs.slide_layouts.get_by_name("Blank")
    builders = [build_summary, build_ideal, build_sources, build_compare, build_work, build_records,
                build_events, build_course, build_time, build_text, build_fields, build_actions, build_closing]
    for build in builders:
        build(prs.slides.add_slide(blank), logo)
    prs.core_properties.title = DECK_TITLE
    prs.save(str(OUT))
    print(f"已生成：{OUT}（共 {len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
