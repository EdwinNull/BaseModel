"""flucase 标准导出（std_v3 / std_v2，parquet）→ CRF 式宽表（一个 Excel 文件，多个 sheet）+ 只含聚合量的分析报告。

以后同格式的导出（新快照、新中心、新编码、新口径、新增的列）都能直接处理，不用改脚本：
- 编码的中文名、单位：先用本脚本的编码表，其次用导出自带的 codes.csv，再次用 events 里最常见的 label / unit；
- 认不出前缀的数据项放进“14_其他数据项”；encounters 里本脚本没有专门处理的列，原样带进“01_就诊与人口学”末尾；
- 中心、场景、候选结局与口径、判据类、标记、复核条目都从数据里取；中心间差异按场景两两比较；
- 输入目录里有几份导出时（--mode auto）：互不重叠（如 GD、GAM 分开导出）就合成一份宽表；有重叠（同一批病例的新旧版本）
  就每份各出一份，不混用版本。--mode merge / each / latest 可改。

输入：flucase “数据导出”页下载的 zip，或解开后的目录（encounters / events / severity_timeline / outcomes / review 的 parquet、
manifest.json、codes.csv、data_dictionary.csv）；目录下的子目录、zip 都会找。同一份导出既有 zip 又有解开的目录，只算一次。
std_v2（observations.parquet）也认。

输出（默认写到输入目录旁边的 crf_wide/；本仓库的 data/ 已 gitignore），文件名带导出版本：
- CRF宽表_<版本>.xlsx：CRF sheet 一行一次就诊，第 1 行中文名、第 2 行字段名
  （pandas：pd.read_excel(文件, sheet_name="05_生命体征", header=1)）；分析 sheet 是聚合量（照实计数）；最后是字段字典。逐就诊，只留本机。
- 分析报告_<版本>.md：同一套分析，只含聚合量；计数 1–9 写作 "<10"、对应比例和分位数不给；不含就诊编号、具体日期和病历文字，
  写出前自检；可以回报给 Claude（检验项目名称、诊断编码只列名称与次数，D56；--no-names 不列）。
- 生成记录_<版本>.json：输入（版本、快照、sha256）、选项、脚本 sha256、行数，便于以后对照是哪份导出、哪版脚本生成的。

口径（都写进“说明”sheet）：
- 基线窗 = 预测时点 t0 前后 24 小时（D18；--window 可改）；基线值 = 窗内时间最早的一个有效值；“全程”含 t0 之后，不能直接当预测特征；
- 主分析人群 = 病原学确诊且 ≥60 岁（D16、D70），≥65 岁亚组（--age-main / --age-sub 可改）；同一患者多次就诊看 season_index（D73），不筛；
- 分级、候选结局取有效层（effective）；候选结局只是候选，用哪一列做标签由研究者定（A 档）；
- 标了 implausible（越界）的数值不进基线和极值（--keep-implausible 保留），条数记在 13_记录质量；
- unrecorded / 空 = 没有记录，不等于“无”。

用法（仓库根目录；系统 python 没有 pyarrow，用 flucase 的 venv）：
  flucase/backend/.venv/Scripts/python.exe scripts/crf_wide/build_crf_workbook.py --input data/export
  flucase/backend/.venv/Scripts/python.exe scripts/crf_wide/build_crf_workbook.py --input data/export/std_v3_20261101.zip
测试（只用虚构数据）：
  flucase/backend/.venv/Scripts/python.exe -m pytest scripts/crf_wide -q -p no:cacheprovider
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd

# ---------------------------------------------------------------- 编码表（与 flucase pipeline/cdm.py LABELS、repo/readers/vocab.py 同一套）
VITALS = {"VS001": ("体温", "℃"), "VS002": ("心率", "次/分"), "VS003": ("呼吸频率", "次/分"), "VS004": ("收缩压", "mmHg"),
          "VS005": ("舒张压", "mmHg"), "VS006": ("SpO2", "%")}
LABS = {"LAB001": ("白细胞", "10^9/L"), "LAB002": ("中性粒细胞绝对值", "10^9/L"), "LAB003": ("淋巴细胞绝对值", "10^9/L"),
        "LAB004": ("血红蛋白", "g/L"), "LAB005": ("血小板", "10^9/L"), "LAB006": ("C 反应蛋白", "mg/L"), "LAB007": ("降钙素原", "ng/mL"),
        "LAB008": ("白细胞介素-6", "pg/mL"), "LAB009": ("ALT", "U/L"), "LAB010": ("AST", "U/L"), "LAB011": ("肌酐", "μmol/L"),
        "LAB012": ("尿素", "mmol/L"), "LAB013": ("D-二聚体", "mg/L"), "LAB014": ("铁蛋白", "ng/mL"), "LAB015": ("LDH", "U/L"),
        "EXT_K": ("钾", "mmol/L"), "EXT_NA": ("钠", "mmol/L"), "EXT_ALB": ("白蛋白", "g/L"), "EXT_PAO2": ("氧分压（血气）", "mmHg"),
        "EXT_PACO2": ("二氧化碳分压（血气）", "mmHg"), "EXT_PF": ("氧合指数", "mmHg"), "EXT_SAO2": ("血氧饱和度（血气）", "%"),
        "EXT_LAC": ("乳酸", "mmol/L"), "EXT_TNI": ("肌钙蛋白 I", "ng/L"), "EXT_CK": ("肌酸激酶", "U/L"), "EXT_TBIL": ("总胆红素", "μmol/L"),
        "EXT_PT": ("凝血酶原时间", "s"), "EXT_FIB": ("纤维蛋白原", "g/L"), "EXT_NEUTP": ("中性粒细胞百分比", "%"),
        "EXT_LYMPHP": ("淋巴细胞百分比", "%"), "LAB_PAO2": ("动脉血氧分压", "mmHg"), "LAB_FIO2": ("吸入氧浓度", "%")}
SYMPTOMS = {"SYM001": "发热", "SYM002": "咳嗽", "SYM003": "咳痰", "SYM004": "咽痛", "SYM005": "肌痛", "SYM006": "乏力",
            "SYM007": "呼吸困难", "SYM008": "胸痛", "SYM009": "喘息", "SYM010": "厌食", "SYM011": "精神状态改变"}
COMORBID = {"COM001": "COPD", "COM002": "哮喘", "COM003": "支气管扩张", "COM004": "冠心病", "COM005": "心力衰竭", "COM006": "糖尿病",
            "COM007": "慢性肾脏病", "COM008": "肿瘤", "COM009": "脑卒中", "COM010": "免疫抑制", "COM011": "高血压", "COM014": "长期卧床/失能",
            "COM012": "肝病", "COM013": "神经肌肉病", "COM015": "其他神经系统疾病", "COM016": "血液系统疾病", "COM017": "肥胖",
            "COM018": "长期家庭氧疗", "COM019": "家庭无创通气"}
TREATMENTS = {"TRT001": "奥司他韦", "TRT002": "玛巴洛沙韦", "TRT003": "金银花口服液", "TRT004": "痰热清注射液", "TRT005": "化湿败毒颗粒",
              "TRT006": "无创通气", "TRT007": "有创通气", "TRT008": "ICU", "TRT009": "ECMO", "TRT010": "其他治疗", "TRT011": "氧疗",
              "TRT012": "升压药", "TRT013": "退热药", "TRT014": "高流量氧疗"}
COMPLICATIONS = {
    "CLP001": "病毒性肺炎", "CLP002": "细菌性或混合性肺炎", "CLP003": "Ⅰ型呼吸衰竭", "CLP004": "Ⅱ型呼吸衰竭", "CLP005": "ARDS",
    "CLP006": "胸腔积液", "CLP007": "气胸或纵隔气肿", "CLP008": "急性心力衰竭", "CLP009": "急性心肌损伤或心肌炎", "CLP010": "新发心律失常",
    "CLP011": "急性冠脉综合征", "CLP012": "脓毒症", "CLP013": "休克", "CLP014": "急性肾损伤", "CLP015": "急性肝损伤",
    "CLP016": "流感相关脑病或脑炎", "CLP017": "意识障碍或谵妄", "CLP018": "DIC", "CLP019": "横纹肌溶解", "CLP020": "MODS",
    "CLP021": "低钾血症", "CLP022": "低钠血症", "CLP023": "低蛋白血症", "CLP024": "急性肝衰竭", "CLPX_PNEU": "肺炎（未分型）",
    "CLPX_RF": "呼吸衰竭（未分型）"}
IMAGING = {"IMG_PNEUMONIA": "影像提示肺炎", "IMG_PROGRESS": "影像较前进展", "IMG_CT": "胸部 CT", "IMG_XR": "胸部 X 线", "IMG_MR": "磁共振",
           "IMG_US": "超声", "IMG_ECG": "心电图", "IMG_ENDO": "内镜", "IMG_OTHER": "其他检查"}
CRF_OUT = {"OUT001": "ICU", "OUT002": "重症", "OUT003": "呼吸机", "OUT004": "ECMO", "OUT005": "死亡", "OUT006": "出院",
           "OUT007": "住院天数", "OUT008": "自动出院"}
OTHER_NAMES = {"ONSET_DAYS": "起病时长", "FEVER_ONSET": "发热起始", "HOSP_ONSET": "院内发病", "HEIGHT": "身高", "WEIGHT": "体重",
               "TCM_SYN": "中医证候", "TCM_TYPE": "中医证型（方案）", "FU_ADMIT": "随访：住院", "FU_LINK": "随访关联", "DISPO": "去向",
               "DEATH": "死亡", "BMI": "BMI", "VACCINE": "本季流感疫苗", "FLU_TEST": "流感检测", "FLU_TYPE": "流感型别", "FLU_DX": "诊断中有流感"}
NAMES = {**{k: v for k, (v, _) in VITALS.items()}, **{k: v for k, (v, _) in LABS.items()}, **SYMPTOMS, **COMORBID, **TREATMENTS,
         **COMPLICATIONS, **IMAGING, **{k: f"原表结局 {v}" for k, v in CRF_OUT.items()}, **OTHER_NAMES}
UNIT = {k: u for d in (VITALS, LABS) for k, (_, u) in d.items()}
# 这些编码在 01–13 里有专门的列；其余的进“14_其他数据项”
COVERED = re.compile(r"^(?:VS\d|VSX:|LAB|EXT_|SYM|COM|TRT|CLP|IMG|DX:|OUT\d|"
                     r"(?:FLU_TEST|FLU_TYPE|FLU_DX|HEIGHT|WEIGHT|HOSP_ONSET|TCM_SYN|TCM_TYPE|FU_ADMIT|FU_LINK|DISPO|DEATH|BMI|VACCINE)$)")

OUTCOMES = ["D0", "O1", "O2", "O3", "O3_7d"]
VARIANTS = ["primary", "physiologic", "guideline", "primary_img", "all", "treatment", "diagnosis"]
LEVELS = ["moderate", "severe", "critical"]
LEVEL_CN = {"none": "未见中型及以上证据", "moderate": "中型", "severe": "重型", "critical": "危重型"}
RANK = {"moderate": 1, "severe": 2, "critical": 3}
KIND_CN = {"physiologic": "生理类", "treatment": "治疗类", "diagnosis": "诊断类"}
SETTING_CN = {"inpatient": "住院", "ed": "门急诊"}
TRI = "yes 有 / no 无 / unrecorded 未记录"
NOT_POS = ("hypothetical", "plan", "refused", "periop")
KEYS = ["site", "setting", "patient_id", "encounter_id"]
KEY_CN = {"site": "中心", "setting": "场景", "patient_id": "患者编号", "encounter_id": "就诊编号"}
T_BUCKETS = [(-math.inf, -168, "t0 前 >7 天"), (-168, -24, "t0 前 1–7 天"), (-24, 0, "t0 前 24 小时内"), (0, 24, "t0 后 24 小时内"),
             (24, 72, "t0 后 1–3 天"), (72, 168, "t0 后 3–7 天"), (168, math.inf, "t0 后 >7 天")]
FORMATS = ("std_v3", "std_v2")
# 问题单计数（A10）用的口径
QSHEET = "docs/汇报/临床负责人问题单_20261024.md"
FEVER_C = 37.3                                     # 发热阈值，腋温（D13；规则集 fever_threshold_c）
ICU_ORGANS = ("CLP005", "CLP008", "CLP014", "CLP016", "CLP018", "CLP020", "CLP024")   # 规则集 v0.3 ICU_ORGAN 认的其他器官衰竭

# 用到的列（缺了照样跑，缺哪些写进分析）
NEED = {
    "encounters": ["site", "patient_id", "encounter_id", "source", "setting", "age", "sex", "bmi", "t0", "t0_precision", "t0_rule",
                   "onset_date", "onset_source", "onset_to_t0_days", "fever_status", "fever_basis", "flu_confirmed", *COMORBID, "obese",
                   "high_risk", "vaccine", "discharge_time", "discharge_precision", "discharge_status", "died", "followup",
                   "index_admitted", "ruleset_version", "icu", "icu_start", "icu_periop", "death_time", "flu_season", "season_index"],
    "events": ["event_id", "site", "encounter_id", "var", "label", "module", "state", "value", "value_text", "unit", "time",
               "time_precision", "t_rel_h", "method", "reviewed", "changed_by_review", "rule_id", "time_kind", "end", "end_precision",
               "o2_status", "marks"],
    "severity_timeline": ["encounter_id", "layer", "row_type", "level", "start", "end", "criterion", "kind", "basis", "precision"],
    "outcomes": ["encounter_id", "setting", "outcome", "variant", "layer", "value", "event_time", "reason", "ruleset_version"],
    "review": ["encounter_id", "item_type", "field", "decision"],
}
ENC_SKIP = {"run_id", "source"}                    # encounters 里不带进宽表的列


@dataclass
class Options:
    window: tuple[float, float] = (-24.0, 24.0)
    keep_implausible: bool = False
    include_unmapped: bool = False
    age_main: float = 60.0
    age_sub: float = 65.0
    sites: list[str] | None = None
    suppress: bool = True
    names: bool = True

    @property
    def default_window(self) -> bool:
        return self.window == (-24.0, 24.0)

    @property
    def default_ages(self) -> bool:
        return (self.age_main, self.age_sub) == (60.0, 65.0)

    def pop_text(self) -> str:
        src = "D16、D70" if self.default_ages else "研究者另定的年龄"
        return f"确诊且 ≥{self.age_main:g} 岁；{src}"


# ---------------------------------------------------------------- 读入
@dataclass
class Bundle:
    """一份导出：一个目录，或 zip 里的一个目录。文件按需读，不常驻内存。"""
    name: str
    src: Path
    prefix: str | None = None                      # zip 里的目录前缀；None = 普通目录
    manifest: dict = field(default_factory=dict)
    sha: str = ""

    @property
    def generated(self) -> str:
        return str(self.manifest.get("generated_at") or "")

    @property
    def version(self) -> str:
        return str(self.manifest.get("version") or self.name)

    def read(self, fn: str) -> bytes | None:
        if self.prefix is None:
            p = self.src / fn
            return p.read_bytes() if p.is_file() else None
        with zipfile.ZipFile(self.src) as z:
            try:
                return z.read(self.prefix + fn)
            except KeyError:
                return None

    def table(self, name: str, columns: list[str] | None = None) -> pd.DataFrame | None:
        raw = self.read(f"{name}.parquet")
        return None if raw is None else pd.read_parquet(io.BytesIO(raw), columns=columns)

    def csv(self, name: str) -> pd.DataFrame | None:
        raw = self.read(name)
        return None if raw is None else pd.read_csv(io.BytesIO(raw), encoding="utf-8-sig")


def _json(b: bytes | None) -> dict:
    try:
        return json.loads(b.decode("utf-8")) if b else {}
    except (ValueError, UnicodeDecodeError):
        return {}


def find_bundles(root: Path) -> list[Bundle]:
    found: list[Bundle] = []
    if root.is_dir():
        for p in sorted(root.rglob("encounters.parquet")):
            if "__MACOSX" in p.parts:
                continue
            rel = p.parent.relative_to(root).as_posix()
            found.append(Bundle(root.name if rel == "." else rel, p.parent))
        zips = sorted(p for p in root.rglob("*.zip") if "__MACOSX" not in p.parts)
    else:
        zips = [root] if root.suffix.lower() == ".zip" else []
    for zp in zips:
        try:
            with zipfile.ZipFile(zp) as z:
                names = z.namelist()
        except zipfile.BadZipFile:
            continue
        for n in names:
            if n.rsplit("/", 1)[-1] == "encounters.parquet" and not n.startswith("__MACOSX"):
                pre = n[: -len("encounters.parquet")]
                found.append(Bundle(f"{zp.name}:{pre.rstrip('/') or '.'}", zp, pre))
    out, seen = [], set()
    for b in found:                                    # zip 与解开的目录是同一份：有 manifest 按 manifest 去重，没有就按 encounters + events 的内容
        raw = b.read("manifest.json")
        b.sha = hashlib.sha256(b.read("encounters.parquet") or b"").hexdigest()
        key = hashlib.sha256(raw).hexdigest() if raw else hashlib.sha256(
            b.sha.encode() + (b.read("events.parquet") or b.read("observations.parquet") or b"")).hexdigest()
        if key in seen:
            continue
        seen.add(key)
        b.manifest = _json(raw)
        out.append(b)
    return sorted(out, key=lambda b: (b.generated, b.name), reverse=True)


def plan_groups(bundles: list[Bundle], mode: str) -> tuple[list[list[Bundle]], str]:
    """几份导出怎么出宽表。bundles 已按生成时间从新到旧排好。"""
    if mode == "latest":
        return [bundles[:1]], "只用最新的一份导出"
    if mode == "each":
        return [[b] for b in bundles], "每份导出各出一份宽表"
    if mode == "merge" or len(bundles) == 1:
        return [bundles], "合并全部导出（同一就诊出现在几份里时保留较新的）" if len(bundles) > 1 else "一份导出"
    ids = [set(b.table("encounters", ["encounter_id"])["encounter_id"].dropna().astype(str)) for b in bundles]
    if any(a & b for a, b in combinations(ids, 2)):
        return [[b] for b in bundles], "几份导出有重叠的就诊（像是同一批病例的新旧版本），每份各出一份宽表；要合并用 --mode merge"
    return [bundles], f"{len(bundles)} 份导出互不重叠（如各中心分开导出），合成一份宽表"


@dataclass
class Data:
    bundles: list[Bundle]
    enc: pd.DataFrame
    ev: pd.DataFrame
    st: pd.DataFrame
    oc: pd.DataFrame
    rv: pd.DataFrame
    dictionary: pd.DataFrame | None
    codes: pd.DataFrame | None
    missing_cols: dict[str, list[str]]
    checks: list[tuple[str, object]]
    voc: "Vocab | None" = None


def _fill(df: pd.DataFrame | None, name: str, missing: dict) -> pd.DataFrame:
    if df is None:
        df = pd.DataFrame(columns=NEED[name])
    lack = [c for c in NEED[name] if c not in df.columns]
    if lack:
        missing.setdefault(name, []).extend(c for c in lack if c not in missing.get(name, []))
        for c in lack:
            df[c] = None                                  # None：与字符串比较得 False，不会冒出 pd.NA
    return df


def _site_col(df: pd.DataFrame | None) -> None:
    if df is not None and "site" not in df.columns and "source" in df.columns:
        df["site"] = df["source"].astype(object).fillna("").astype(str).str.removeprefix("REPO_")


def combine(bundles: list[Bundle], sites: list[str] | None = None) -> Data:
    """几份导出合成一份（同一就诊保留较新的那份）。bundles 已按生成时间从新到旧排好。"""
    parts: dict[str, list[pd.DataFrame]] = {k: [] for k in NEED}
    missing: dict[str, list[str]] = {}
    checks: list[tuple[str, object]] = []
    kept: set[str] = set()
    dup = 0
    for b in bundles:
        fmt = str(b.manifest.get("format") or "")
        if fmt and fmt not in FORMATS:
            checks.append((f"{b.name}：manifest 写的格式是 {fmt}，不是 {' / '.join(FORMATS)}", "照样按 std_v3 的列读"))
        enc = b.table("encounters")
        ev = b.table("events")
        if ev is None and (obs := b.table("observations")) is not None:        # std_v2
            ev = obs.rename(columns={"obs_id": "event_id"})
        _site_col(enc)
        _site_col(ev)
        if sites and "site" in enc.columns:
            enc = enc[enc["site"].isin(sites)]
        ids = set(enc["encounter_id"])
        dup += len(ids & kept)
        mine = ids - kept
        kept |= mine
        for name, df in (("encounters", enc), ("events", ev), ("severity_timeline", b.table("severity_timeline")),
                         ("outcomes", b.table("outcomes")), ("review", b.table("review"))):
            if df is None:
                continue
            if "encounter_id" in df.columns:
                df = df[df["encounter_id"].isin(mine)]
            parts[name].append(df)
        rc = b.manifest.get("row_counts") or {}
        for t, df in (("encounters", enc), ("events", ev)):
            want = rc.get(f"{t}.parquet")
            if want is not None and df is not None and not sites:
                checks.append((f"{b.name}：{t} 行数与 manifest 一致", "是" if int(want) == len(df) else f"否（manifest {want}，实际 {len(df)}）"))
    frames = {k: _fill(pd.concat(v, ignore_index=True) if v else None, k, missing) for k, v in parts.items()}
    enc = frames["encounters"].sort_values(["site", "setting", "encounter_id"], kind="stable").reset_index(drop=True)
    if len(bundles) > 1:
        checks.append(("同一就诊编号出现在不止一份导出（保留较新的）", dup))
    checks.append(("同一份导出里就诊编号重复（只留第一行）", int(enc["encounter_id"].duplicated().sum())))
    enc = enc.drop_duplicates("encounter_id").reset_index(drop=True)
    ev = frames["events"]
    checks.append(("events 里就诊编号不在 encounters 的行", int((~ev["encounter_id"].isin(enc["encounter_id"])).sum())))
    checks.append(("events 记录编号重复", int(ev["event_id"].duplicated().sum())))
    first = lambda fn: next((x for b in bundles if (x := b.csv(fn)) is not None), None)  # noqa: E731
    return Data(bundles, enc, ev, frames["severity_timeline"], frames["outcomes"], frames["review"], first("data_dictionary.csv"),
                first("codes.csv"), missing, checks)


# ---------------------------------------------------------------- 预处理
def to_dt(s: pd.Series) -> pd.Series:
    return pd.to_datetime(s.astype("string"), format="mixed", errors="coerce")


def _mode_of(df: pd.DataFrame, col: str) -> dict[str, str]:
    x = df.dropna(subset=[col])
    if not len(x):
        return {}
    g = x.groupby(["var", col]).size().reset_index(name="n").sort_values(["var", "n"], ascending=[True, False]).drop_duplicates("var")
    return dict(zip(g["var"].astype(str), g[col].astype(str)))


class Vocab:
    """编码 → 中文名、单位：本脚本的编码表优先，其次导出自带的 codes.csv，再次 events 里最常见的 label / unit。"""

    def __init__(self, ev: pd.DataFrame, codes: pd.DataFrame | None):
        self.names = dict(NAMES)
        if codes is not None and {"code", "label"} <= set(codes.columns):
            for c, lab in zip(codes["code"], codes["label"]):
                if isinstance(c, str) and isinstance(lab, str) and lab and lab != c:
                    self.names.setdefault(c, lab)
        for c, lab in _mode_of(ev, "label").items():
            if lab and lab != c:
                self.names.setdefault(c, lab)
        self.units = {**_mode_of(ev[ev["value"].notna()], "unit"), **UNIT}

    def name(self, code: str) -> str:
        return self.names.get(code, code)

    def unit(self, code: str) -> str:
        return self.units.get(code, "")


def prepare(d: Data) -> None:
    enc, ev = d.enc, d.ev
    for c in ("age", "bmi", "onset_to_t0_days"):
        enc[c] = pd.to_numeric(enc[c], errors="coerce")
    enc["_t0"] = to_dt(enc["t0"])
    setting = enc["setting"].astype(object).fillna("?").astype(str)
    enc["_grp"] = (enc["site"].astype(object).fillna("?").astype(str) + " " + setting.map(SETTING_CN).fillna(setting)).astype(object)
    enc["_set"] = setting.map(SETTING_CN).fillna(setting).astype(object)
    ev = ev[ev["encounter_id"].isin(enc["encounter_id"])].copy()
    ev["value"] = pd.to_numeric(ev["value"], errors="coerce")
    ev["t_rel_h"] = pd.to_numeric(ev["t_rel_h"], errors="coerce")
    ev["var"] = ev["var"].astype(object)                  # object：比较、str 方法都不会出 pd.NA
    ev["_marks"] = ev["marks"].astype(object).fillna("").astype(str)
    ev["_impl"] = ev["_marks"].str.contains("implausible", regex=False)
    ev["_pos"] = ev["state"] == "pos"
    ev["_date"] = ev["time"].astype("string").str[:10]
    ev["_grp"] = ev["encounter_id"].map(enc.set_index("encounter_id")["_grp"])
    d.ev = ev.sort_values(["encounter_id", "var", "t_rel_h", "event_id"], na_position="last", kind="stable").reset_index(drop=True)
    d.voc = Vocab(d.ev, d.codes)


def rel_h(times: pd.Series, eids: pd.Series, enc: pd.DataFrame) -> pd.Series:
    t0 = eids.map(enc.set_index("encounter_id")["_t0"])
    return ((to_dt(times) - t0).dt.total_seconds() / 3600).round(2)


def join_unique(s: pd.Series, limit: int = 8, sep: str = "；") -> str | None:
    vals = list(dict.fromkeys(str(x) for x in s.dropna() if str(x).strip()))
    if not vals:
        return None
    return sep.join(vals[:limit]) + (f"{sep}…（共 {len(vals)} 种）" if len(vals) > limit else "")


def _xs(s, code: str):
    """多重索引（就诊编号, 编码）里取一个编码；没有就给空的。"""
    if len(s) and code in s.index.get_level_values("var"):
        return s.xs(code, level="var")
    return pd.DataFrame(columns=s.columns) if isinstance(s, pd.DataFrame) else pd.Series(dtype=object)


# ---------------------------------------------------------------- CRF sheet
class Sheet:
    """一张 CRF sheet：按就诊编号对齐的列，带中文名与说明（字段字典用）。"""

    def __init__(self, name: str, enc: pd.DataFrame):
        self.name = name
        self.index = pd.Index(enc["encounter_id"], name="encounter_id")
        self.cols: dict[str, pd.Series] = {}
        self.meta: dict[str, tuple[str, str]] = {}
        self._df: pd.DataFrame | None = None

    def add(self, col: str, values, cn: str, note: str = "", count: bool = False) -> None:
        """count=True：计数列，没有记录的写 0。"""
        v = values.reindex(self.index) if isinstance(values, pd.Series) else pd.Series(values, index=self.index)
        self.cols[col] = v.fillna(0).astype(int) if count else v
        self.meta[col] = (cn, note)
        self._df = None

    @property
    def df(self) -> pd.DataFrame:
        if self._df is None:
            self._df = pd.concat(self.cols, axis=1).copy() if self.cols else pd.DataFrame(index=self.index)
        return self._df


def numeric_block(sh: Sheet, ev: pd.DataFrame, codes: list[str], o: Options, voc: Vocab, vitals: bool) -> None:
    x = ev[ev["_pos"] & ev["var"].isin(codes) & ev["value"].notna()]
    if not o.keep_implausible:
        x = x[~x["_impl"]]
    g = x.groupby(["encounter_id", "var"])
    allmin, allmax, n, days = g["value"].min(), g["value"].max(), g.size(), g["_date"].nunique()
    w = x[x["t_rel_h"].between(*o.window)]
    base = w.drop_duplicates(["encounter_id", "var"], keep="first").set_index(["encounter_id", "var"])
    wg = w.groupby(["encounter_id", "var"])["value"]
    wmin, wmax = wg.min(), wg.max()
    air = x[x["o2_status"].astype(object) == "空气"]
    air_all, air_win = air.groupby("encounter_id")["value"].min(), air[air["t_rel_h"].between(*o.window)].groupby("encounter_id")["value"].min()
    wtxt = f"t0{o.window[0]:+g}h 至 {o.window[1]:+g}h"
    for code in codes:
        nm, unit = voc.name(code), voc.unit(code)
        u = f"（{unit}）" if unit else ""
        b = _xs(base, code)
        sh.add(f"{code}_base", b.get("value", pd.Series(dtype=float)), f"{nm} 基线值{u}", f"{wtxt} 内时间最早的一个值（events.value，state=pos）")
        sh.add(f"{code}_base_h", b.get("t_rel_h", pd.Series(dtype=float)), f"{nm} 基线值相对 t0 小时", "events.t_rel_h；只有日期的按当天 00:00")
        if code == "VS006":
            sh.add("VS006_base_o2", b.get("o2_status", pd.Series(dtype=object)), "SpO2 基线值的吸氧状态", "events.o2_status；GD 按医嘱推断")
        if vitals:
            sh.add(f"{code}_win_min", _xs(wmin, code), f"{nm} 基线窗最低{u}", wtxt)
            sh.add(f"{code}_win_max", _xs(wmax, code), f"{nm} 基线窗最高{u}", wtxt)
            if code == "VS006":
                sh.add("VS006_win_min_air", air_win, "SpO2 基线窗吸空气最低（%）", "o2_status = 空气")
        sh.add(f"{code}_all_min", _xs(allmin, code), f"{nm} 全程最低{u}", "本次就诊全部记录（含 t0 之后，不能直接当预测特征）")
        sh.add(f"{code}_all_max", _xs(allmax, code), f"{nm} 全程最高{u}", "本次就诊全部记录（含 t0 之后）")
        if code == "VS006":
            sh.add("VS006_all_min_air", air_all, "SpO2 全程吸空气最低（%）", "o2_status = 空气；含 t0 之后")
        sh.add(f"{code}_n", _xs(n, code), f"{nm} 记录条数", "有数值的记录条数（全程）", count=True)
        if vitals:
            sh.add(f"{code}_days", _xs(days, code), f"{nm} 有记录的天数", "不同日期数（全程）", count=True)


def _tri(x: pd.DataFrame) -> pd.Series:
    """每个（就诊, 编码）：有一条 pos 记 yes，只有 neg 记 no，其余 unrecorded。"""
    g = x.assign(_p=x["state"] == "pos", _n=x["state"] == "neg").groupby(["encounter_id", "var"])[["_p", "_n"]].any()
    return pd.Series(np.where(g["_p"], "yes", np.where(g["_n"], "no", "unrecorded")), index=g.index)


def state_block(sh: Sheet, ev: pd.DataFrame, codes: dict[str, str], with_marks: bool = False, with_n: bool = False) -> None:
    x = ev[ev["var"].isin(codes)]
    st = _tri(x)
    p = x[x["_pos"]]
    first = p.groupby(["encounter_id", "var"])["t_rel_h"].min()
    n = p.groupby(["encounter_id", "var"]).size()
    marks = p.groupby(["encounter_id", "var"])["_marks"].agg(lambda s: join_unique(s.str.split(",").explode(), sep=","))
    for code, nm in codes.items():
        sh.add(code, _xs(st, code).reindex(sh.index).fillna("unrecorded"), nm, f"{TRI}；有一条 state=pos 记为 yes，只有 neg 记为 no")
        sh.add(f"{code}_first_h", _xs(first, code), f"{nm} 首次“有”相对 t0 小时", "空 = 没有“有”的记录或时间未知；文书来源多为书写时间")
        if with_n:
            sh.add(f"{code}_n", _xs(n, code), f"{nm} 记录条数", "", count=True)
        if with_marks:
            sh.add(f"{code}_marks", _xs(marks, code), f"{nm} 标记", "dx_icd_only 只认到编码 / dx_name_only 只认到名称 等")


def generic_block(sh: Sheet, ev: pd.DataFrame, codes: list[str], voc: Vocab) -> None:
    """认不出前缀的数据项：有没有、第一个数值、文字、首次时间、条数（只列这个编码实际有的东西）。"""
    x = ev[ev["var"].isin(codes)]
    st = _tri(x)
    p = x[x["_pos"]]
    num = p[p["value"].notna()].drop_duplicates(["encounter_id", "var"]).set_index(["encounter_id", "var"])["value"]
    txt = p.groupby(["encounter_id", "var"])["value_text"].agg(lambda s: join_unique(s, limit=3))
    first = p.groupby(["encounter_id", "var"])["t_rel_h"].min()
    n = p.groupby(["encounter_id", "var"]).size()
    for code in codes:
        nm, unit = voc.name(code), voc.unit(code)
        sh.add(code, _xs(st, code).reindex(sh.index).fillna("unrecorded"), nm, f"{TRI}（导出里的编码，本脚本没有专门处理）")
        if len(_xs(num, code)):
            sh.add(f"{code}_value", _xs(num, code), f"{nm} 第一个数值" + (f"（{unit}）" if unit else ""), "按相对 t0 小时最早的一条")
        if _xs(txt, code).notna().any():
            sh.add(f"{code}_text", _xs(txt, code), f"{nm} 文字", "最多列 3 种")
        sh.add(f"{code}_first_h", _xs(first, code), f"{nm} 首次“有”相对 t0 小时", "")
        sh.add(f"{code}_n", _xs(n, code), f"{nm} 记录条数", "", count=True)


def _codes_present(ev: pd.DataFrame, pattern: str, known: dict) -> list[str]:
    extra = sorted(set(ev.loc[ev["var"].str.match(pattern, na=False), "var"]) - set(known))
    return [*known, *extra]


def _dict_desc(d: Data, table: str, col: str) -> str | None:
    dic = d.dictionary
    if dic is None or not {"table", "column", "description"} <= set(dic.columns):
        return None
    hit = dic.loc[(dic["table"] == table) & (dic["column"] == col), "description"]
    return str(hit.iloc[0]) if len(hit) else None


def build_sheets(d: Data, o: Options) -> list[Sheet]:
    enc, ev, voc, sheets = d.enc, d.ev, d.voc, []
    e = enc.set_index("encounter_id")
    pos = ev[ev["_pos"]]

    def first_text(code: str) -> pd.Series:
        return pos[pos["var"] == code].groupby("encounter_id")["value_text"].agg(join_unique)

    def first_num(code: str) -> pd.Series:
        return pos[(pos["var"] == code) & pos["value"].notna()].groupby("encounter_id")["value"].first()

    def has(code: str) -> pd.Series:
        return pd.Series("yes", index=pd.Index(pos.loc[pos["var"] == code, "encounter_id"].unique()))

    # 01 就诊与人口学
    sh = Sheet("01_就诊与人口学", enc)
    for c, cn, note in [("age", "年龄（岁）", ""), ("sex", "性别", "M / F"), ("bmi", "BMI（kg/m²）", "有效层"), ("obese", "肥胖", "BMI ≥28.0 或诊断写明（D64）")]:
        sh.add(c, e[c], cn, note)
    sh.add("height_cm", first_num("HEIGHT"), "身高（cm）", "events HEIGHT 第一个值")
    sh.add("weight_kg", first_num("WEIGHT"), "体重（kg）", "events WEIGHT 第一个值")
    for c, cn, note in [("t0", "预测时点 t0", "入院 / 就诊时间；院内发病取首次流感阳性（D48）"), ("t0_precision", "t0 精度", "minute / day / note / unknown"),
                        ("t0_rule", "t0 取法", ""), ("flu_season", "流感季", "ISO 第 40 周至次年第 39 周（D73）"),
                        ("season_index", "同一患者多次就诊的索引状态", "index / repeat_within_14d / repeat_in_season / not_eligible（D73）")]:
        sh.add(c, e[c], cn, note)
    conf = e["flu_confirmed"].astype(object).eq("yes")
    sh.add("pop_main", np.where(conf & (e["age"] >= o.age_main), "yes", np.where(e["age"].isna() | ~conf, "unknown_or_no", "no")),
           f"主分析人群（{o.pop_text()}）", f"yes / no（年龄 <{o.age_main:g}）/ unknown_or_no（未确诊、只有阴性或年龄缺失）")
    sh.add("pop_sub", np.where(conf & (e["age"] >= o.age_sub), "yes", "no"), f"≥{o.age_sub:g} 岁亚组（确诊且 ≥{o.age_sub:g} 岁）",
           "D70" if o.default_ages else "研究者另定的年龄")
    for c, cn, note in [("onset_date", "起病日期", "有效层"), ("onset_source", "起病日期来源", ""), ("onset_to_t0_days", "起病到 t0 天数", "按日期相减；负数 = 起病晚于 t0（待核，R4）")]:
        sh.add(c, e[c], cn, note)
    sh.add("hosp_onset", has("HOSP_ONSET"), "院内发病", "events HOSP_ONSET；空 = 没有记录")
    for c, cn, note in [("fever_status", "发热（三态）", "yes / no / unknown（D13 阈值见规则集）"), ("fever_basis", "发热判定依据", ""),
                        ("vaccine", "本季流感疫苗", TRI), ("high_risk", "高危人群类别", "诊疗方案 p7–8（D64）；空 = 没有证据，不等于不是高危")]:
        sh.add(c, e[c], cn, note)
    sh.add("tcm_syndrome", first_text("TCM_SYN"), "中医证候", "events TCM_SYN 文字")
    sh.add("tcm_type", first_text("TCM_TYPE"), "中医证型（方案）", "events TCM_TYPE 文字")
    sh.add("ruleset_version", e["ruleset_version"], "规则集版本", "")
    known = set(NEED["encounters"]) | ENC_SKIP | set(KEYS)
    for c in [c for c in enc.columns if c not in known and not c.startswith("_") and not re.fullmatch(r"COM\d+", c)]:
        sh.add(c, e[c], _dict_desc(d, "encounters", c) or c, "encounters 原列（本脚本没有专门处理，原样带出）")
    sheets.append(sh)

    # 02 流感检测
    sh = Sheet("02_流感检测", enc)
    ft = ev[ev["var"] == "FLU_TEST"]
    fpos, fneg = ft[ft["state"] == "pos"], ft[ft["state"] == "neg"]
    txt = ft["value_text"].astype(object).fillna("").astype(str)
    meth = np.select([txt.str.contains("核酸|PCR", regex=True).to_numpy(bool), txt.str.contains("抗原", regex=False).to_numpy(bool)],
                     ["核酸", "抗原"], "其他或未写明")
    # flucase gd_raw-6 / gam_*-4 起读取器把方法写进中文名“流感检测（抗原 / 核酸 / NGS）”（D80：两项只写型别的认作抗原、加 NGS），有就用它
    lab_meth = ft["label"].astype(object).fillna("").astype(str).str.extract(r"流感检测（(.+)）", expand=False)
    meth = np.where(lab_meth.notna().to_numpy(), lab_meth.fillna("").to_numpy(), meth)
    sh.add("flu_confirmed", e["flu_confirmed"], "病原学确诊", "yes 检测阳性 / no 只有阴性 / unrecorded 来源没有检测记录（D16）")
    sh.add("flu_test_pos_n", fpos.groupby("encounter_id").size(), "阳性检测条数", "", count=True)
    sh.add("flu_test_neg_n", fneg.groupby("encounter_id").size(), "阴性检测条数", "", count=True)
    fp = fpos.sort_values("t_rel_h", na_position="last").drop_duplicates("encounter_id").set_index("encounter_id")
    sh.add("flu_pos_first_time", fp["time"], "首次阳性时间", "空 = 没有阳性或时间未知")
    sh.add("flu_pos_first_h", fp["t_rel_h"], "首次阳性相对 t0 小时", "")
    sh.add("flu_methods", pd.Series(meth, index=ft.index).groupby(ft["encounter_id"]).agg(join_unique), "检测方法",
           "读取器写的方法（抗原 / 核酸 / NGS，D80）；旧导出按检测名称文字归类：核酸 / 抗原 / 其他或未写明")
    sh.add("flu_type", ev[(ev["var"] == "FLU_TYPE") & ev["_pos"]].groupby("encounter_id")["value_text"].agg(join_unique), "型别", "events FLU_TYPE")
    sh.add("flu_test_text", ft.groupby("encounter_id")["value_text"].agg(join_unique), "检测名称与结果（原文）", "")
    sh.add("flu_dx", has("FLU_DX"), "诊断中有流感", "yes；空 = 没有记录")
    sh.add("flu_dx_text", first_text("FLU_DX"), "流感诊断名称", "")
    sheets.append(sh)

    # 03 基础病：encounters 里的 COM 列（有效层），加上只在 events 里出现的 COM 编码
    sh = Sheet("03_基础病", enc)
    com_enc = [c for c in COMORBID if c in enc.columns] + sorted(c for c in enc.columns if re.fullmatch(r"COM\d+", c) and c not in COMORBID)
    for c in com_enc:
        sh.add(c, e[c].astype(object).where(e[c].notna(), "unrecorded"), voc.name(c), f"encounters（有效层）；{TRI}，未记录不等于无")
    com_ev = sorted(set(ev.loc[ev["var"].str.match(r"^COM", na=False), "var"]) - set(com_enc))
    if com_ev:
        state_block(sh, ev, {c: voc.name(c) for c in com_ev})
    risk = e["high_risk"].astype("string").fillna("")
    sh.add("high_risk_n", risk.apply(lambda s: len([x for x in s.split("；") if x])), "高危类别数", "high_risk 按“；”计数")
    sheets.append(sh)

    # 04 症状
    sh = Sheet("04_症状", enc)
    state_block(sh, ev, {c: voc.name(c) for c in _codes_present(ev, r"^SYM", SYMPTOMS)})
    sheets.append(sh)

    # 05 生命体征
    sh = Sheet("05_生命体征", enc)
    numeric_block(sh, ev, _codes_present(ev, r"^VS\d", VITALS), o, voc, vitals=True)
    for code in sorted(set(ev.loc[ev["var"].str.startswith("VSX:", na=False), "var"])):
        x = pos[pos["var"] == code]
        sh.add(f"{code}_first", x.drop_duplicates("encounter_id").set_index("encounter_id")["value_text"], f"{code[4:]} 首次记录", "未映射的体征项目（文字）")
        sh.add(f"{code}_n", x.groupby("encounter_id").size(), f"{code[4:]} 记录条数", "", count=True)
    sheets.append(sh)

    # 06 检验
    sh = Sheet("06_检验", enc)
    labs = [c for c in _codes_present(ev, r"^(LAB|EXT_)", LABS) if not c.startswith("LABX:")]
    lx = pos[pos["var"].str.startswith("LABX:", na=False)]
    if o.include_unmapped:
        labs += sorted(set(lx.loc[lx["value"].notna(), "var"]))
    numeric_block(sh, ev, labs, o, voc, vitals=False)
    sh.add("labx_n", lx.groupby("encounter_id").size(), "未映射检验记录条数",
           "LABX:<原名称>" + ("，有数值的已按项目列在前面" if o.include_unmapped else "，不进宽表（--include-unmapped 可列出），名称清单见分析 sheet"), count=True)
    sheets.append(sh)

    # 07 影像
    sh = Sheet("07_影像", enc)
    state_block(sh, ev, {c: voc.name(c) for c in _codes_present(ev, r"^IMG", IMAGING)}, with_n=True)
    sheets.append(sh)

    # 08 治疗
    sh = Sheet("08_治疗", enc)
    trt = _codes_present(ev, r"^TRT", TREATMENTS)
    x = ev[ev["var"].isin(trt)]
    p = x[x["_pos"]]
    gp = p.groupby(["encounter_id", "var"])
    other = x[x["state"].isin(NOT_POS)].groupby(["encounter_id", "var"])["state"].agg(lambda s: join_unique(s, sep=","))
    firsts = p.drop_duplicates(["encounter_id", "var"], keep="first").set_index(["encounter_id", "var"])
    end_last = p.assign(_e=to_dt(p["end"])).groupby(["encounter_id", "var"])["_e"].max().dt.strftime("%Y-%m-%d %H:%M")
    names = gp["value_text"].agg(lambda s: join_unique(s, limit=5))
    for code in trt:
        nm = voc.name(code)
        f = _xs(firsts, code)
        sh.add(code, pd.Series("yes", index=f.index) if len(f) else pd.Series(dtype=object), nm, "yes = 有 state=pos 的记录；空 = 没有记录（不等于没用）")
        sh.add(f"{code}_other", _xs(other, code), f"{nm} 非确定语境", "hypothetical 假设 / plan 计划 / refused 拒绝 / periop 围术期（不计入“有”）")
        sh.add(f"{code}_start", f.get("time", pd.Series(dtype=object)), f"{nm} 首次开始", "最早的一条（按相对 t0 小时）")
        sh.add(f"{code}_start_h", f.get("t_rel_h", pd.Series(dtype=float)), f"{nm} 首次开始相对 t0 小时", "负数 = t0 前已开始")
        sh.add(f"{code}_end", _xs(end_last, code), f"{nm} 最晚结束", "空 = 未记录结束时间")
        sh.add(f"{code}_n", _xs(gp.size(), code), f"{nm} 记录条数", "", count=True)
        sh.add(f"{code}_names", _xs(names, code), f"{nm} 医嘱或药品名称", "原文，最多列 5 种")
    sheets.append(sh)

    # 09 并发症与诊断
    sh = Sheet("09_并发症与诊断", enc)
    state_block(sh, ev, {c: voc.name(c) for c in _codes_present(ev, r"^CLP", COMPLICATIONS)}, with_marks=True)
    dx = pos[pos["var"].str.startswith("DX:", na=False)]
    dx_txt = dx["var"].str[3:] + " " + dx["label"].astype("string").fillna("")
    sh.add("dx_raw", dx_txt.groupby(dx["encounter_id"]).agg(lambda s: join_unique(s, limit=30)), "原始诊断（编码 名称）", "DX:<院内编码>，最多列 30 个")
    sh.add("dx_raw_n", dx.groupby("encounter_id")["var"].nunique(), "原始诊断个数", "", count=True)
    sheets.append(sh)

    # 10 结局与去向
    sh = Sheet("10_结局与去向", enc)
    for c, cn, note in [("discharge_time", "离院时间", "GD 为锚点日期 + 住院天数推算"), ("discharge_precision", "离院时间精度", ""),
                        ("discharge_status", "去向", "门诊为就诊类型与是否之后住院")]:
        sh.add(c, e[c], cn, note)
    sh.add("los_days", ((to_dt(e["discharge_time"]) - e["_t0"]).dt.total_seconds() / 86400).round(2), "住院天数（离院 − t0）", "天；门诊为空")
    for c, cn, note in [("died", "院内死亡", ""), ("death_time", "死亡时间", "")]:
        sh.add(c, e[c], cn, note)
    sh.add("death_h", rel_h(e["death_time"], e.index.to_series(), enc), "死亡相对 t0 小时", "")
    for c, cn, note in [("icu", "进入 ICU", TRI + "（“呼吸与危重症医学科”不算）"), ("icu_start", "首次进入 ICU 时间", "")]:
        sh.add(c, e[c], cn, note)
    sh.add("icu_start_h", rel_h(e["icu_start"], e.index.to_series(), enc), "首次进入 ICU 相对 t0 小时", "")
    for c, cn, note in [("icu_periop", "进入 ICU 都是围术期", "照样算危重型（D66、D72），只供计数"),
                        ("followup", "随访是否完整", "complete / unknown（门急诊 O3 用）"), ("index_admitted", "门急诊当次收住院", "")]:
        sh.add(c, e[c], cn, note)
    fu = pos[pos["var"] == "FU_ADMIT"].drop_duplicates("encounter_id").set_index("encounter_id")
    sh.add("fu_admit_time", fu["time"], "之后住院时间（随访）", "events FU_ADMIT")
    sh.add("fu_admit_h", fu["t_rel_h"], "之后住院相对 t0 小时", "")
    sh.add("fu_link", has("FU_LINK"), "关联到全院记录", "events FU_LINK")
    sh.add("dispo_text", first_text("DISPO"), "去向原文", "events DISPO")
    for code in _codes_present(ev, r"^OUT\d", CRF_OUT):
        sh.add(code, first_num(code) if code == "OUT007" else has(code), voc.name(code),
               "原表住院天数（只作对照）" if code == "OUT007" else "原表勾选，yes / 空（只作对照，D24：OUT002 不读）")
    sheets.append(sh)

    # 11 分级时间线（有效层）
    sh = Sheet("11_分级时间线", enc)
    st = d.st[d.st["layer"] == "effective"].copy()
    st["_rank"] = st["level"].map(RANK)
    st["_h"] = rel_h(st["start"], st["encounter_id"], enc)
    bands, crit = st[st["row_type"] == "band"], st[st["row_type"] == "criterion"].sort_values(["encounter_id", "_h"], na_position="last")
    unt, sens = st[st["row_type"] == "untimed"], st[st["row_type"] == "sensitivity"]
    rk = bands.groupby("encounter_id")["_rank"].max()
    rk_any = pd.concat([bands, unt]).groupby("encounter_id")["_rank"].max()
    inv = {v: k for k, v in RANK.items()}
    sh.add("sev_max", rk.reindex(sh.index).map(inv).fillna("none").map(LEVEL_CN), "最高级别（有时间的证据）", "分级色带（band）；按全部判据")
    sh.add("sev_max_any", rk_any.reindex(sh.index).map(inv).fillna("none").map(LEVEL_CN), "最高级别（含时间未知的证据）", "band + untimed")
    for lv in LEVELS:
        cn = LEVEL_CN[lv]
        b = bands[bands["_rank"] >= RANK[lv]].sort_values("_h").drop_duplicates("encounter_id").set_index("encounter_id")
        c = crit[crit["_rank"] >= RANK[lv]].drop_duplicates("encounter_id").set_index("encounter_id")
        sh.add(f"{lv}_first", b["start"], f"首次达到{cn}及以上的时间", "色带开始时间")
        sh.add(f"{lv}_first_h", b["_h"], f"首次达到{cn}及以上相对 t0 小时", "")
        sh.add(f"{lv}_first_crit", c["criterion"], f"首次达到{cn}及以上的判据", "最早命中的判据（规则集编号与名称）")
        sh.add(f"{lv}_first_basis", c["basis"], "该判据的依据", "guideline 方案原文 / proxy 代理 / weak 弱代理")
    for kind in [*KIND_CN, *sorted(set(crit["kind"].dropna().astype(str)) - set(KIND_CN))]:
        k = crit[(crit["kind"] == kind) & (crit["_rank"] >= 2)]
        sh.add(f"severe_{kind}_first_h", k.groupby("encounter_id")["_h"].min(), f"{KIND_CN.get(kind, kind)}判据首次重型及以上相对 t0 小时",
               "判据类分开（D21、D28）")
    sh.add("criteria_hit", crit.groupby("encounter_id")["criterion"].agg(join_unique), "命中过的判据", "")
    sh.add("untimed", (unt["level"].map(LEVEL_CN).fillna(unt["level"]) + "：" + unt["criterion"].astype("string")).groupby(unt["encounter_id"]).agg(join_unique),
           "时间未知的阳性证据", "untimed_positive")
    sh.add("sensitivity_hit", sens.groupby("encounter_id")["criterion"].agg(join_unique), "只进敏感性口径的判据", "如影像进展弱代理（D63）")
    sheets.append(sh)

    # 12 候选结局（有效层）：结局与口径都从数据里取，已知的排在前面
    sh = Sheet("12_候选结局", enc)
    oc = d.oc[d.oc["layer"] == "effective"]
    piv = oc.drop_duplicates(["encounter_id", "outcome", "variant"]).pivot(index="encounter_id", columns=["outcome", "variant"], values="value")
    for oc_name in [*OUTCOMES, *sorted(set(oc["outcome"].dropna()) - set(OUTCOMES))]:
        for v in [*VARIANTS, *sorted(set(oc["variant"].dropna()) - set(VARIANTS))]:
            if (oc_name, v) in piv.columns:
                sh.add(f"{oc_name}_{v}", piv[(oc_name, v)], f"{oc_name} 候选结局（{v}）",
                       "event / event_unclean / unknown / excluded / no_event；候选，不是标签（A 档）")
        det = oc[(oc["outcome"] == oc_name) & (oc["variant"] == "primary") & oc["event_time"].notna()].drop_duplicates("encounter_id").set_index("encounter_id")
        if len(det):
            sh.add(f"{oc_name}_event_time", det["event_time"], f"{oc_name} 事件时间（primary）", "")
            sh.add(f"{oc_name}_reason", det["reason"], f"{oc_name} 判定理由（primary）", "")
    au = d.oc[d.oc["layer"] == "auto"].drop_duplicates(["encounter_id", "outcome", "variant"]).set_index(["encounter_id", "outcome", "variant"])["value"]
    ef = oc.drop_duplicates(["encounter_id", "outcome", "variant"]).set_index(["encounter_id", "outcome", "variant"])["value"]
    diff = (au.reindex(ef.index) != ef).groupby(level="encounter_id").sum()
    sh.add("auto_vs_effective_diff", diff, "自动层与有效层不同的结局个数", "复核改动带来的差异", count=True)
    sheets.append(sh)

    # 13 记录质量
    sh = Sheet("13_记录质量", enc)
    cnt = lambda m: ev[m].groupby("encounter_id").size()  # noqa: E731
    sh.add("n_events", ev.groupby("encounter_id").size(), "记录条数", "events 全部状态", count=True)
    sh.add("n_pos", pos.groupby("encounter_id").size(), "“有”的记录条数", "state = pos", count=True)
    for m in ("structured", "rule", "llm", "review"):
        sh.add(f"n_{m}", pos[pos["method"] == m].groupby("encounter_id").size(), f"来源方法 {m}", "state = pos", count=True)
    for pr in ("minute", "day", "note", "unknown"):
        sh.add(f"prec_{pr}", pos[pos["time_precision"] == pr].groupby("encounter_id").size(), f"时间精度 {pr}", "state = pos", count=True)
    timed_kind = pos["module"].isin(["vital", "lab", "treatment", "complication", "imaging"]) | (pos["var"] == "FLU_TEST")
    sh.add("n_untimed", pos[timed_kind & pos["t_rel_h"].isna()].groupby("encounter_id").size(), "时间未知（或 t0 缺失）的“有”记录",
           "只算体征、检验、治疗、并发症、影像、流感检测（结局、型别等本来就没有时间）", count=True)
    for mk, cn in (("implausible", "越界值"), ("value_fixed", "已自动修正的数值"), ("o2_inferred", "吸氧状态按医嘱推断")):
        sh.add(f"n_{mk}", cnt(ev["_marks"].str.contains(mk, regex=False)), cn, f"marks 含 {mk}", count=True)
    sh.add("n_censored", cnt(ev["_marks"].str.contains(r"(?:^|,)(?:lt|gt)(?:,|$)", regex=True)), "删失值（<、>）", "marks 含 lt / gt", count=True)
    sh.add("n_reviewed", cnt(ev["reviewed"].fillna(False).astype(bool)), "复核过的记录", "", count=True)
    sh.add("n_changed_by_review", cnt(ev["changed_by_review"].fillna(False).astype(bool)), "复核改过的记录", "", count=True)
    timed = pos[pos["t_rel_h"].notna()]
    sh.add("t_rel_min_h", timed.groupby("encounter_id")["t_rel_h"].min(), "最早记录相对 t0 小时", "")
    sh.add("t_rel_max_h", timed.groupby("encounter_id")["t_rel_h"].max(), "最晚记录相对 t0 小时", "")
    sh.add("n_before_30d", timed[timed["t_rel_h"] < -720].groupby("encounter_id").size(), "t0 前 30 天以前的记录", "待核：可能挂错就诊或年份笔误",
           count=True)
    sh.add("vital_days", pos[pos["var"].str.match(r"^VS\d", na=False)].groupby("encounter_id")["_date"].nunique(), "有体征记录的天数", "", count=True)
    sh.add("lab_days", pos[pos["var"].str.match(r"^(LAB|EXT_)", na=False)].groupby("encounter_id")["_date"].nunique(), "有检验记录的天数", "",
           count=True)
    sheets.append(sh)

    # 14 其他数据项：导出里有、上面没有专门处理的编码（新加的数据项会自动出现在这里）
    sh = Sheet("14_其他数据项", enc)
    rest = sorted(c for c in set(ev["var"].dropna()) if not COVERED.match(str(c)))
    generic_block(sh, ev, rest, voc)
    sheets.append(sh)
    return sheets


# ---------------------------------------------------------------- 分析（聚合量）
class Fmt:
    """聚合量的写法：suppress=True 时计数 1–9 写作 "<10"、对应比例与分位数不给（报告用）；否则照实（工作簿用）。"""

    def __init__(self, suppress: bool):
        self.s = suppress

    def n(self, k) -> object:
        k = int(k)
        return "<10" if self.s and 0 < k < 10 else k

    def kp(self, k, n) -> object:
        k, n = int(k), int(n)
        if n == 0:
            return "—"
        if self.s and 0 < k < 10:
            return "<10"
        return f"{k} ({k / n:.0%})"

    def q(self, s: pd.Series, digits: int = 1) -> object:
        s = pd.to_numeric(s, errors="coerce").dropna()
        if not len(s) or (self.s and len(s) < 10):
            return "—"
        a, b, c = np.percentile(s, [25, 50, 75])
        return f"{b:.{digits}f} [{a:.{digits}f}, {c:.{digits}f}]"


@dataclass
class Table:
    title: str
    df: pd.DataFrame
    note: str = ""


def _smd(x1: pd.Series, x2: pd.Series):
    x1, x2 = pd.to_numeric(x1, errors="coerce").dropna(), pd.to_numeric(x2, errors="coerce").dropna()
    if len(x1) < 2 or len(x2) < 2:
        return None
    sd = math.sqrt((x1.var() + x2.var()) / 2)
    return None if sd == 0 else round((x1.mean() - x2.mean()) / sd, 3)


def question_counts(d: Data, S: dict[str, pd.DataFrame], o: Options, f: Fmt, names: bool) -> list[Table]:
    """问题单（QSHEET）上要先在数据上数的数，按中心 × 场景。题号跟问题单走；“其中”行的比例分母是上一级。"""
    enc, ev = d.enc, d.ev
    e = enc.set_index("encounter_id")
    ids = e.index.to_series()
    grps = list(dict.fromkeys(enc["_grp"]))
    by = {g: enc.loc[enc["_grp"] == g, "encounter_id"].to_numpy() for g in grps}
    pos = ev[ev["_pos"]]
    s02, s03, s04, s05, s08, s09, s11 = (S[k] for k in ("02_流感检测", "03_基础病", "04_症状", "05_生命体征", "08_治疗", "09_并发症与诊断",
                                                         "11_分级时间线"))
    rows: list[list] = []

    def put(q: str, label: str, mask: pd.Series, parent: pd.Series | None = None, note: str = "") -> pd.Series:
        mask = mask.reindex(e.index).fillna(False).astype(bool)
        if parent is not None:
            mask &= parent
        rows.append([q, label, *[f.kp(mask.loc[by[g]].sum(), parent.loc[by[g]].sum() if parent is not None else len(by[g])) for g in grps], note])
        return mask

    def per_group(q: str, label: str, fn, note: str = "") -> None:
        rows.append([q, label, *[fn(g) for g in grps], note])

    # 1 发热
    no_hist = s04["SYM001"].astype(object).ne("yes")
    wmax = s05["VS001_win_max"]
    cool = put("1", f"基线窗（t0 {o.window[0]:+g} 至 {o.window[1]:+g} 小时）体温都 <{FEVER_C:g}℃ 或没测，且病史没写发热",
               (wmax.isna() | (wmax < FEVER_C)) & no_hist, note="体温 = 05_生命体征 VS001_win_max；病史 = 04_症状 SYM001")
    put("1", "  其中基线窗内没测体温", wmax.isna(), parent=cool)
    put("1", "整个就诊都不发热（fever_status = no）", e["fever_status"] == "no", note="fever_status 按整次就诊：任何时候测到 ≥阈值或病史写了发热都算发热")
    # 2 COPD 与 SpO2
    copd = put("2", "有 COPD（COM001 = yes）", s03["COM001"] == "yes")
    put("2", "  其中吸空气 SpO2 曾 ≤93%", s05["VS006_all_min_air"] <= 93, parent=copd, note="全程；o2_status = 空气")
    # 3 病例从哪里找
    conf_yes, dx = e["flu_confirmed"] == "yes", s02["flu_dx"] == "yes"
    put("3", "诊断写了流感、没有阳性检测", dx & ~conf_yes, note="flu_dx = yes 且 flu_confirmed 不是 yes")
    put("3", "有阳性检测、诊断没写流感", conf_yes & ~dx)
    # 4 弱阳性
    ft = ev[ev["var"] == "FLU_TEST"]
    ftp = ft[ft["_pos"]]
    weak_row = ftp["value_text"].astype(object).fillna("").astype(str).str.contains("弱阳性|±", regex=True)
    weak_ids, strong_ids = set(ftp.loc[weak_row, "encounter_id"]), set(ftp.loc[~weak_row, "encounter_id"])
    weak = put("4", "有弱阳性（±）检测结果", ids.isin(weak_ids), note="结果原文含“弱阳性”或“±”")
    put("4", "  其中没有别的阳性检测（只靠弱阳性确诊）", ids.isin(weak_ids - strong_ids), parent=weak)
    # 5 14 天内再住院
    rep = put("5", "14 天内再就诊（season_index = repeat_within_14d）", e["season_index"] == "repeat_within_14d")
    put("5", "  其中这次就诊达到重型及以上", s11["sev_max"].isin([LEVEL_CN["severe"], LEVEL_CN["critical"]]), parent=rep, note="有时间的证据（sev_max）")
    # 6 中医证候
    syn = pos[pos["var"] == "TCM_SYN"]
    typed = set(pos.loc[pos["var"] == "TCM_TYPE", "encounter_id"])
    put("6", "有中医证候记录", ids.isin(set(syn["encounter_id"])))
    per_group("6", "  证候写法（种）", lambda g: f.n(syn.loc[syn["_grp"] == g, "value_text"].nunique()), "不同写法的个数")
    per_group("6", "  其中已对到方案证型（种）",
              lambda g: f.n(syn.loc[(syn["_grp"] == g) & syn["encounter_id"].isin(typed), "value_text"].nunique()),
              "按就诊近似：同一次就诊认出了方案证型（TCM_TYPE），就算这次的写法已对应")
    # 8 新开始吸氧，及给氧前后 24 小时内是否同时有低氧（对齐方案 §7 第 11、12 项）
    st = d.st
    o2 = st[(st["layer"] == "effective") & (st["row_type"] == "criterion") & st["criterion"].astype(str).str.startswith("O2_NEW")]
    o2h = rel_h(o2["start"], o2["encounter_id"], enc).groupby(o2["encounter_id"]).min()
    vals = pos[pos["value"].notna() & (o.keep_implausible | ~pos["_impl"])]
    signs = {"吸空气 SpO2 ≤93%": vals[(vals["var"] == "VS006") & (vals["o2_status"].astype(object) == "空气") & (vals["value"] <= 93)],
             "PaO2/FiO2 ≤300": vals[(vals["var"] == "EXT_PF") & (vals["value"] <= 300)],
             "呼吸 ≥30 次/分": vals[(vals["var"] == "VS003") & (vals["value"] >= 30)]}
    near = {k: set(x.loc[(x["t_rel_h"] - x["encounter_id"].map(o2h)).abs() <= 24, "encounter_id"]) for k, x in signs.items()}
    new = put("8", "新开始吸氧（判据 O2_NEW 命中，有效层）", ids.isin(set(o2["encounter_id"])),
              note="“必要时”吸氧、长期家庭氧疗者规则集已排除（D67、D71）；时间取判据首次命中")
    h = o2h.reindex(e.index)
    put("8", "  其中首次在 t0 后 24 小时内（含 t0 前）", h <= 24, parent=new)
    put("8", "  其中首次在 t0 后 24 小时以后", h > 24, parent=new)
    put("8", "  其中时间未知", h.isna(), parent=new)
    timed = new & h.notna()
    hy = put("8", "  给氧前后 24 小时内同时有低氧（分母：有给氧时间的）", ids.isin(set().union(*near.values())), parent=timed,
             note="低氧 = 吸空气 SpO2 ≤93%、PaO2/FiO2 ≤300 或呼吸 ≥30 次/分（方案重型判据）；越界值不算")
    put("8", "  给氧前后 24 小时内没有低氧（分母：有给氧时间的）", ~hy, parent=timed)
    for k, s in near.items():
        put("8", f"    低氧里有{k}", ids.isin(s), parent=hy)
    # 9 ICU 与器官衰竭
    icu = put("9", "进入 ICU（icu = yes）", e["icu"] == "yes")
    organ = s09[[c for c in ICU_ORGANS if c in s09.columns]].eq("yes").any(axis=1)
    put("9", "  其中同时有其他器官衰竭", organ, parent=icu, note="ARDS、急性心衰、急性肾损伤、流感相关脑病、DIC、MODS、急性肝衰竭（规则集 v0.3 ICU_ORGAN）")
    # 10 ICU 科室名称：导出里没有
    per_group("10", "ICU 科室名称（问题单附件 2）", lambda g: "—", "std_v3 导出里没有科室名称，要从 flucase 的读取结果里取（对齐方案 §7 第 13 项）")
    # 11–13
    put("11", "进入 ICU 的就诊里，都是术后转入的（icu_periop）", e["icu_periop"].fillna(False).astype(bool), parent=icu,
        note="分母是第 9 题的进入 ICU；转入前 48 小时内有手术或病历写术后转入")
    for code in ("TRT006", "TRT007", "TRT014"):
        put("12", d.voc.name(code), s08[code] == "yes")
    put("13", f"确诊且 ≥{o.age_main:g} 岁（主分析人群）", conf_yes & (e["age"] >= o.age_main))
    put("13", f"确诊且 ≥{o.age_sub:g} 岁（亚组）", conf_yes & (e["age"] >= o.age_sub))
    # 14 SpO2 吸氧状态：按记录条数
    sp = pos[pos["var"] == "VS006"]
    inf, air = sp["_marks"].str.contains("o2_inferred", regex=False), sp["o2_status"].astype(object) == "空气"
    per_group("14", "SpO2 记录推断为吸空气（条，分母：全部 SpO2 记录）",
              lambda g: f.kp((inf & air)[sp["_grp"] == g].sum(), (sp["_grp"] == g).sum()), "按医嘱推断（marks 含 o2_inferred）")
    per_group("14", "SpO2 记录的吸氧状态是按医嘱推断的（条）", lambda g: f.kp(inf[sp["_grp"] == g].sum(), (sp["_grp"] == g).sum()))
    out = [Table("A10 问题单计数（按中心 × 场景）", pd.DataFrame(rows, columns=["题", "计数", *grps, "说明"]),
                 f"给 {QSHEET}；题号跟问题单走；“其中”行的比例分母是上一级；按就诊计数（第 14 题按记录条数）；"
                 "问题单只用住院（门诊只用第 3、13 题），其余组照算供参考")]
    if names:
        name = ft["value_text"].astype(object).fillna("（空）").astype(str).str.split("[：:]", n=1, regex=True).str[0]
        r = [[k, *[f.n(x.loc[x["_grp"] == g, "encounter_id"].nunique()) for g in grps], f.n(len(x)), f.n(x["_pos"].sum())]
             for k, x in sorted(ft.assign(_name=name).groupby("_name"), key=lambda kv: -len(kv[1]))]
        out.append(Table("A10b 流感检测名称（问题单附件 1）", pd.DataFrame(r, columns=["检测名称", *[f"{g} 就诊" for g in grps], "记录", "其中阳性记录"]),
                         "名称取结果原文冒号前的部分；只列名称与次数（D56）"))
        r = [[k, "是" if x["encounter_id"].isin(typed).any() else "否", *[f.n(x.loc[x["_grp"] == g, "encounter_id"].nunique()) for g in grps]]
             for k, x in sorted(syn.groupby("value_text"), key=lambda kv: -kv[1]["encounter_id"].nunique())]
        out.append(Table("A10c 中医证候写法（问题单附件 3）", pd.DataFrame(r, columns=["写法", "已对到方案证型", *[f"{g} 就诊" for g in grps]]),
                         "已对应按就诊近似；只列写法与次数"))
    return out


def analyse(d: Data, sheets: list[Sheet], o: Options, f: Fmt, names: bool) -> tuple[list[Table], list[str]]:
    enc, ev, voc, out = d.enc, d.ev, d.voc, []
    grps = list(dict.fromkeys(enc["_grp"]))
    by = {g: enc[enc["_grp"] == g] for g in grps}
    set_of = dict(zip(enc["_grp"], enc["_set"]))
    S = {s.name: s.df for s in sheets}
    pos = ev[ev["_pos"]]
    var_set = set(ev["var"].dropna())
    conf = lambda x: (x["flu_confirmed"] == "yes") & (x["age"] >= o.age_main)  # noqa: E731
    named = lambda c: names or ":" not in str(c)  # noqa: E731   带冒号的编码里是医院原名（LABX:、VSX:），--no-names 时报告不列
    new_codes = [c for c in S["14_其他数据项"].columns if c in var_set and named(c)]

    # A1 导出与结构
    rows = []
    for b in d.bundles:
        m = b.manifest
        rows += [("导出", b.name), ("  格式 / 版本", f"{m.get('format', '?')} / {m.get('version', '?')}"), ("  生成时间", m.get("generated_at")),
                 ("  快照", f"#{(m.get('snapshot') or {}).get('id')}，冻结 {(m.get('snapshot') or {}).get('frozen_at')}"),
                 ("  中心", "、".join((m.get("options") or {}).get("sites") or [])), ("  规则集", "、".join(m.get("ruleset_versions") or [])),
                 ("  读取格式（中心 → 格式@版本：就诊数）", json.dumps(m.get("reader_formats_by_site") or {}, ensure_ascii=False)),
                 ("  隐私正则命中（总数）", (m.get("privacy_audit") or {}).get("total_hits")),
                 ("  manifest 行数", json.dumps(m.get("row_counts") or {}, ensure_ascii=False))]
    if o.sites:
        rows.append(("只取这些中心（--sites）", "、".join(o.sites)))
    rows += [(k, v) for k, v in d.checks]
    rows += [(f"缺少的列：{t}", "、".join(c)) for t, c in d.missing_cols.items()]
    eff = d.oc[d.oc["layer"] == "effective"]
    have = eff.groupby("encounter_id")["outcome"].agg(set)
    exp = eff.assign(_s=eff["encounter_id"].map(enc.set_index("encounter_id")["setting"])).groupby("_s")["outcome"].agg(set)
    lack = sum(1 for eid, s in zip(enc["encounter_id"], enc["setting"]) if not exp.get(s, set()) <= have.get(eid, set()))
    no_oc = sorted({str(s) for s in enc["setting"].dropna()} - set(exp.index.astype(str)))
    unknown_lv = sorted(set(d.st["level"].dropna().astype(str)) - set(RANK))
    rows += [("候选结局缺行的就诊（同一场景别的就诊有、它没有的结局）", f.n(lack)),
             ("没有任何候选结局的场景", "、".join(no_oc) or "无"),
             ("分级时间线里不认识的级别（不进最高级别）", "、".join(unknown_lv) or "无"),
             ("没有任何 events 的就诊", f.n((~enc["encounter_id"].isin(ev["encounter_id"])).sum())),
             ("没有分级时间线行的就诊（未见中型及以上证据时本来就没有）", f.n((~enc["encounter_id"].isin(d.st["encounter_id"])).sum())),
             ("候选结局自动层与有效层不同的行（复核改动）", f.n(S["12_候选结局"]["auto_vs_effective_diff"].sum())),
             ("14_其他数据项里的编码", "、".join(new_codes) or "无"),
             ("基线窗（小时，相对 t0）", f"[{o.window[0]:g}, {o.window[1]:g}]")]
    out.append(Table("A1 导出与结构检查", pd.DataFrame(rows, columns=["项", "值"]).astype({"值": "string"})))

    # A2 人群与基本特征
    def row(label, fn):
        return [label, *[fn(by[g]) for g in grps]]
    rows = [row("就诊数", lambda x: f.n(len(x))), row("患者数", lambda x: f.n(x["patient_id"].nunique())),
            row("年龄 中位 [P25, P75]", lambda x: f.q(x["age"], 0)), row("年龄缺失", lambda x: f.kp(x["age"].isna().sum(), len(x)))]
    for a in sorted({o.age_main, o.age_sub, 80.0}):
        rows.append(row(f"年龄 ≥{a:g}", lambda x, a=a: f.kp((x["age"] >= a).sum(), len(x))))
    rows += [row("男性", lambda x: f.kp((x["sex"] == "M").sum(), len(x))), row("性别缺失", lambda x: f.kp(x["sex"].isna().sum(), len(x))),
             row("BMI 有记录", lambda x: f.kp(x["bmi"].notna().sum(), len(x))), row("BMI 中位 [P25, P75]", lambda x: f.q(x["bmi"])),
             row("肥胖 yes", lambda x: f.kp((x["obese"] == "yes").sum(), len(x))),
             row("起病到 t0 天数 有记录", lambda x: f.kp(x["onset_to_t0_days"].notna().sum(), len(x))),
             row("起病到 t0 天数 中位 [P25, P75]", lambda x: f.q(x["onset_to_t0_days"], 0)),
             row(f"主分析人群（{o.pop_text()}）", lambda x: f.kp(conf(x).sum(), len(x))),
             row("主分析人群中索引就诊（season_index = index）", lambda x: f.kp((conf(x) & (x["season_index"] == "index")).sum(), conf(x).sum()))]
    cats = {"fever_status": "发热", "flu_confirmed": "病原学确诊", "vaccine": "本季疫苗", "season_index": "索引状态（D73）",
            "t0_rule": "t0 取法", "t0_precision": "t0 精度", "onset_source": "起病日期来源", "discharge_status": "去向", "followup": "随访",
            "icu": "进入 ICU", "flu_season": "流感季"}
    for c, cn in cats.items():
        for v in sorted(enc[c].astype("string").fillna("（空）").unique()):
            rows.append(row(f"{cn}：{v}", lambda x, c=c, v=v: f.kp((x[c].astype("string").fillna("（空）") == v).sum(), len(x))))
    rows += [row("住院天数（离院 − t0）中位 [P25, P75]", lambda x: f.q(S["10_结局与去向"].loc[x["encounter_id"], "los_days"])),
             row("院内死亡", lambda x: f.kp(x["died"].fillna(False).astype(bool).sum(), len(x))),
             row("ICU 都是围术期", lambda x: f.kp(x["icu_periop"].fillna(False).astype(bool).sum(), len(x))),
             row("门急诊当次收住院", lambda x: f.kp(x["index_admitted"].fillna(False).astype(bool).sum(), len(x)))]
    risk = enc["high_risk"].astype("string").fillna("").str.split("；").explode()
    for v in sorted(set(risk) - {""}):
        rows.append(row(f"高危类别：{v}", lambda x, v=v: f.kp(x["high_risk"].astype("string").fillna("").str.split("；").apply(lambda s: v in s).sum(), len(x))))
    out.append(Table("A2 人群与基本特征（按中心 × 场景）", pd.DataFrame(rows, columns=["指标", *grps]),
                     f"比例的分母是该组就诊数；主分析人群 = {o.pop_text()}；season_index 导出不筛（D73）"))

    # A3 数据项覆盖（体征、检验）与 A4 数值分布
    vs, lab = S["05_生命体征"], S["06_检验"]
    vcodes = [c[:-5] for c in vs.columns if c.endswith("_base")]
    codes = vcodes + [c[:-5] for c in lab.columns if c.endswith("_base")]
    src_of = {c: (vs if c in vcodes else lab) for c in codes}
    num = pos[pos["value"].notna()]
    empty = [c for c in codes if not (src_of[c][f"{c}_n"] > 0).any()]
    codes = [c for c in codes if c not in empty and named(c)]
    r3, r4 = [], []
    for code in codes:
        src = src_of[code]
        x = num[num["var"] == code]
        units = "、".join(sorted(set(x["unit"].dropna().astype(str)))) or "—"
        a = [code, voc.name(code), UNIT.get(code, ""), units]
        b = [code, voc.name(code)]
        for g in grps:
            ids = by[g]["encounter_id"]
            nn = src.loc[ids, f"{code}_n"]
            a += [f.kp((nn > 0).sum(), len(ids)), f.kp(src.loc[ids, f"{code}_base"].notna().sum(), len(ids)), f.q(nn[nn > 0], 0)]
            b += [f.q(src.loc[ids, f"{code}_base"], 2), f.q(x.loc[x["_grp"] == g, "value"], 2)]
        a += [f.n(x["_impl"].sum()), f.kp((x["time_precision"] == "minute").sum(), len(x)) if len(x) else "—"]
        r3.append(a)
        r4.append(b)
    cols3 = ["编码", "名称", "统一单位", "实际单位"] + [f"{g}｜{k}" for g in grps for k in ("有记录", "基线窗有值", "条数/就诊")] + ["越界条数", "分钟精度"]
    none_txt = f"；全组都没有记录、不列出的编码：{'、'.join(empty)}" if empty else ""
    out.append(Table("A3 体征与检验覆盖", pd.DataFrame(r3, columns=cols3),
                     "有记录 = 至少一条有数值的记录的就诊；基线窗有值 = t0 前后窗内有值；条数/就诊 = 有记录的就诊里条数的中位 [P25, P75]；"
                     "统一单位为空 = 本脚本编码表里没有这个编码" + none_txt))
    cols4 = ["编码", "名称"] + [f"{g}｜{k}" for g in grps for k in ("基线值", "全部记录")]
    out.append(Table("A4 数值分布（中位 [P25, P75]）", pd.DataFrame(r4, columns=cols4), "越界值不计入（--keep-implausible 可保留）"))

    # A5 中心间差异：同一场景里有两个及以上中心时两两比较
    smd_rows: list[tuple[float, str]] = []
    e_ = enc.set_index("encounter_id")
    items = [("age", "年龄", e_["age"]), ("male", "男性（比例）", e_["sex"].eq("M").astype(float).where(e_["sex"].notna()))]
    items += [(f"{c}_base", f"{voc.name(c)} 基线值", src_of[c][f"{c}_base"]) for c in codes]

    def ms(x: pd.Series):
        x = pd.to_numeric(x, errors="coerce").dropna()
        return "—" if not len(x) or (f.s and len(x) < 10) else f"{x.mean():.2f} ± {x.std():.2f}"
    for setting in dict.fromkeys(set_of[g] for g in grps):
        gs = [g for g in grps if set_of[g] == setting]
        if len(gs) < 2:
            continue
        pairs = list(combinations(gs, 2))[:6]
        r5 = []
        for key, cn, s in items:
            xs = {g: s.reindex(by[g]["encounter_id"]) for g in gs}
            ks = {g: int(xs[g].notna().sum()) for g in gs}
            if not any(ks.values()):
                continue
            r = [key, cn]
            for g in gs:
                r += [f.kp(ks[g], len(xs[g])), ms(xs[g])]
            for g1, g2 in pairs:
                v = None if f.s and (ks[g1] < 10 or ks[g2] < 10) else _smd(xs[g1], xs[g2])      # 报告：任一边 <10 个值时 SMD 不给
                r.append(v)
                if v is not None:
                    smd_rows.append((v, f"{cn}（{g1} vs {g2}）"))
            r5.append(r)
        cols5 = ["字段", "名称", *[f"{g} {k}" for g in gs for k in ("有值", "均值 ± SD")], *[f"SMD {a} vs {b}" for a, b in pairs]]
        out.append(Table(f"A5 中心间差异（{setting}）", pd.DataFrame(r5, columns=cols5),
                         "基线值比较；SMD = (均值1 − 均值2) / √((SD1² + SD2²)/2)，|SMD| >0.1 常作不平衡的提示，>0.2 较明显"))

    # A6 症状、基础病、并发症、影像、治疗、其他数据项
    r6, empty6 = [], []
    pat = r"SYM\d+|COM\d+|CLP\w+?|IMG_\w+?|TRT\d+"
    for sheet in ("04_症状", "03_基础病", "09_并发症与诊断", "07_影像", "08_治疗", "14_其他数据项"):
        t = S[sheet]
        if sheet == "14_其他数据项":
            cs = [c for c in t.columns if c in var_set]
        else:
            cs = [c for c in t.columns if re.fullmatch(pat, c) and not c.endswith(("_n", "_h", "_marks"))]
        for c in cs:
            v_all = t[c].astype(object).fillna("unrecorded")
            if not v_all.isin(["yes", "no"]).any():
                empty6.append(c)
                continue
            if not named(c):
                continue
            r = [sheet[3:], c, voc.name(c)]
            for g in grps:
                v = v_all.reindex(by[g]["encounter_id"])
                r += [f.kp((v == "yes").sum(), len(v)), f.n((v == "no").sum()) if sheet != "08_治疗" else "—"]
            r6.append(r)
    out.append(Table("A6 各数据项有记录的就诊",
                     pd.DataFrame(r6, columns=["类", "编码", "名称", *[f"{g}｜{k}" for g in grps for k in ("有", "无")]]),
                     "症状、基础病、并发症与诊断、影像、治疗、其他数据项；有 = 至少一条 state=pos；无 = 只有否定记录；其余为未记录（不等于无）。治疗没有“无”的记录"
                     + (f"；全组都没有记录、不列出的编码：{'、'.join(c for c in empty6 if named(c))}" if empty6 else "")))
    t = S["08_治疗"]
    r6b = []
    for c in [c for c in t.columns if re.fullmatch(r"TRT\d+", c)]:
        h = t[f"{c}_start_h"]
        for g in grps:
            hh = h.reindex(by[g]["encounter_id"]).dropna()
            n_yes = (t.loc[by[g]["encounter_id"], c] == "yes").sum()
            if not n_yes:
                continue
            cut = pd.cut(hh, [b[0] for b in T_BUCKETS] + [math.inf], right=False, labels=[b[2] for b in T_BUCKETS])
            r6b.append([c, voc.name(c), g, f.n(n_yes), f.kp(n_yes - len(hh), n_yes),
                        *[f.n((cut == b[2]).sum()) for b in T_BUCKETS], f.kp(t.loc[by[g]["encounter_id"], f"{c}_end"].notna().sum(), n_yes)])
    out.append(Table("A6b 治疗首次开始时间相对 t0", pd.DataFrame(r6b, columns=["编码", "名称", "组", "有记录就诊", "开始时间未知", *[b[2] for b in T_BUCKETS], "记了结束时间"]),
                     "按每次就诊该治疗最早的一条"))

    # A7 分级与候选结局
    sv = S["11_分级时间线"]
    r7 = [[lv, *[f.kp((sv.loc[by[g]["encounter_id"], "sev_max"] == lv).sum(), len(by[g])) for g in grps]]
          for lv in [LEVEL_CN[x] for x in ("none", *LEVELS)]]
    r7 += [[f"含时间未知证据：{lv}", *[f.kp((sv.loc[by[g]["encounter_id"], "sev_max_any"] == lv).sum(), len(by[g])) for g in grps]]
           for lv in [LEVEL_CN[x] for x in ("none", *LEVELS)]]
    for lv in ("severe", "critical"):
        for lo, hi, lab_ in T_BUCKETS:
            r7.append([f"首次{LEVEL_CN[lv]}及以上：{lab_}",
                       *[f.n(sv.loc[by[g]["encounter_id"], f"{lv}_first_h"].between(lo, hi, inclusive="left").sum()) for g in grps]])
    out.append(Table("A7 分级（有效层）", pd.DataFrame(r7, columns=["指标", *grps]), "首次达到时间按色带开始时间；时间未知的证据不在时间分布里"))
    st = d.st[(d.st["layer"] == "effective") & d.st["row_type"].isin(["criterion", "untimed", "sensitivity"])]
    st = st.assign(_grp=st["encounter_id"].map(enc.set_index("encounter_id")["_grp"]))
    r7b = []
    for (crit, kind, basis, rt), x in st.groupby(["criterion", "kind", "basis", "row_type"], dropna=False):
        r7b.append([crit, KIND_CN.get(kind, kind if isinstance(kind, str) else "—"), basis if isinstance(basis, str) else "—", rt,
                    *[f.n(x.loc[x["_grp"] == g, "encounter_id"].nunique()) for g in grps]])
    out.append(Table("A7b 判据命中的就诊数（有效层）", pd.DataFrame(r7b, columns=["判据", "判据类", "依据", "行类型", *grps]),
                     "criterion = 首次命中；untimed = 时间未知；sensitivity = 只进敏感性口径"))
    ocg = eff.assign(_grp=eff["encounter_id"].map(enc.set_index("encounter_id")["_grp"]))
    vals = [*["event", "event_unclean", "unknown", "excluded", "no_event"],
            *sorted(set(ocg["value"].dropna().astype(str)) - {"event", "event_unclean", "unknown", "excluded", "no_event"})]
    r7c = []
    for (oname, v), x in ocg.groupby(["outcome", "variant"]):
        for g in grps:
            y = x[x["_grp"] == g]
            if len(y):
                r7c.append([oname, v, g, f.n(len(y)), *[f.kp((y["value"] == k).sum(), len(y)) for k in vals]])
    out.append(Table("A7c 候选结局（有效层；候选，不是标签）", pd.DataFrame(r7c, columns=["结局", "口径", "组", "就诊数", *vals]),
                     "event_unclean = 带 same_day 或时间未知证据；unknown = 不能判断（如门诊没能关联全院记录）"))

    # A8 记录层
    r8 = []
    for (g, mod), x in ev.groupby(["_grp", "module"], dropna=False):
        xp = x[x["_pos"]]
        r8.append([g, mod, f.n(len(x)), f.n(len(xp)), f.n(x["encounter_id"].nunique()),
                   *[f.kp((x["state"] == s).sum(), len(x)) for s in ("neg", "none", *NOT_POS)],
                   *[f.kp((xp["method"] == m).sum(), len(xp)) for m in ("structured", "rule", "llm", "review")],
                   *[f.kp((xp["time_precision"] == p).sum(), len(xp)) for p in ("minute", "day", "note", "unknown")],
                   f.kp(xp["t_rel_h"].isna().sum(), len(xp))])
    out.append(Table("A8 记录层（events）按组 × 类别", pd.DataFrame(r8, columns=[
        "组", "类别", "记录", "“有”的记录", "就诊", "neg", "none（未提及）", *NOT_POS, "structured", "rule", "llm", "review",
        "minute", "day", "note", "unknown", "时间未知"]), "状态比例分母为全部记录；方法、精度、时间未知的分母为“有”的记录"))
    r8b = []
    for (g, mod), x in pos.groupby(["_grp", "module"], dropna=False):
        cut = pd.cut(x["t_rel_h"], [b[0] for b in T_BUCKETS] + [math.inf], right=False, labels=[b[2] for b in T_BUCKETS])
        r8b.append([g, mod, f.n(len(x)), *[f.kp((cut == b[2]).sum(), len(x)) for b in T_BUCKETS], f.kp(cut.isna().sum(), len(x))])
    out.append(Table("A8b “有”的记录的时间分布（相对 t0）", pd.DataFrame(r8b, columns=["组", "类别", "记录", *[b[2] for b in T_BUCKETS], "时间未知"])))
    tokens = sorted({t for s in ev["_marks"].unique() for t in str(s).split(",") if t})     # 先看有哪些标记，再逐个计数
    r8c = []
    for m in tokens:
        hit = ev.loc[ev["_marks"].str.contains(rf"(?:^|,){re.escape(m)}(?:,|$)", regex=True), "_grp"]
        r8c.append([m, *[f.n((hit == g).sum()) for g in grps]])
    r8c.append(["复核过 / 复核改过", *[f"{f.n(ev.loc[ev['_grp'] == g, 'reviewed'].fillna(False).astype(bool).sum())} / "
                                     f"{f.n(ev.loc[ev['_grp'] == g, 'changed_by_review'].fillna(False).astype(bool).sum())}" for g in grps]])
    if len(d.rv):
        for (it, dec), x in d.rv.groupby(["item_type", "decision"], dropna=False):
            r8c.append([f"复核结论 {it} · {dec}", *[f.n(x["encounter_id"].map(enc.set_index("encounter_id")["_grp"]).eq(g).sum()) for g in grps]])
    out.append(Table("A8c 标记与复核（记录条数）", pd.DataFrame(r8c, columns=["标记", *grps]),
                     "lt/gt 删失、implausible 越界、value_fixed 自动修正、o2_inferred 吸氧状态按医嘱推断、periop 围术期、dx_icd_only / dx_name_only"))
    o2 = pos[pos["var"] == "VS006"]
    r8d = [[st_, *[f.kp((o2.loc[o2["_grp"] == g, "o2_status"].astype("string").fillna("（空）") == st_).sum(), (o2["_grp"] == g).sum()) for g in grps]]
           for st_ in sorted(o2["o2_status"].astype("string").fillna("（空）").unique())]
    out.append(Table("A8d SpO2 记录的吸氧状态", pd.DataFrame(r8d, columns=["吸氧状态", *grps])))
    if names:
        un = pos[pos["var"].str.match(r"^(LABX:|VSX:)", na=False) | (pos["module"] == "other") & ~pos["var"].str.startswith("DX:", na=False)]
        r8e = [[v, *[f.n(x.loc[x["_grp"] == g, "encounter_id"].nunique()) for g in grps], f.n(len(x))]
               for v, x in sorted(un.groupby("var"), key=lambda kv: -len(kv[1]))[:40]]
        out.append(Table("A8e 未进统一编码的项目（名称与次数，前 40）", pd.DataFrame(r8e, columns=["项目", *[f"{g} 就诊" for g in grps], "记录"]),
                         "LABX: 未映射检验、VSX: 未映射体征、其他 other 类；只列名称与次数（D56）"))
        dx = pos[pos["var"].str.startswith("DX:", na=False)]
        r8f = [[v[3:], x["label"].iloc[0], *[f.n(x.loc[x["_grp"] == g, "encounter_id"].nunique()) for g in grps]]
               for v, x in sorted(dx.groupby("var"), key=lambda kv: -kv[1]["encounter_id"].nunique())[:40]]
        out.append(Table("A8f 原始诊断编码（就诊数，前 40）", pd.DataFrame(r8f, columns=["编码", "名称", *grps]), "只列编码、名称与次数（D56）"))

    # A9 待核问题
    e = e_
    q = S["13_记录质量"]
    oe = S["10_结局与去向"]
    vn = [c for c in S["05_生命体征"].columns if re.fullmatch(r"VS\d+_n", c)]
    ln = [c for c in S["06_检验"].columns if c.endswith("_n") and c != "labx_n"]
    vb = [c for c in S["05_生命体征"].columns if re.fullmatch(r"VS\d+_base", c)]
    flags = [("t0 缺失", e["_t0"].isna(), "预测时点缺失时相对时间都算不出"),
             ("年龄缺失", e["age"].isna(), ""), ("性别缺失", e["sex"].isna(), ""),
             ("起病晚于 t0（起病到 t0 天数 <0）", e["onset_to_t0_days"] < 0, "R4：起病时间矛盾"),
             ("起病到 t0 >30 天", e["onset_to_t0_days"] > 30, "可能是慢性症状被当成起病"),
             ("离院早于 t0", oe["los_days"] < 0, ""), ("住院天数 >60", oe["los_days"] > 60, ""),
             ("院内死亡但没有死亡时间", e["died"].fillna(False).astype(bool) & e["death_time"].isna(), ""),
             ("去向写院内死亡但 died 不是 true", e["discharge_status"].astype("string").str.contains("死亡", na=False) & ~e["died"].fillna(False).astype(bool), ""),
             ("进入 ICU 但没有时间", (e["icu"] == "yes") & e["icu_start"].isna(), "R1 类：决定结局的事件没有时间"),
             ("没有任何生命体征", S["05_生命体征"][vn].sum(axis=1) == 0, ""),
             ("住院但没有任何检验", (e["setting"] == "inpatient") & (S["06_检验"][ln].sum(axis=1) == 0), "GD 检验表只有就诊号时会挂不上住院"),
             ("基线窗内没有任何体征", S["05_生命体征"][vb].notna().sum(axis=1) == 0, ""),
             ("有越界值", q["n_implausible"] > 0, "不进基线与极值"), ("有 t0 前 30 天以前的记录", q["n_before_30d"] > 0, "可能挂错就诊或年份笔误"),
             ("有时间未知的“有”记录", q["n_untimed"] > 0, ""),
             ("分级有时间未知的阳性证据", S["11_分级时间线"]["untimed"].notna(), "untimed_positive"),
             ("病原学确诊 unrecorded", e["flu_confirmed"] == "unrecorded", "来源没有检测记录，不当阴性（D16）"),
             ("门急诊随访 unknown", (e["setting"] == "ed") & (e["followup"] == "unknown"), "O3 记为未知"),
             ("主分析人群里的非索引就诊", conf(e) & (e["season_index"] != "index"), "D73：同一病程或同季更晚")]
    r9 = []
    for title, mask, note in flags:
        mask = mask.reindex(e.index).fillna(False).astype(bool)
        r9.append([title, *[f.n(mask.reindex(by[g]["encounter_id"]).sum()) for g in grps], note])
    bad_units = []
    for code, u in UNIT.items():
        x = num[(num["var"] == code) & num["unit"].notna()]
        other = sorted(set(x["unit"].astype(str)) - {u})
        if other:
            bad_units.append(f"{code} {voc.name(code)}：{'、'.join(other)}（应为 {u}）")
    multi = [f"{c} {voc.name(c)}：{'、'.join(sorted(set(x.astype(str))))}" for c, x in num.loc[~num["var"].isin(UNIT) & num["unit"].notna()].groupby("var")["unit"]
             if x.nunique() > 1 and named(c)]
    r9.append(["单位与统一单位不同的数据项", *[""] * len(grps), "；".join(bad_units) or "无"])
    r9.append(["编码表里没有、导出里单位不止一种的数据项", *[""] * len(grps), "；".join(multi) or "无"])
    out.append(Table("A9 待核问题（就诊数）", pd.DataFrame(r9, columns=["问题", *grps, "说明"]), "只是提示，处理与否由研究者定"))
    out += question_counts(d, S, o, f, names)

    # 要点
    key = ["就诊数：" + "，".join(f"{g} {f.n(len(by[g]))}" for g in grps),
           f"主分析人群（{o.pop_text()}）：" + "，".join(f"{g} {f.kp(conf(by[g]).sum(), len(by[g]))}" for g in grps)]
    if smd_rows:
        big = sorted(smd_rows, key=lambda r: -abs(r[0]))[:6]
        key.append("中心间差异最大的字段（|SMD|）：" + "；".join(f"{lab_} {v:+.2f}" for v, lab_ in big))
    gap = []
    for code in codes:
        cov = {g: src_of[code].loc[by[g]["encounter_id"], f"{code}_n"].gt(0).mean() for g in grps if len(by[g])}
        if len(cov) >= 2 and max(cov.values()) - min(cov.values()) >= 0.5:
            gap.append(f"{voc.name(code)}（" + " / ".join(f"{g} {v:.0%}" if not f.s or len(by[g]) >= 10 else f"{g} —" for g, v in cov.items()) + "）")
    if gap:
        key.append("覆盖率组间相差 ≥50 个百分点：" + "；".join(gap[:12]) + ("…" if len(gap) > 12 else ""))
    nz = [r for r in r9[:-2] if any(v not in (0, "0") for v in r[1:-1])]
    if nz:
        key.append("待核问题（就诊数）：" + "；".join(f"{r[0]}（" + "，".join(f"{g} {v}" for g, v in zip(grps, r[1:-1]) if v not in (0, "0")) + "）" for r in nz))
    if bad_units:
        key.append("单位与统一单位不同：" + "；".join(bad_units))
    if new_codes:
        key.append(f"本脚本没有专门处理、放进 14_其他数据项的编码 {len(new_codes)} 个：" + "、".join(new_codes[:20]) + ("…" if len(new_codes) > 20 else ""))
    return out, key


# ---------------------------------------------------------------- 写出
def _cell(v):
    if v is None or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, (np.generic,)):
        v = v.item()
        if isinstance(v, float) and math.isnan(v):
            return None
    if isinstance(v, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        return v[:32000]
    return v


def _sheet_name(title: str, used: set[str]) -> str:
    head, _, rest = title.partition(" ")
    name = re.sub(r"[\\/*?:\[\]]", "", f"{head}_{rest[:24]}" if rest else head)[:31]
    base, k = name, 2
    while name in used:
        name = f"{base[:28]}_{k}"
        k += 1
    used.add(name)
    return name


def write_workbook(path: Path, d: Data, sheets: list[Sheet], tables: list[Table], key: list[str], info: list[tuple[str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    sub_font = Font(color="666666", size=9)
    wb = Workbook(write_only=True)
    used: set[str] = set()

    def hcell(ws, v, font=bold, fill=head_fill):
        c = WriteOnlyCell(ws, value=v)
        c.font, c.alignment = font, Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        return c

    ws = wb.create_sheet(_sheet_name("说明", used))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 120
    ws.append([hcell(ws, "CRF 式宽表（flucase 标准导出 → 每次就诊一行）"), None])
    for k, v in info:
        ws.append([hcell(ws, k, fill=None), _cell(v)])
    ws.append([])
    ws.append([hcell(ws, "要点（自动生成，供核对）"), None])
    for k in key:
        ws.append([None, k])

    keys = d.enc.set_index("encounter_id")[["site", "setting", "patient_id"]]
    for sh in sheets:
        ws = wb.create_sheet(_sheet_name(sh.name, used))
        df = keys.join(sh.df, how="right").reset_index()[[*KEYS, *sh.df.columns]]
        ws.freeze_panes = "E3"
        for i, _ in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 22 if i == 4 else 13 if i > 4 else 9
        ws.append([hcell(ws, KEY_CN.get(c) or sh.meta.get(c, (c, ""))[0]) for c in df.columns])
        ws.append([hcell(ws, c, font=sub_font, fill=None) for c in df.columns])
        for row in df.itertuples(index=False, name=None):
            ws.append([_cell(v) for v in row])

    for t in tables:
        ws = wb.create_sheet(_sheet_name(t.title, used))
        ws.append([hcell(ws, t.title)])
        if t.note:
            ws.append([WriteOnlyCell(ws, value=t.note)])
        ws.append([hcell(ws, str(c)) for c in t.df.columns])
        for i, _ in enumerate(t.df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 30 if i <= 2 else 16
        for row in t.df.itertuples(index=False, name=None):
            ws.append([_cell(v) for v in row])

    ws = wb.create_sheet(_sheet_name("字段字典", used))
    ws.append([hcell(ws, x) for x in ("sheet", "字段名", "中文名", "说明")])
    for i, w in enumerate((18, 30, 34, 90), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for k in KEYS:
        ws.append(["（每个 CRF sheet）", k, KEY_CN[k], "encounters 原列"])
    for sh in sheets:
        for c, (cn, note) in sh.meta.items():
            if c in sh.df.columns:
                ws.append([sh.name, c, cn, note])
    if d.dictionary is not None:
        ws.append([])
        ws.append([hcell(ws, "附：导出自带的数据字典（data_dictionary.csv）")])
        cols = [c for c in ("table", "column", "description", "values", "missing") if c in d.dictionary.columns]
        ws.append([hcell(ws, c) for c in cols])
        for row in d.dictionary[cols].itertuples(index=False, name=None):
            ws.append([_cell(v) for v in row])
    wb.save(path)


def _md_table(df: pd.DataFrame) -> str:
    def esc(v):
        v = _cell(v)
        return "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(esc(c) for c in df.columns) + " |", "|" + "---|" * len(df.columns)]
    lines += ["| " + " | ".join(esc(v) for v in r) + " |" for r in df.itertuples(index=False, name=None)]
    return "\n".join(lines)


DATE_RE = re.compile(r"(?<!\d)(?:19|20)\d\d[-/.年](?:0?[1-9]|1[0-2])[-/.月](?:0?[1-9]|[12]\d|3[01])(?!\d)")


def leak_check(text: str, d: Data) -> list[str]:
    """报告写出前的自检：不能带就诊编号、患者编号（含去掉中心前缀的部分）和具体日期（导出、冻结、生成时间除外）。"""
    found = []
    ids = set(d.enc["encounter_id"].dropna().astype(str)) | set(d.enc["patient_id"].dropna().astype(str))
    ids |= {i.split(":", 1)[1] for i in ids if ":" in i}
    ids = sorted((i for i in ids if len(i) >= 5), key=len, reverse=True)
    hit = 0
    for k in range(0, len(ids), 500):                    # 整词匹配，避免把更长的数字串算进来
        pat = re.compile(r"(?<![0-9A-Za-z])(?:" + "|".join(map(re.escape, ids[k:k + 500])) + r")(?![0-9A-Za-z])")
        hit += len(set(pat.findall(text)))
    if hit:
        found.append(f"含就诊或患者编号 {hit} 个")
    meta = ("生成时间", "冻结")
    n = sum(len(DATE_RE.findall(line)) for line in text.splitlines() if not any(m in line for m in meta))
    if n:
        found.append(f"含具体日期 {n} 处")
    return found


def write_report(path: Path, tables: list[Table], key: list[str], info: list[tuple[str, str]], d: Data, check: bool = True) -> list[str]:
    md = ["# 导出分析报告（只含聚合量）", "",
          "- 计数 1–9 写作 \"<10\"，对应的比例、分位数和 SMD 不给；不含就诊编号、具体日期和病历文字。",
          "- 生成脚本：`scripts/crf_wide/build_crf_workbook.py`；逐就诊的 CRF 宽表在同目录的 xlsx 里，只留本机。", ""]
    md += [f"- {k}：{v}" for k, v in info if k in ("输入", "怎么组合", "导出版本", "规则集", "基线窗", "主分析人群", "越界值", "未映射检验", "生成时间")]
    md += ["", "## 要点（自动生成，供研究者核对）", ""] + [f"- {k}" for k in key]
    for t in tables:
        md += ["", f"## {t.title}", ""]
        if t.note:
            md += [f"> {t.note}", ""]
        md.append(_md_table(t.df))
    text = "\n".join(md) + "\n"
    problems = leak_check(text, d) if check else []
    if not problems:
        path.write_text(text, encoding="utf-8")
    return problems


# ---------------------------------------------------------------- 主流程
@dataclass
class Result:
    tag: str
    xlsx: Path | None
    md: Path | None
    record: Path | None
    n_enc: int
    n_ev: int
    problems: list[str] = field(default_factory=list)
    skipped: str = ""


def _tag(group: list[Bundle], o: Options, used: set[str]) -> str:
    tag = re.sub(r"[^\w.+-]+", "_", "+".join(dict.fromkeys(b.version for b in group)))[:80] or "export"
    if o.sites:
        tag += "_" + "-".join(o.sites)
    base, k = tag, 2
    while tag in used:
        tag = f"{base}_{k}"
        k += 1
    used.add(tag)
    return tag


def build_one(group: list[Bundle], how: str, out_dir: Path, tag: str, o: Options, skip_existing: bool = False) -> Result:
    xlsx, md, rec = out_dir / f"CRF宽表_{tag}.xlsx", out_dir / f"分析报告_{tag}.md", out_dir / f"生成记录_{tag}.json"
    if skip_existing and xlsx.exists():
        return Result(tag, xlsx, md if md.exists() else None, rec if rec.exists() else None, 0, 0, skipped="宽表已存在（--skip-existing）")
    d = combine(group, o.sites)
    if not len(d.enc):
        return Result(tag, None, None, None, 0, 0, skipped="所选中心在这份导出里没有就诊")
    prepare(d)
    sheets = build_sheets(d, o)
    info = [("输入", f"{len(group)} 份导出：{'、'.join(b.name for b in group)}"), ("怎么组合", how),
            ("导出版本", "、".join(dict.fromkeys(b.version for b in group))),
            ("规则集", "、".join(sorted({str(v) for v in d.enc["ruleset_version"].dropna()}))),
            ("中心 × 场景", "，".join(f"{g} {n}" for g, n in d.enc["_grp"].value_counts(sort=False).items())),
            ("基线窗", f"t0 {o.window[0]:+g} 至 {o.window[1]:+g} 小时（{'D18' if o.default_window else '不是 D18 的默认 ±24 小时'}）；"
                    "基线值 = 窗内时间最早的一个有效值"),
            ("主分析人群", f"{o.pop_text()}；≥{o.age_sub:g} 岁亚组；只打标记，不筛"),
            ("越界值", "保留" if o.keep_implausible else "标了 implausible 的数值不进基线与极值（条数见 13_记录质量）"),
            ("未映射检验", "有数值的按项目列进 06_检验" if o.include_unmapped else "不进宽表，只计条数（--include-unmapped 可列出）"),
            ("生成时间", datetime.now().astimezone().isoformat(timespec="seconds")),
            ("读法", "CRF sheet 第 1 行中文名、第 2 行字段名；pandas：pd.read_excel(文件, sheet_name=..., header=1)；各 sheet 用 encounter_id 连接"),
            ("口径提醒", "“全程”最低 / 最高含 t0 之后，不能直接当预测特征；unrecorded / 空 = 没有记录，不等于无；"
                     "分级与候选结局取有效层，候选结局不是标签（A 档）；season_index 不筛（D73）；14_其他数据项是本脚本没有专门处理的编码"),
            ("数据纪律", "本文件逐就诊，含病例ID与真实日期：只留研究者本机，不提交、不发给在线 LLM / API（含 Claude）")]
    tables_x, key_x = analyse(d, sheets, o, Fmt(False), names=True)
    write_workbook(xlsx, d, sheets, tables_x, key_x, info)
    tables_s, key_s = analyse(d, sheets, o, Fmt(o.suppress), names=o.names)
    problems = write_report(md, tables_s, key_s, info, d, check=o.suppress)
    record = {"generated_at": dict(info)["生成时间"], "script": "scripts/crf_wide/build_crf_workbook.py",
              "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(), "how": how,
              "options": {"window": list(o.window), "keep_implausible": o.keep_implausible, "include_unmapped": o.include_unmapped,
                          "age_main": o.age_main, "age_sub": o.age_sub, "sites": o.sites, "suppress": o.suppress, "names": o.names},
              "inputs": [{"name": b.name, "format": b.manifest.get("format"), "version": b.manifest.get("version"),
                          "generated_at": b.manifest.get("generated_at"), "snapshot": (b.manifest.get("snapshot") or {}).get("id"),
                          "encounters_sha256": b.sha} for b in group],
              "rows": {"encounters": len(d.enc), "events": len(d.ev)}, "sheets": [s.name for s in sheets],
              "outputs": {"xlsx": xlsx.name, "md": None if problems else md.name}, "report_check": problems or "通过"}
    rec.write_text(json.dumps(record, ensure_ascii=False, indent=1), encoding="utf-8")
    return Result(tag, xlsx, None if problems else md, rec, len(d.enc), len(d.ev), problems)


def run(root: Path, out_dir: Path | None = None, o: Options | None = None, mode: str = "auto", skip_existing: bool = False) -> list[Result]:
    o = o or Options()
    root = Path(root).resolve()
    bundles = find_bundles(root)
    if not bundles:
        raise FileNotFoundError(f"在 {root} 下没有找到 encounters.parquet（标准导出解开后的目录，或导出的 zip）")
    groups, how = plan_groups(bundles, mode)
    out_dir = Path(out_dir or root.parent / "crf_wide").resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    used: set[str] = set()
    return [build_one(g, how, out_dir, _tag(g, o, used), o, skip_existing) for g in groups]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="flucase 标准导出（std_v3 / std_v2）→ CRF 式宽表（一个 Excel，多 sheet）+ 聚合分析报告")
    ap.add_argument("--input", type=Path, default=Path("data/export"), help="导出目录或 zip；目录下的子目录和 zip 都会找（默认 data/export）")
    ap.add_argument("--out-dir", type=Path, default=None, help="输出目录（默认输入旁边的 crf_wide/）")
    ap.add_argument("--mode", choices=("auto", "merge", "each", "latest"), default="auto",
                    help="几份导出怎么出宽表：auto 互不重叠就合并、有重叠就各出一份（默认）；merge 全部合并（同一就诊保留较新的）；"
                         "each 每份各出一份；latest 只用最新的一份")
    ap.add_argument("--sites", default=None, help="只要这些中心，逗号分隔，如 GD,GAM")
    ap.add_argument("--window", type=float, nargs=2, default=(-24.0, 24.0), metavar=("LO", "HI"), help="基线窗，相对 t0 的小时（默认 -24 24，D18）")
    ap.add_argument("--age-main", type=float, default=60.0, help="主分析人群的年龄下限（默认 60，D70）")
    ap.add_argument("--age-sub", type=float, default=65.0, help="亚组的年龄下限（默认 65，D70）")
    ap.add_argument("--keep-implausible", action="store_true", help="越界值也进基线与极值")
    ap.add_argument("--include-unmapped", action="store_true", help="未映射检验（LABX:）有数值的也按项目列进 06_检验")
    ap.add_argument("--skip-existing", action="store_true", help="同名宽表已经存在就跳过（目录里攒了很多份导出时用）")
    ap.add_argument("--no-suppress", action="store_true", help="报告里也照实计数（报告不再适合外传，也不做自检）")
    ap.add_argument("--no-names", action="store_true", help="报告里不列未映射项目名称与诊断编码")
    a = ap.parse_args(argv)
    if not a.input.exists():
        print(f"找不到输入：{a.input.resolve()}", file=sys.stderr)
        return 1
    o = Options(window=(float(a.window[0]), float(a.window[1])), keep_implausible=a.keep_implausible, include_unmapped=a.include_unmapped,
                age_main=a.age_main, age_sub=a.age_sub, sites=[s.strip() for s in a.sites.split(",") if s.strip()] if a.sites else None,
                suppress=not a.no_suppress, names=not a.no_names)
    try:
        results = run(a.input, a.out_dir, o, a.mode, a.skip_existing)
    except FileNotFoundError as e:
        print(e, file=sys.stderr)
        return 1
    code = 0
    for r in results:
        if r.skipped:
            print(f"[{r.tag}] 跳过：{r.skipped}")
            continue
        print(f"[{r.tag}] 就诊 {r.n_enc}，记录 {r.n_ev}")
        print(f"  宽表：{r.xlsx}")
        if r.problems:
            print(f"  报告没有写出：自检发现{'、'.join(r.problems)}（请把这一行回报，不要贴报告内容）", file=sys.stderr)
            code = 2
        else:
            print(f"  报告：{r.md}" + ("" if o.suppress else "（照实计数，不适合外传）"))
        print(f"  生成记录：{r.record}")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
