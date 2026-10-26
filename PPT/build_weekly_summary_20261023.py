# -*- coding: utf-8 -*-
"""
生成《老年流感转重研究 · 本周总结与周末计划》（2026-10-23）。

输入：PPT/hospital_优化版.pptx（模板，只读，不会被修改；沿用其主题、封面和版式）
      PPT/flucase_病例时间线_虚构.png（flucase 测试教程里的截图，虚构演示数据）
输出：PPT/周总结与周末计划_20261023.pptx

样式与工具函数沿用 PPT/build_hospital_optimized.py：标题栏 + 章节标签 + 青色分隔线 + 关键信息条
+ 浅黄结论框 + 页脚页码；字体微软雅黑（西文、东亚都写）；配色同原稿。模板里的校徽已被删掉，这里也不加。

数字都是聚合量，来源（全部在 BaseModel 与 flucase 仓库内）：
- flucase README“当前状态”（10-17 至 10-23 各行）：功能、性能、测试数、部署、决策 D50–D73；
- flucase docs/GD与GAM对齐及统一导出方案_20261022.md：规则集 v0.2 / v0.3、std_v3、§7 核对清单、§8 待定；
- flucase docs/多中心统一格式与数据集建设_分析与建议_20261022.md：两院差异、站点可区分性 0.709、GD CRF 99.4% 单时间点；
- flucase 分支 medical-timeseries-dataset-standards 的 docs/医学时序数据集的格式与划分_讲解_20261023.html；
- BaseModel docs/实验/实验进度.md（GAM→GD v1.0）、docs/claude_project/02_项目背景知识.md §6.3（六模型实验，服务器聚合结果）。
10-23 晚计划变动：年底出结果暂不考虑，P1 与年底投稿的内容已从本页去掉。
没有读任何真实病例；截图里的病例为虚构。

运行（BaseModel 根目录）：python PPT/build_weekly_summary_20261023.py
依赖：python-pptx、lxml
"""

import uuid
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "hospital_优化版.pptx"
SHOT = HERE / "flucase_病例时间线_虚构.png"
OUT = HERE / "周总结与周末计划_20261023.pptx"

# ---------------------------------------------------------------------------
# 全局样式（与 build_hospital_optimized.py 一致）
# ---------------------------------------------------------------------------
FONT = "微软雅黑"
DECK_TITLE = "老年流感转重研究 · 本周总结与周末计划"

DARK = "1F4E94"     # 深蓝：顶部色条、卡片标题栏
MID = "4075BE"      # 中蓝：标题栏
CYAN = "00A4CB"     # 青色：分隔线、本周新增
LIGHT = "EBF3FB"    # 浅蓝：关键信息条
PALE = "F5F8FC"     # 更浅的蓝灰：行底色
BORDER = "CDDAE8"   # 卡片描边
NAVY = "11376A"     # 深蓝文字
BODY = "273140"     # 正文
MUTED = "646E79"    # 说明文字、页脚
NOTE_BG = "FCEFC6"  # 浅黄：底部结论框
WHITE = "FFFFFF"
RED = "C0392B"      # 重症 / 要你动手
AMBER = "B7791F"    # 待决定
GREY = "8A94A6"     # 外部 / 弱代理 / 可选
MILD = "D9786C"     # 中型
CRIT = "8E1F14"     # 危重型
BEFORE = "C5D0DE"   # 图表：改前

CL = 0.55           # 版心左边界（英寸）
CW = 12.23          # 版心宽度
SW = 13.333         # 幻灯片宽度

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------------------
# 文字工具（沿用 build_hospital_optimized.py）
# ---------------------------------------------------------------------------
def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_typeface(rPr, font=FONT):
    """西文、东亚、复杂文种字体都设为 font（只写 latin 时中文会回退成宋体）。"""
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
        for attr in ("panose", "charset", "pitchFamily"):
            if attr in el.attrib:
                del el.attrib[attr]


def style_run(run, size, color=BODY, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    set_typeface(rPr)


def set_bullet(paragraph, color=MID, indent=0.2):
    """项目符号（悬挂缩进），不在文字里写“•”。"""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def add_break(paragraph, size):
    br = paragraph._p.add_br()
    rPr = etree.SubElement(br, qn("a:rPr"), lang="zh-CN", sz=str(int(round(size * 100))))
    set_typeface(rPr)


def fill_text(tf, paras, size=14, color=BODY, bold=False, align="l", line_spacing=1.08, space_after=0):
    """paras：字符串；或列表，每段为字符串 / [(文字, {选项})] / {"runs":…, "bullet":…, "size":…}。"""
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = spec if isinstance(spec, dict) else {"runs": spec}
        runs = opts.get("runs", "")
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = ALIGN[opts.get("align", align)]
        p.line_spacing = opts.get("line_spacing", line_spacing)
        p.space_before = Pt(opts.get("space_before", 0))
        p.space_after = Pt(opts.get("space_after", space_after))
        p_size = opts.get("size", size)
        p_color = opts.get("color", color)
        p_bold = opts.get("bold", bold)
        for text, ropt in runs:
            r_size = ropt.get("size", p_size)
            for k, part in enumerate(text.split("\n")):
                if k:
                    add_break(p, r_size)
                if not part:
                    continue
                r = p.add_run()
                r.text = part
                style_run(r, r_size, ropt.get("color", p_color), ropt.get("bold", p_bold), ropt.get("italic", False))
        if opts.get("bullet"):
            set_bullet(p, color=opts.get("bullet_color", MID))


def setup_frame(tf, anchor="t", margins=(0.05, 0.03, 0.05, 0.03), wrap=True):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = ANCHOR[anchor]


def add_text(slide, x, y, w, h, paras, size=14, color=BODY, bold=False, align="l", anchor="t",
             margins=(0.05, 0.03, 0.05, 0.03), line_spacing=1.08, space_after=0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    setup_frame(tb.text_frame, anchor, margins)
    fill_text(tb.text_frame, paras, size, color, bold, align, line_spacing, space_after)
    if name:
        tb.name = name
    return tb


def text_width(s, size_pt):
    """粗略估计单行文字宽度（英寸），用来给徽章、标签定宽。"""
    em = 0.0
    for ch in s:
        if ord(ch) >= 0x2E80:
            em += 1.0
        elif ch == " ":
            em += 0.3
        elif ch in "·•":
            em += 0.5
        else:
            em += 0.58
    return em * size_pt / 72


def bullets(items, gap=4, color=MID):
    """把字符串或 [(文字, 选项)] 列表转成带项目符号的段落。"""
    return [{"runs": t, "bullet": True, "space_after": gap, "bullet_color": color} for t in items]


def lead(label, text, label_color=NAVY):
    """“粗体标签 + 正文”的一段。"""
    return [(label, {"bold": True, "color": label_color}), (text, {})]


# ---------------------------------------------------------------------------
# 形状工具（沿用）
# ---------------------------------------------------------------------------
def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, alpha=16, blur=5, dist=1.5):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)),
                          dist=str(int(dist * 12700)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="1F3B66")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE,
             radius=None, shadow=False, dash=None, name=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    else:
        shp.line.fill.background()
    if radius is not None:
        shp.adjustments[0] = radius
    if shadow:
        add_shadow(shp)
    if name:
        shp.name = name
    return shp


def shape_text(shape, paras, size=14, color=BODY, bold=False, align="c", anchor="m",
               margins=(0.04, 0.02, 0.04, 0.02), wrap=True):
    setup_frame(shape.text_frame, anchor, margins, wrap)
    fill_text(shape.text_frame, paras, size, color, bold, align)
    return shape


def add_oval(slide, cx, cy, d, fill=None, line=None, line_w=1.25, dash=None):
    return add_rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, line_w=line_w,
                    shape=MSO_SHAPE.OVAL, dash=dash)


def add_line(slide, x1, y1, x2, y2, color, width=1.0, arrow_end=False, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow_end:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


# ---------------------------------------------------------------------------
# 表格工具（沿用）
# ---------------------------------------------------------------------------
def set_cell_border(cell, color=None, width=0.75, sides="B"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for i, (tag, side) in enumerate((("a:lnL", "L"), ("a:lnR", "R"), ("a:lnT", "T"), ("a:lnB", "B"))):
        ln = etree.Element(qn(tag))
        if color and side in sides:
            ln.set("w", str(int(width * 12700)))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def add_table(slide, x, y, col_w, row_h, rows, size=13, header_size=13.5, header_fill=DARK,
              line_color="DDE6F0", margins=(0.1, 0.05, 0.1, 0.05)):
    nrows, ncols = len(rows), len(col_w)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tbl.vert_banding = False
    for i, wv in enumerate(col_w):
        tbl.columns[i].width = Inches(wv)
    for i, hv in enumerate(row_h):
        tbl.rows[i].height = Inches(hv)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            spec = val if isinstance(val, dict) else {"text": val}
            cell = tbl.cell(r, c)
            is_head = r == 0
            fill = spec.get("fill", header_fill if is_head else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            set_cell_border(cell, None if is_head else line_color, 0.75, "B")
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [Inches(m) for m in margins]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text(tf, spec.get("text", ""),
                      size=spec.get("size", header_size if is_head else size),
                      color=spec.get("color", WHITE if is_head else BODY),
                      bold=spec.get("bold", is_head),
                      align=spec.get("align", "l"))
    return tbl


def merge_rows(tbl, col, r0, r1, paras, fill="DCE7F5"):
    """合并某一列的若干行，写入新文字（同模板第 13 页的“批次”列）。"""
    cell = tbl.cell(r0, col)
    cell.merge(tbl.cell(r1, col))
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    tf = cell.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    tf.paragraphs[0].text = ""
    fill_text(tf, paras, 13, NAVY, align="c")
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


# ---------------------------------------------------------------------------
# 页面框架（沿用模板：顶部色条、标题栏、章节标签、分隔线、关键信息条、结论框、页脚、页码）
# ---------------------------------------------------------------------------
def add_slide_number(slide, x=11.68, y=7.06, w=1.1, h=0.28):
    number = slide.part.package.presentation_part.presentation.slides.index(slide) + 1
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "页码"
    setup_frame(tb.text_frame, "m", (0, 0, 0, 0))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{%s}" % str(uuid.uuid4()).upper(), type="slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="zh-CN", sz="1050")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=MUTED)
    set_typeface(rPr)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = str(number)


def frame(slide, badge, title, key, ask=None, ask_lead=None, notes=None):
    add_rect(slide, 0, 0, SW, 0.10, fill=DARK, name="顶部色条")
    add_rect(slide, 0, 0.24, 10.30, 0.58, fill=MID, name="标题栏")
    bw = text_width(badge, 14) + 0.34
    b = add_rect(slide, 0.35, 0.33, bw, 0.40, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.25, name="章节标签")
    shape_text(b, badge, 14, DARK, bold=True)
    tx = 0.35 + bw + 0.16
    add_text(slide, tx, 0.24, 10.30 - tx - 0.15, 0.58, title, size=22, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0), name="标题")
    add_rect(slide, 0.50, 0.98, 12.28, 0.04, fill=CYAN, name="分隔线")
    add_rect(slide, CL, 1.12, CW, 0.50, fill=LIGHT, name="关键信息条")
    add_text(slide, CL + 0.2, 1.12, CW - 0.4, 0.50, key, size=16, color=NAVY, bold=True, anchor="m",
             name="关键信息")
    if ask:
        add_rect(slide, CL, 6.42, CW, 0.46, fill=NOTE_BG, name="结论框")
        runs = [(ask_lead, {"bold": True}), (ask, {})] if ask_lead else [(ask, {})]
        add_text(slide, CL + 0.2, 6.42, CW - 0.4, 0.46, [runs], size=15, color=NAVY, anchor="m",
                 name="结论")
    add_text(slide, CL, 7.06, 9.0, 0.28, DECK_TITLE, size=10.5, color=MUTED, anchor="m",
             margins=(0, 0, 0, 0), name="页脚")
    add_slide_number(slide)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def card(slide, x, y, w, h, header, header_h=0.44, header_size=16, header_fill=DARK, shadow=True):
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=shadow)
    add_rect(slide, x, y, w, header_h, fill=header_fill)
    add_text(slide, x + 0.15, y, w - 0.3, header_h, header, size=header_size, color=WHITE, bold=True,
             anchor="m", margins=(0, 0, 0, 0))


def stat_tile(slide, x, y, w, h, big, label, big_size=20, label_size=12):
    """白底小卡片：上面大数字，下面一行说明。"""
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, x + 0.16, y + 0.06, w - 0.32, h * 0.52, big, size=big_size, color=DARK, bold=True,
             anchor="b", margins=(0, 0, 0, 0))
    add_text(slide, x + 0.16, y + 0.06 + h * 0.52, w - 0.32, h * 0.44 - 0.1, label, size=label_size,
             color=MUTED, anchor="t", margins=(0, 0.02, 0, 0))


def chip(slide, x, y, text, kind, size=12.5, h=0.36):
    """判据小块：kind = guideline（方案原文）/ proxy（代理）/ weak（弱代理）。返回宽度。

    宽度按字数估计后留足余量（≥、%、/ 等符号在雅黑里比估计的宽），并关掉自动换行，避免一个判据折成两行。
    """
    w = text_width(text, size) * 1.06 + 0.42
    if kind == "guideline":
        s = add_rect(slide, x, y, w, h, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        shape_text(s, text, size, WHITE, bold=True, wrap=False)
    elif kind == "proxy":
        s = add_rect(slide, x, y, w, h, fill=WHITE, line=MID, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.3)
        shape_text(s, text, size, NAVY, bold=True, wrap=False)
    else:
        s = add_rect(slide, x, y, w, h, fill=WHITE, line=GREY, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.3, dash=MSO_LINE_DASH_STYLE.DASH)
        shape_text(s, text, size, MUTED, bold=True, wrap=False)
    return w


def chart_fonts(ch):
    """图表文字补上东亚字体（与模板一致）。"""
    for latin in list(ch._chartSpace.iter(qn("a:latin"))):
        parent = latin.getparent()
        if parent.find(qn("a:ea")) is None:
            ea = etree.Element(qn("a:ea"))
            ea.set("typeface", FONT)
            latin.addnext(ea)


def hide_value_axis(va):
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va._element.get_or_add_delete_().set("val", "1")
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE


# ---------------------------------------------------------------------------
# 演讲者备注
# ---------------------------------------------------------------------------
NOTES = {
    "cover": "封面。本周指 10-17（周六）至 10-23（周五）。内容：本周成果、现状判断、周末两天的安排、下周衔接。"
             "全部数字为聚合量，来自 BaseModel 与 flucase 仓库里的文档；没有读任何真实病例。",
    "summary": "约 1 分钟。先给结论：系统已经跑在数据前面，周末最值钱的事是把真实数据跑一遍、把计数交回来。"
               "10-23 晚计划变动，年底出结果暂不考虑，周末只排医院线和六模型实验的事。四个问题对应后面的页码。",
    "overview": "约 1.5 分钟。左边是 BaseModel 已有的探索：本机的时序基础模型权重、GAM→GD v1.0 的结果（实验进度.md）；"
                "右边是医院数据的变化：GD CRF 几乎没有时序，这周 flucase 读进来的 GD 原始导出是护理 + 监护的分钟级记录。"
                "底部是共用的家底。年底出结果暂不考虑（10-23 计划变动）。",
    "week": "约 1.5 分钟。按天讲：10-17 部署；10-18 拿到 GD 原始导出后写读取器（D50–D58）并做性能改造；"
            "10-19、10-20 试用后的修改和提速；10-21 医生端体征图；10-21 / 10-22 六模型实验本人负责的三个模型跑完；"
            "10-22 规则集 v0.2、v0.3 与 std_v3；10-23 确认 D69–D73 并部署。"
            "三步：读得进、判得齐、交得出。深蓝圆点是当天部署的日子。",
    "sixmodel": "约 1.5 分钟。MC-MED 冻结 benchmark，decomp_60，六项生命体征、0–15 分钟输入窗；每个模型只训练一次，"
                "四套测试集用同一 checkpoint 打分，以 CSN 为单位 paired bootstrap 1,000 次。数字为 5 个 seed 的均值，"
                "来自服务器 week04 结果（docs/claude_project/02_项目背景知识.md §6.3 的聚合表）。"
                "只说“机制与观测策略敏感性存在关联”，不做因果主张。OPL-MT-MNAR 的结果等合作同学。",
    "gd": "约 1.5 分钟。GD 原始导出是一整套 20 个 xlsx（13 类表），按表头认表、按 ENCT_ID / 住院号挂接。"
          "性能数字均为虚构数据（618 次住院、79–88 万条记录）上的本机实测，来自 flucase README 10-18 至 10-21 各行；"
          "真实 GD 约 40 万条。检验只有 JZH 时不入库，要等 GD 补编号。",
    "rules": "约 2 分钟。依据《流行性感冒诊疗方案（2025年版）》p6–7、p12。深蓝实心是方案原文判据，白底蓝框是代理，"
             "虚线是弱代理（只进 primary_img 敏感性口径）。D66、D67 之后，重型主要靠“新开始氧疗”、危重型主要靠“进入 ICU”，"
             "这两个代理在两院的检出机会不同（GD 来自医嘱和科室，GAM 来自病历文字），所以要用 §7 第 11–14 项的真实计数检验。",
    "export": "约 1.5 分钟。流程条里青色是本周新增或升级的环节。std_v3 是 std_v2 的超集；默认 Parquet，可选 CSV、Excel；"
              "facts 默认不导、可勾选；只给管理员，下载写审计和数据使用台账。五份文档：前三份 10-22 起草，"
              "方案 10-22 起草、10-23 补 D69–D73，讲解页 10-23。",
    "numbers": "约 1 分钟。测试数按 README 各行记录：10-16 晚 124（D48、D49），10-18 155（GD 原始导出 143 + 性能改造），"
               "10-19 157，10-20 159，10-22 203（v0.2 170 + v0.3 与 std_v3），10-23 204；服务器上 206 个通过。"
               "部署：10-17、10-20、10-21、10-23。规则集本周新增 v0.2、v0.3，v0、v0.1 保留不动。"
               "截图来自 10-16 的研究者测试教程，虚构演示数据。",
    "status": "约 1.5 分钟。四栏：已完成（但只在虚构数据上）、只有你能做的、等你决定的、等外部的。"
              "要点是第二栏：Claude 不读真实病例，真实数据上的导入、重算、计数只能由你来做。",
    "must": "约 2 分钟。三件事都只有你能做。① 先导一份重导前的基线，否则没法比较规则改动前后；"
            "② 只回报计数，按 §7 的格式；③ §8 第 11 项是 A 档，第 14 项同意即可开工；"
            "六模型实验的三项 B 档要和合作同学一起定，先把选项写清楚发过去。",
    "schedule": "约 1.5 分钟。不赶年底以后，核心约 6 小时：周六 4、周日 2；下午和晚上的事都是可选。"
                "下面左边是可以交给 Claude 并行做的（只用虚构数据），右边是有余力再做的。",
    "checks": "约 1.5 分钟。§7 共 14 项（方案 §7）。对齐报告 alignment_report.json 直接给出第 2、4 项和第 11、13 项的一部分；"
              "第 11 项其余、第 12、14 项需要一个只出计数的小脚本。暂缓的几项依赖阶段 1 还没做的功能或审计脚本。",
    "decide": "约 1.5 分钟。A 档决定归你，这里只给选项。第 11 项影响标签定义；第 14 项是技术做法，同意即可开工。"
              "下面五项来自分析稿 §6 与 10-23 讲解页，不急，下周组会前定。",
    "next": "约 1.5 分钟。关键路径：计数交回 → 定稿口径与读取器 → 医生试用 → 流感高峰（12 月至次年 1 月）前上线前瞻收集。"
            "六模型实验等合作同学的 OPL-MT-MNAR 结果，口径对齐后再做汇总和配对比较。日期是建议，不是已定排期。",
    "closing": "结束页。周末必做三件事见第 11 页，核对清单见第 13 页，待定决定见第 14 页。",
}


# ---------------------------------------------------------------------------
# 各页内容
# ---------------------------------------------------------------------------
def build_summary(slide):
    frame(slide, "要点", "四个问题与一句话回答",
          "核心判断：系统已经跑在数据前面，周末最值钱的是把真实数据跑一遍、把计数交回来。",
          ask_lead="计划调整（10-23）：", ask="年底出结果暂不考虑；周末只排医院线和六模型实验的事。",
          notes=NOTES["summary"])
    items = [
        ("01", "本周做成了什么？",
         "flucase 读得进 GD 原始导出、两院同口径判得齐、std_v3 交得出；六模型实验跑完本人三个模型。", "第 4–9 页"),
        ("02", "现在卡在哪里？",
         "全部只在虚构数据上验证过；v0.3 靠“新开始氧疗”“进 ICU”定级，要用真实计数检验。", "第 10 页"),
        ("03", "周末必须做什么？",
         "服务器重导并按 v0.3 重算；交回 §7 核对计数；定两项决定，和合作同学对齐六模型口径。", "第 11–14 页"),
        ("04", "下周怎么接上？",
         "计数交回后改规则、开阶段 1 其余项；约医生做可用性测试；六模型结果齐了再汇总。", "第 15 页"),
    ]
    h, gap = 0.92, 0.14
    for i, (num, q, a, page) in enumerate(items):
        y = 1.86 + i * (h + gap)
        add_rect(slide, 0.72, y, 11.9, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 radius=0.12, shadow=True)
        nb = add_rect(slide, 0.95, y + 0.16, 0.72, 0.60, fill=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        shape_text(nb, num, 20, WHITE, bold=True)
        add_text(slide, 1.92, y + 0.07, 9.0, 0.42, q, size=18, color=NAVY, bold=True, anchor="m")
        add_text(slide, 1.92, y + 0.49, 9.15, 0.36, a, size=14, color=BODY, anchor="m")
        add_line(slide, 11.18, y + 0.2, 11.18, y + h - 0.2, BORDER, 1)
        add_text(slide, 11.22, y, 1.32, h, page, size=13, color=MID, bold=True, align="c", anchor="m")


def build_overview(slide):
    frame(slide, "全景", "BaseModel 的探索与医院数据",
          "BaseModel 攒下了模型、文献和两院对齐的经验；这周 flucase 把 GD 原始导出读了进来。",
          ask_lead="可以先做：", ask="用只出计数的小脚本看 GD 原始导出每次住院的体征密度，判断时序基础模型能不能用回医院数据。",
          notes=NOTES["overview"])
    y0, ph = 1.82, 3.45
    pw = (CW - 0.25) / 2
    panels = [
        (CL, "BaseModel · 已有的探索", [
            ("16 个", "时序基础模型权重已在本机（Chronos、TimesFM、Moirai、Toto 等）"),
            ("0.733", "GAM 训练 → GD 外部验证，LR 的 AUROC（0.627–0.826）"),
            ("0.709", "只用 15 个输入变量就能区分两院（站点可区分性 AUROC）"),
        ], "09-29 的结论：在稀疏的 GD CRF 上，时序模型赢不了 LR"),
        (CL + pw + 0.25, "医院数据 · 从稀疏到可用", [
            ("99.4%", "GD CRF 的变量序列只有 1 个时间点，撑不起时序模型"),
            ("10 / 38", "可识别的转重事件（GAM 163 次住院 / GD 292 次住院）"),
            ("约 40 万条", "GD 原始导出：护理 + 监护的分钟级记录，flucase 已能读"),
        ], "这周起，GD 第一次有了密集的时序数据（待真实计数确认）"),
    ]
    for px, head, rows, status in panels:
        card(slide, px, y0, pw, ph, head, header_h=0.48, header_size=15)
        for i, (big, desc) in enumerate(rows):
            ry = y0 + 0.62 + i * 0.78
            add_text(slide, px + 0.18, ry, 2.3, 0.66, big, size=18, color=DARK, bold=True, anchor="m",
                     margins=(0, 0, 0, 0))
            add_text(slide, px + 2.55, ry, pw - 2.7, 0.66, desc, size=12.5, color=BODY, anchor="m",
                     margins=(0, 0, 0, 0), line_spacing=1.05)
            if i < 2:
                add_line(slide, px + 0.18, ry + 0.72, px + pw - 0.18, ry + 0.72, "E3EAF3", 0.75)
        add_rect(slide, px + 0.15, y0 + ph - 0.5, pw - 0.3, 0.38, fill=LIGHT)
        add_text(slide, px + 0.3, y0 + ph - 0.5, pw - 0.6, 0.38, status, size=12, color=NAVY, bold=True,
                 anchor="m", margins=(0, 0, 0, 0))
    tiles = [("24 篇", "论文精读笔记"), ("3 份", "调研报告 + 本周数据集讲解页"),
             ("5 个模型", "六模型实验已出结果（本人 3 个）"), ("D1–D73", "决策记录（本周新增 24 条）")]
    tw = (CW - 3 * 0.2) / 4
    for i, (big, label) in enumerate(tiles):
        stat_tile(slide, CL + i * (tw + 0.2), 5.42, tw, 0.85, big, label, big_size=19, label_size=11.5)


def build_week(slide):
    frame(slide, "本周", "七天三步：读得进、判得齐、交得出",
          "10-17 至 10-23：从读入 GD 原始导出，到两院按同一口径分级，再到统一格式导出。",
          ask_lead="节奏：", ask="4 次部署到专用服务器；每一步先在虚构数据上测试，你确认后再合并。",
          notes=NOTES["week"])
    ly = 2.55
    add_line(slide, 0.85, ly, 12.45, ly, DARK, 2.5, arrow_end=True)
    days = [
        ("10-17 周六", "部署到专用服务器\n124 个测试通过", True),
        ("10-18 周日", "读 GD 原始导出\nD50–D58；性能改造", False),
        ("10-19 周一", "试用后三处修改\n放弃导入、翻页", False),
        ("10-20 周二", "压缩包读取提速\n部署", True),
        ("10-21 周三", "医生端体征图\n修预览显示；部署", True),
        ("10-22 周四", "规则集 v0.2→v0.3\nstd_v3、导出页", False),
        ("10-23 周五", "确认 D69–D73\n部署，206 个测试", True),
    ]
    x0, step = 1.35, (11.95 - 1.35) / 6
    for i, (date, label, deployed) in enumerate(days):
        mx = x0 + i * step
        add_oval(slide, mx, ly, 0.26, fill=DARK if deployed else CYAN, line=WHITE, line_w=2)
        add_text(slide, mx - 0.86, 1.86, 1.72, 0.5, date, size=13.5, color=NAVY, bold=True, align="c", anchor="b")
        add_text(slide, mx - 0.86, 2.76, 1.72, 0.62, label, size=12, color=MUTED, align="c")
    add_oval(slide, 0.72, 3.56, 0.16, fill=DARK, line=WHITE, line_w=1.5)
    add_text(slide, 0.86, 3.42, 2.2, 0.28, "当天部署到专用服务器", size=10.5, color=MUTED, anchor="m",
             margins=(0.04, 0, 0, 0))
    add_oval(slide, 3.0, 3.56, 0.16, fill=CYAN, line=WHITE, line_w=1.5)
    add_text(slide, 3.14, 3.42, 1.6, 0.28, "开发与试用", size=10.5, color=MUTED, anchor="m",
             margins=(0.04, 0, 0, 0))
    # 同一周的另一条线：六模型实验（10-21 / 10-22，服务器上跑）
    sx0, sx1 = x0 + 4 * step - 0.86, x0 + 5 * step + 0.86
    trk = add_rect(slide, sx0, 3.41, sx1 - sx0, 0.3, fill=PALE, line=MID, line_w=1,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    shape_text(trk, "六模型实验：本人 3 个模型 × 5 个 seed 跑完", 11, NAVY, bold=True, wrap=False)

    cards = [
        ("① 读得进：GD 原始导出", ["一整套 20 个 xlsx、13 类表，按表头认表、自动挂接",
                              "文件夹、zip、rar、7z 都能整包上传",
                              "约 40 万条的真实数据量级，读取放在子进程里跑"]),
        ("② 判得齐：两院同一口径", ["两院共用一套文字规则、ICU 科室规则、诊断名称规则",
                              "规则集 v0.3：每条判据标“原文 / 代理 / 弱代理”",
                              "人群、主口径、多次就诊的规则已定（D69–D73）"]),
        ("③ 交得出：std_v3 统一导出", ["两院同表、同列、同编码、同取值",
                                "“数据导出”页一键下载：Parquet / CSV / Excel",
                                "附对齐报告、字典、manifest，同一快照可复现"]),
    ]
    cw, gap = (CW - 2 * 0.25) / 3, 0.25
    for i, (head, lines) in enumerate(cards):
        x = CL + i * (cw + gap)
        card(slide, x, 3.82, cw, 2.42, head, header_h=0.46, header_size=15)
        add_text(slide, x + 0.16, 4.38, cw - 0.32, 1.8, bullets(lines, gap=6), size=13, color=BODY,
                 margins=(0, 0, 0, 0))


def build_gd(slide):
    frame(slide, "成果一", "GD 原始导出：读得进，也读得动",
          "GD 病历系统的原始导出整套读入；数据量上来以后，页面和内存都压下来了。",
          ask_lead="口径说明：", ask="性能数字是虚构数据（618 次住院、79–88 万条记录）上的本机实测；真实 GD 约 40 万条。",
          notes=NOTES["gd"])
    lw = 5.5
    card(slide, CL, 1.82, lw, 4.42, "读法：一整套文件一起读", header_h=0.46, header_size=15)
    add_text(slide, CL + 0.18, 2.4, lw - 0.36, 3.8, bullets([
        lead("认表：", "按表头认出 13 类表——入院记录、医嘱、生命体征、检验、检查、病程等"),
        lead("挂接：", "按 ENCT_ID / ENCT_UUID / 住院号挂成一次次住院"),
        lead("上传：", "整个文件夹，或 zip、rar、7z 压缩包（D57）"),
        lead("清洗：", "体温漏小数点自动 ÷10，原值保留（D55）；SpO2 的吸氧状态按医嘱推断"),
        lead("诊断：", "编码与名称拆开读；名称缺失时，查院内 ICD 字典（3.7 万条）补上"),
        lead("缺口：", "只有 JZH 的检验先不入库，等 GD 补编号"),
    ], gap=7), size=13, color=BODY, margins=(0, 0, 0, 0))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    add_rect(slide, rx, 1.82, rw, 2.98, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, rx + 0.2, 1.9, 4.2, 0.38, "响应时间（秒）：改前 → 改后", size=15, color=NAVY, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    # 自绘图例（按“改前、改后”的阅读顺序）
    lx = rx + rw - 2.05
    add_rect(slide, lx, 2.02, 0.18, 0.14, fill=BEFORE)
    add_text(slide, lx + 0.22, 1.94, 0.7, 0.3, "改前", size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))
    add_rect(slide, lx + 0.95, 2.02, 0.18, 0.14, fill=DARK)
    add_text(slide, lx + 1.17, 1.94, 0.7, 0.3, "改后", size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))

    cd = CategoryChartData()
    # 横向条形图自下而上排类别：倒序写入，让第一项显示在最上面；
    # 同一类别内第一个系列画在下方，所以先写“改后”、再写“改前”，读起来是“改前在上、改后在下”
    cats = ["病例列表首次打开", "导入预览抽查一次就诊", "医生病例页体征图（最密一例）", "病例时间线（最密一例）"]
    before = [8.9, 4.4, 3.7, 2.6]
    after = [0.1, 1.7, 1.4, 0.4]
    cd.categories = list(reversed(cats))
    cd.add_series("改后", list(reversed(after)))
    cd.add_series("改前", list(reversed(before)))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(rx + 0.1), Inches(2.3), Inches(rw - 0.25),
                                Inches(2.45), cd)
    gf.name = "响应时间条形图"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(12)
    ch.font.name = FONT
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.overlap = -8
    plot.vary_by_categories = False
    for ser, col in zip(plot.series, [DARK, BEFORE]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)
        ser.format.line.fill.background()
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format = '0.0" 秒"'
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(11.5)
        dl.font.bold = True
        dl.font.color.rgb = rgb(NAVY if col == DARK else MUTED)
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = 10.8
    hide_value_axis(va)
    ca = ch.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = rgb(BODY)
    chart_fonts(ch)

    tiles = [("1.4 → 0.3 GB", "服务进程常驻内存"), ("1.13 GB", "读 87.6 万条时的峰值"),
             ("174 → 148 秒", "整套读取一次的用时")]
    tw = (rw - 2 * 0.18) / 3
    for i, (big, label) in enumerate(tiles):
        stat_tile(slide, rx + i * (tw + 0.18), 4.95, tw, 1.29, big, label, big_size=18, label_size=11.5)


def build_rules(slide):
    frame(slide, "成果二", "两院同一口径的重症判定：规则集 v0.3",
          "依据《流行性感冒诊疗方案（2025年版）》，每条判据都标明来源：方案原文、代理或弱代理。",
          ask_lead="最大的风险：", ask="重型主要靠“新开始氧疗”，危重型主要靠“进 ICU”和机械通气；两院检出机会不同，需真实计数检验。",
          notes=NOTES["rules"])
    # 图例
    lx = CL
    for text, kind in (("方案原文", "guideline"), ("代理", "proxy"), ("弱代理（只进敏感性分析）", "weak")):
        w = chip(slide, lx, 1.8, text, kind, size=11, h=0.3)
        lx += w + 0.15
    add_text(slide, lx + 0.05, 1.8, 3.2, 0.3, "括号里是对应的研究者决定", size=10.5, color=MUTED, anchor="m",
             margins=(0, 0, 0, 0))

    area_w = 8.3
    x_chip0 = CL + 1.45
    rows = [
        ("中型", MILD, 2.24, [[("影像或诊断提示肺炎", "guideline"), ("气促 / 呼吸困难（不用咳嗽）", "guideline")]]),
        ("重型", RED, 2.94, [[("RR ≥ 30 次/分", "guideline"), ("吸空气 SpO2 ≤ 93%", "guideline"),
                              ("PaO2/FiO2 ≤ 300", "guideline")],
                             [("新开始普通氧疗（D67）", "proxy"), ("高流量氧疗（D60）", "proxy"),
                              ("呼吸衰竭诊断", "proxy")]]),
        ("危重型", CRIT, 4.04, [[("有创或无创机械通气（D61）", "guideline"), ("休克", "guideline"),
                                ("ICU + 器官衰竭（D62）", "guideline")],
                               [("进入 ICU（D66）", "proxy"), ("ECMO", "proxy"), ("升压药", "proxy")]]),
        ("敏感性", GREY, 5.14, [[("24–48 小时影像进展 > 50% 的弱代理（D63）", "weak")]]),
    ]
    for name, col, y, lines in rows:
        h = 0.2 + len(lines) * 0.36 + (len(lines) - 1) * 0.12
        add_rect(slide, CL, y, area_w, h, fill=PALE)
        blk = add_rect(slide, CL, y, 1.25, h, fill=col, shape=MSO_SHAPE.RECTANGLE)
        shape_text(blk, name, 16, WHITE, bold=True)
        for li, line in enumerate(lines):
            cx = x_chip0
            cy = y + 0.1 + li * 0.48
            for text, kind in line:
                cx += chip(slide, cx, cy, text, kind) + 0.12
    add_text(slide, CL, 5.78, area_w, 0.5,
             "ICU + 器官衰竭的“器官”含 ARDS 与急性肝衰竭（D62）；影像进展只进 primary_img 口径，不进分级本身（D63）。",
             size=10.5, color=MUTED, margins=(0, 0, 0, 0))

    rx, rw = CL + area_w + 0.25, CW - area_w - 0.25
    card(slide, rx, 1.8, rw, 4.44, "口径（D69–D73 已定）", header_h=0.46, header_size=15)
    add_text(slide, rx + 0.16, 2.36, rw - 0.32, 3.84, bullets([
        lead("主口径：", "primary = 治疗类 + 诊断类（D69）"),
        lead("敏感性：", "生理类、只用方案原文、加影像进展、O3 的 7 天窗"),
        lead("人群：", "确诊且 ≥60 岁为主分析，≥65 岁亚组（D70）"),
        lead("“新开始”：", "长期家庭氧疗、家庭无创者不算（D71）"),
        lead("术后转入 ICU：", "照算危重型，单独计数（D72）"),
        lead("多次就诊：", "每人每流感季取第一次（D73，ISO 第 40 周起）"),
    ], gap=6), size=12.5, color=BODY, margins=(0, 0, 0, 0))


def build_export(slide):
    frame(slide, "成果三", "std_v3 统一导出，配五份多中心文档",
          "不要求医院按规范导出：flucase 适配各院已有导出，两院导出同一标准格式（10-22 决定）。",
          ask_lead="还没做：", ask="来源类型与检验报告时间两列、mappings 与取值表、按就诊分批写（看真实峰值内存再定）。",
          notes=NOTES["export"])
    steps = [("原始导出", False), ("读取器", True), ("病例库", False), ("分级 v0.3", True), ("复核", False),
             ("冻结快照", False), ("std_v3", True), ("研究者下载", True)]
    n = len(steps)
    sw = (CW + (n - 1) * 0.08) / n
    alt = 0
    for i, (s, new) in enumerate(steps):
        x = CL + i * (sw - 0.08)
        shp_type = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        if new:
            fill = CYAN
        else:
            fill = DARK if alt % 2 == 0 else MID
            alt += 1
        blk = add_rect(slide, x, 1.8, sw, 0.6, fill=fill, shape=shp_type)
        blk.adjustments[0] = 0.28
        shape_text(blk, s, 12.5, WHITE, bold=True, margins=(0.08, 0, 0.04, 0))
    add_text(slide, CL, 2.44, CW, 0.28, "青色为本周新增或升级的环节：GD 原始导出读取器与统一文字规则、规则集 v0.3、std_v3 与“数据导出”页。",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))

    lw = 6.0
    add_text(slide, CL, 2.78, lw, 0.34, "std_v3 导出包里有什么", size=14.5, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    code = lambda t: {"text": t, "bold": True, "color": NAVY, "size": 12}
    rows = [
        ["文件", "一行是什么"],
        [code("encounters"), "一次就诊：中心、ICU 起止、流感季索引"],
        [code("events"), "一条记录：规则号、时间含义、吸氧状态"],
        [code("vitals · labs · treatments\ndiagnoses · flu_tests"), "按领域拆好的视图（默认导出）"],
        [code("severity_timeline\noutcomes"), "分级时间线；各口径候选结局"],
        [code("facts"), "自动层与证据（可勾选，默认不导）"],
        [code("alignment_report.json"), "按中心的覆盖率、来源、精度、一致率"],
    ]
    add_table(slide, CL, 3.14, [2.55, 3.45], [0.36, 0.38, 0.38, 0.54, 0.54, 0.38, 0.38], rows, size=12,
              header_size=12.5, margins=(0.1, 0.03, 0.08, 0.03))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    add_text(slide, rx, 2.78, rw, 0.34, "本周新写的五份文档", size=14.5, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    docs = [
        ("多中心统一格式与数据集建设：分析与建议", "“统一格式”其实是三件事：结构统一、语义统一、数据集治理"),
        ("取数规范 v1", "两院表结构的参考；按研究者决定，不发给医院"),
        ("统一数据规范 v1（草案）", "以后“共享版”（去标识）的发布层，兼容 MEDS"),
        ("GD 与 GAM 对齐及统一导出方案", "D59–D73 的依据；§7 真实数据核对清单"),
        ("医学时序数据集的格式与划分（讲解页）", "MIMIC-IV、eICU、HiRID 等怎么定格式、给切分"),
    ]
    dh, dg = 0.56, 0.06
    for i, (t, sub) in enumerate(docs):
        y = 3.14 + i * (dh + dg)
        add_rect(slide, rx, y, rw, dh, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1,
                 shadow=True)
        nb = add_rect(slide, rx + 0.12, y + 0.1, 0.36, 0.36, fill=DARK if i < 4 else CYAN,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        shape_text(nb, str(i + 1), 13, WHITE, bold=True)
        add_text(slide, rx + 0.62, y + 0.03, rw - 0.75, 0.27, t, size=13, color=NAVY, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, rx + 0.62, y + 0.29, rw - 0.75, 0.24, sub, size=11, color=MUTED, anchor="m",
                 margins=(0, 0, 0, 0))


def build_sixmodel(slide):
    frame(slide, "成果四", "六模型观测删减敏感性实验（与合作同学分工）",
          "随机删减观测时，五个模型的 AUPRC 都在下降，删得越多降得越多；群体指标不变时，个体预测也会翻转。",
          ask_lead="待定（B 档）：", ask="类别加权口径怎么标；Raindrop seed 2028 要不要重跑；配对比较前是否先登记。",
          notes=NOTES["sixmodel"])
    lw = 6.3
    add_rect(slide, CL, 1.82, lw, 4.42, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.9, lw - 0.4, 0.36, "ΔAUPRC（相对原始测试集，5 个 seed 均值）", size=15, color=NAVY,
             bold=True, anchor="m", margins=(0, 0, 0, 0))
    add_text(slide, CL + 0.2, 2.26, lw - 0.4, 0.28, "MC-MED · decomp_60 · 同一 checkpoint 打四套测试集",
             size=10.5, color=MUTED, anchor="m", margins=(0, 0, 0, 0))
    series = [("drop20", "9DBBE3"), ("drop40", DARK), ("continuous5min", CYAN)]
    lx = CL + 0.2
    for name, col in series:
        add_rect(slide, lx, 2.66, 0.18, 0.14, fill=col)
        add_text(slide, lx + 0.23, 2.58, 1.6, 0.3, name, size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))
        lx += 0.23 + text_width(name, 11.5) + 0.35
    cd = CategoryChartData()
    cd.categories = ["FF-LSTM", "GRU-D", "BRITS", "mTAND", "Raindrop"]
    cd.add_series("drop20", (-0.006, -0.020, -0.016, -0.010, -0.008))
    cd.add_series("drop40", (-0.016, -0.036, -0.033, -0.024, -0.024))
    cd.add_series("continuous5min", (-0.005, -0.009, -0.011, 0.004, -0.001))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(CL + 0.1), Inches(2.9), Inches(lw - 0.2),
                                Inches(2.8), cd)
    gf.name = "六模型 ΔAUPRC 柱状图"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(12)
    ch.font.name = FONT
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.overlap = -5
    plot.vary_by_categories = False
    for ser, (name, col) in zip(plot.series, series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)
        ser.format.line.fill.background()
        ser.invert_if_negative = False
        if name == "drop40":
            dl = ser.data_labels
            dl.show_value = True
            dl.number_format = '0.000'
            dl.number_format_is_linked = False
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.size = Pt(10)
            dl.font.bold = True
            dl.font.color.rgb = rgb(NAVY)
    va = ch.value_axis
    va.minimum_scale = -0.045
    va.maximum_scale = 0.01
    va.major_unit = 0.01
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E3EAF3")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    va.tick_labels.number_format = '0.00'
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(10.5)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    ca = ch.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = rgb(BODY)
    chart_fonts(ch)
    add_text(slide, CL + 0.2, 5.74, lw - 0.4, 0.46,
             "本人负责 BRITS、mTAND、Raindrop；合作同学负责 FF-LSTM、GRU-D，OPL-MT-MNAR 结果待出。",
             size=10.5, color=MUTED, margins=(0, 0, 0, 0))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    mine = lambda t: {"text": t, "bold": True, "color": DARK}
    other = lambda t: {"text": t, "bold": True, "color": NAVY}
    c = lambda t: {"text": t, "align": "c"}
    rows = [
        ["模型", {"text": "原始 AUPRC", "align": "c"}, {"text": "原始 AUROC", "align": "c"},
         {"text": "drop40 翻转率", "align": "c"}],
        [other("FF-LSTM"), c("0.275"), c("0.711"), c("6.4%")],
        [other("GRU-D"), c("0.298"), c("0.728"), c("7.3%")],
        [mine("BRITS"), c("0.314"), c("0.734"), c("9.2%")],
        [mine("mTAND"), c("0.273"), c("0.692"), c("6.2%")],
        [mine("Raindrop"), c("0.269"), c("0.705"), c("9.7%")],
    ]
    add_table(slide, rx, 1.82, [1.4, 1.3, 1.3, rw - 4.0], [0.4, 0.36, 0.36, 0.36, 0.36, 0.36], rows, size=12.5,
              header_size=12.5, margins=(0.1, 0.03, 0.08, 0.03))
    card(slide, rx, 4.34, rw, 1.9, "读结果要注意", header_h=0.42, header_size=14, header_fill=AMBER)
    add_text(slide, rx + 0.16, 4.84, rw - 0.32, 1.36, bullets([
        "加权 BCE（pos_weight ≈ 7.08）把概率整体推高，Brier、ECE 不能直接解读",
        "训练 seed 带来的波动大于单个 run 的 bootstrap CI",
        "BRITS 是离线双向模型，不能说成在线预测；只说关联、不说因果",
    ], gap=4, color=AMBER), size=12, color=BODY, margins=(0, 0, 0, 0))


def build_numbers(slide):
    frame(slide, "数据", "本周的数字：测试、部署、决策",
          "工程上是稳的：后端测试从 124 个增加到 204 个，服务器上 206 个全部通过。",
          ask_lead="工程纪律：", ask="每次合并前在虚构数据上跑全部测试；部署后在服务器上再跑一遍，并做正式模式自检。",
          notes=NOTES["numbers"])
    lw = 6.0
    add_rect(slide, CL, 1.82, lw, 2.92, fill=WHITE, line=BORDER, shadow=True)
    add_text(slide, CL + 0.2, 1.9, lw - 0.4, 0.38, "后端测试数（虚构数据，全部通过）", size=15, color=NAVY,
             bold=True, anchor="m", margins=(0, 0, 0, 0))
    cd = CategoryChartData()
    cd.categories = ["10-16", "10-18", "10-19", "10-20", "10-22", "10-23"]
    cd.add_series("后端测试数", (124, 155, 157, 159, 203, 204))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(CL + 0.1), Inches(2.3), Inches(lw - 0.2),
                                Inches(2.38), cd)
    gf.name = "测试数折线图"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(12)
    ch.font.name = FONT
    plot = ch.plots[0]
    ser = plot.series[0]
    ser.smooth = False
    ser.format.line.color.rgb = rgb(DARK)
    ser.format.line.width = Pt(2.5)
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 8
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(DARK)
    ser.marker.format.line.color.rgb = rgb(WHITE)
    dl = ser.data_labels
    dl.show_value = True
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(12)
    dl.font.bold = True
    dl.font.color.rgb = rgb(NAVY)
    va = ch.value_axis
    va.minimum_scale = 100
    va.maximum_scale = 230
    va.major_unit = 50
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E3EAF3")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    ca = ch.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(BORDER)
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = rgb(BODY)
    chart_fonts(ch)

    tiles = [("4 次", "部署到专用服务器", "10-17、20、21、23"), ("24 条", "研究者决策", "D50–D73"),
             ("5 份", "新文档", "10-22 四份，10-23 一份"), ("2 版", "新规则集", "v0.2、v0.3（v0.1 保留）")]
    tw, th = (lw - 0.18) / 2, 0.66
    for i, (big, label, sub) in enumerate(tiles):
        x = CL + (i % 2) * (tw + 0.18)
        y = 4.88 + (i // 2) * (th + 0.12)
        add_rect(slide, x, y, tw, th, fill=WHITE, line=BORDER, shadow=True)
        add_text(slide, x + 0.14, y, 0.9, th, big, size=18, color=DARK, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 1.02, y, tw - 1.1, th,
                 [{"runs": label, "size": 12, "color": BODY, "bold": True},
                  {"runs": sub, "size": 10.5, "color": MUTED}],
                 anchor="m", margins=(0, 0, 0, 0))

    rx, rw = CL + lw + 0.25, CW - lw - 0.25
    ih = rw / 1.6
    add_rect(slide, rx - 0.04, 1.78, rw + 0.08, ih + 0.08, fill=WHITE, line=BORDER, shadow=True)
    pic = slide.shapes.add_picture(str(SHOT), Inches(rx), Inches(1.82), Inches(rw), Inches(ih))
    pic.name = "flucase 截图（虚构演示数据）"
    add_text(slide, rx, 1.82 + ih + 0.12, rw, 0.62,
             [[("flucase · 候选分级 · 病例时间线：", {"bold": True, "color": NAVY}),
               ("体温曲线、发病 48 小时底色、O3 的 7 / 14 天观察窗。虚构演示数据，10-16 版界面。", {})]],
             size=11.5, color=MUTED, margins=(0, 0, 0, 0))


def build_status(slide):
    frame(slide, "现状", "现状：结构已同构，卡在真实数据验证",
          "系统已经跑在数据前面：本周所有新功能都只在虚构数据上验证过。",
          ask_lead="所以：", ask="周末不加功能，先把真实数据跑一遍、把计数交回来——这一步只有你能做（Claude 不读真实病例）。",
          notes=NOTES["status"])
    cols = [
        ("已完成（虚构数据上）", DARK, [
            "GD 原始导出读取",
            "两院共用文字规则",
            "ICU 科室、诊断名称规则",
            "规则集 v0.2、v0.3",
            "std_v3 与导出页",
            "服务器 206 个测试通过",
        ]),
        ("只有你能做", RED, [
            "重新导入三类原始导出",
            "版本对比，切到 v0.3 重算",
            "前后各导一次 std_v3",
            "§7 第 2、4、8、11–14 项",
            "约医生做可用性测试",
        ]),
        ("等你决定", AMBER, [
            "§8 第 11 项：再住院的事件进不进标签（A 档）",
            "§8 第 14 项：阶段 1 其余项开不开工",
            "六模型三项 B 档（和合作同学）",
            "多中心五项（不急）",
        ]),
        ("等外部", GREY, [
            "GD 补检验编号（JZH 对照）",
            "合作同学的 OPL-MT-MNAR 结果",
            "临床负责人确认口径（对外沟通暂缓）",
            "复核医生人选",
        ]),
    ]
    cw, gap = (CW - 3 * 0.2) / 4, 0.2
    for i, (head, col, items) in enumerate(cols):
        x = CL + i * (cw + gap)
        card(slide, x, 1.82, cw, 2.98, head, header_h=0.5, header_size=15, header_fill=col)
        add_text(slide, x + 0.16, 2.44, cw - 0.32, 2.3, bullets(items, gap=6, color=col), size=12.5,
                 color=BODY, margins=(0, 0, 0, 0))

    # 底部：flucase 上线顺序（D45）的进度
    add_text(slide, CL, 4.94, 4.0, 0.3, "flucase 上线顺序（D45）", size=13.5, color=NAVY, bold=True, anchor="m",
             margins=(0, 0, 0, 0))
    steps = [("① 部署前开发准备", "完成", DARK, WHITE),
             ("② 研究者试用", "进行中：GD 原始导出已试导", CYAN, WHITE),
             ("③ 部署到专用服务器", "完成：10-17 起，已更新 3 次", DARK, WHITE),
             ("④ 医生测试", "计划中：可用性测试与试用", WHITE, NAVY)]
    n = len(steps)
    sw = (CW + (n - 1) * 0.08) / n
    for i, (name, state, fill, tc) in enumerate(steps):
        x = CL + i * (sw - 0.08)
        shp_type = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        blk = add_rect(slide, x, 5.3, sw, 0.82, fill=fill, line=None if fill != WHITE else MID, line_w=1.25,
                       shape=shp_type)
        blk.adjustments[0] = 0.22
        shape_text(blk, [{"runs": name, "bold": True, "size": 13.5},
                         {"runs": state, "size": 11, "bold": False}], 13, tc, margins=(0.22, 0, 0.12, 0))


def build_must(slide):
    frame(slide, "周末必做", "三件只有你能做的事",
          "做完这三件，下周一 Claude 就能接着改规则、开阶段 1 的其余项；六模型的汇总也不用再等口径。",
          ask_lead="数据纪律：", ask="交回的只有计数、名称清单和问题代码；不发病例内容、编号和截图，也不发给任何在线大模型。",
          notes=NOTES["must"])
    tasks = [
        ("① 服务器重导与重算", "约 1.5 小时 · 周六上午", [
            "先冻结并导出一份 std_v3（重导前基线）",
            "重新导入 GD 原始导出、GAM 住院和门诊，新规则才生效",
            "规则集页做 v0.2 / v0.3 版本对比，再切到 v0.3 重算",
            "再导出一份 std_v3；就诊数应不变，只有“更新”",
        ], "重导前、后两份 std_v3 导出包"),
        ("② 核对计数，只回报数字", "约 2.5 小时 · 周六下午", [
            "从两份对齐报告抄第 2、4 项，以及第 11、13 项的部分计数",
            "手工补第 8 项名称清单、第 13 项各科室住院数、“必要时”氧疗行数",
            "第 11 项其余、第 12、14 项用只出计数的小脚本（可让 Claude 写，你在本机跑）",
        ], "一张中心 × 项目的计数表"),
        ("③ 定下来、约好", "约 2 小时 · 周日上午", [
            "§8 第 11 项（A 档）、第 14 项（同意即可开工）",
            "给合作同学一页说明：六模型三项 B 档各列选项，等 OPL-MT-MNAR 结果出来就能汇总",
            "约 1–2 位医生下周做可用性测试（虚构样例，5 项任务）",
        ], "两项决定、给合作同学的说明、医生测试时间"),
    ]
    cw, gap = (CW - 2 * 0.25) / 3, 0.25
    for i, (head, when, items, out) in enumerate(tasks):
        x = CL + i * (cw + gap)
        y, h = 1.82, 4.42
        card(slide, x, y, cw, h, head, header_h=0.52, header_size=16.5)
        pill = add_rect(slide, x + 0.16, y + 0.66, cw - 0.32, 0.36, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        radius=0.5)
        shape_text(pill, when, 12.5, NAVY, bold=True)
        add_text(slide, x + 0.16, y + 1.16, cw - 0.32, 2.4, bullets(items, gap=7), size=12.5, color=BODY,
                 margins=(0, 0, 0, 0))
        add_rect(slide, x + 0.16, y + h - 0.72, cw - 0.32, 0.56, fill=PALE, line=BORDER)
        add_text(slide, x + 0.28, y + h - 0.72, cw - 0.56, 0.56, [lead("产出：", out)], size=12.5, color=BODY,
                 anchor="m", margins=(0, 0, 0, 0))


def build_schedule(slide):
    frame(slide, "周末安排", "周六跑真实数据，周日定事、约人",
          "不赶年底以后，两天合计约 6 小时有效工作就够了；下午和晚上的都是可选，其余时间休息。",
          ask_lead="不建议：", ask="计数回来之前给 flucase 加新功能；为赶进度熬夜；把任何病例内容发给在线大模型。",
          notes=NOTES["schedule"])
    x0, x1 = 1.85, 12.6
    t0, t1 = 9.0, 21.0
    px = lambda t: x0 + (t - t0) / (t1 - t0) * (x1 - x0)
    for hh in range(9, 22, 2):
        add_text(slide, px(hh) - 0.4, 1.8, 0.8, 0.26, f"{hh:02d}:00", size=10.5, color=MUTED, align="c",
                 anchor="m", margins=(0, 0, 0, 0))
        add_line(slide, px(hh), 2.08, px(hh), 3.92, "DDE6F0", 0.75, dash=MSO_LINE_DASH_STYLE.DASH)
    days = [
        ("周六\n10-24", 2.12, [
            (10.0, 11.5, "重导与重算", MID, False),
            (14.0, 16.5, "核对计数、填回报表", MID, False),
            (19.5, 21.0, "文献（可选）", None, True),
        ]),
        ("周日\n10-25", 3.06, [
            (10.0, 11.0, "两项决定\n约医生", CYAN, False),
            (11.0, 12.0, "六模型\n口径说明", DARK, False),
            (14.0, 15.5, "GD 观测密度\n计数（可选）", None, True),
            (19.5, 21.0, "自学（可选）", None, True),
        ]),
    ]
    rh = 0.8
    for label, y, blocks in days:
        add_rect(slide, CL, y, 1.18, rh, fill=LIGHT)
        add_text(slide, CL, y, 1.18, rh, label, size=13.5, color=NAVY, bold=True, align="c", anchor="m")
        add_rect(slide, x0, y, x1 - x0, rh, fill=PALE)
        for s, e, text, col, optional in blocks:
            bx, bw = px(s) + 0.02, px(e) - px(s) - 0.04
            if optional:
                b = add_rect(slide, bx, y + 0.1, bw, rh - 0.2, fill=WHITE, line=GREY, line_w=1.25,
                             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, dash=MSO_LINE_DASH_STYLE.DASH)
                shape_text(b, text, 12, MUTED, bold=True)
            else:
                b = add_rect(slide, bx, y + 0.1, bw, rh - 0.2, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                             radius=0.12)
                shape_text(b, text, 12 if "\n" not in text else 11, WHITE, bold=True)
    # 图例与合计
    lx = x0
    for text, col, dashed in (("flucase 真实数据", MID, False), ("决定与约人", CYAN, False), ("六模型实验", DARK, False),
                              ("可选", None, True)):
        if dashed:
            add_rect(slide, lx, 4.04, 0.3, 0.18, fill=WHITE, line=GREY, line_w=1, dash=MSO_LINE_DASH_STYLE.DASH)
        else:
            add_rect(slide, lx, 4.04, 0.3, 0.18, fill=col)
        add_text(slide, lx + 0.36, 3.98, 1.9, 0.3, text, size=11.5, color=BODY, anchor="m", margins=(0, 0, 0, 0))
        lx += 0.42 + text_width(text, 11.5) + 0.35
    add_text(slide, x1 - 4.6, 3.98, 4.6, 0.3, "核心约 6 小时：周六 4 · 周日 2（不含可选）", size=11.5,
             color=NAVY, bold=True, align="r", anchor="m", margins=(0, 0, 0, 0))

    bw = (CW - 0.25) / 2
    card(slide, CL, 4.46, bw, 1.8, "可以交给 Claude 并行做（只用虚构数据）", header_h=0.42, header_size=14)
    add_text(slide, CL + 0.16, 4.96, bw - 0.32, 1.26, bullets([
        "只出计数的核对脚本（§7 第 11、12、14 项），你在本机跑",
        "GD 观测密度的计数脚本：t0 前后 48 小时各体征测了几次",
        "阶段 1 其余项（你同意后）：GD 认转院、检查报告认肺炎、共用对应表",
    ], gap=4), size=12.5, color=BODY, margins=(0, 0, 0, 0))
    card(slide, CL + bw + 0.25, 4.46, bw, 1.8, "有余力再做", header_h=0.42, header_size=14, header_fill=GREY)
    add_text(slide, CL + bw + 0.41, 4.96, bw - 0.32, 1.26, bullets([
        "每周文献：可接着讲解页读 MEDS / ACES、YAIB（任务定义比模型影响更大）",
        "对外沟通草稿 C1–C3 按 v0.3 更新，定何时发（D27）",
        "用这份 PPT 向导师或组里同步本周进展",
    ], gap=4, color=GREY), size=12.5, color=BODY, margins=(0, 0, 0, 0))


def build_checks(slide):
    frame(slide, "周末核对", "§7 核对清单：这个周末先做 7 项",
          "§7 共 14 项，其中 7 项这个周末就能从 std_v3 导出包、导入回执和分级页里取到。",
          ask_lead="回报格式：", ask="按中心 × 项目的计数（1–9 写 <10）；名称清单不含病例内容；重导前、后各一份。",
          notes=NOTES["checks"])
    num = lambda t: {"text": t, "bold": True, "color": DARK, "align": "c", "size": 14}
    src = lambda t: {"text": t, "color": BODY}
    rep = lambda t: {"text": t, "bold": True, "color": NAVY}
    rows = [
        ["优先", "项", "核对什么", "从哪里取", "回报什么"],
        ["", num("11"), "新开始氧疗：O2_NEW 就诊数（起点 24 小时内 / 之后）、GD 氧疗医嘱时长分档、“必要时”氧疗行数",
         src("对齐报告、回执、小脚本"), rep("按中心计数")],
        ["", num("12"), "给氧前后 24 小时内是否有低氧：吸空气 SpO2 ≤ 93%、PaO2/FiO2 ≤ 300 或 RR ≥ 30",
         src("events + 小脚本"), rep("有 / 没有 各几例")],
        ["", num("13"), "ICU 科室名称：认成 ICU 的、被排除的（如“呼吸与危重症医学科”）；术后转入住院数",
         src("对齐报告、科室清单"), rep("名称 + 住院数")],
        ["", num("14"), "长期家庭氧疗（COM018）、家庭无创（COM019）的住院数与被排除的命中数",
         src("分级页的排除命中"), rep("按中心计数")],
        ["", num("2"), "统一文字规则前后，各症状、基础病、并发症的阳性数", src("两份对齐报告"),
         rep("中心 × 数据项")],
        ["", num("4"), "GD 诊断编码规则与名称规则的一致性", src("对齐报告"), rep("一致 / 只编码 / 只名称")],
        ["", num("8"), "两院的流感检测名称、ICU 科室名称", src("labs、encounters"), rep("名称清单")],
    ]
    tbl = add_table(slide, CL, 1.82, [0.95, 0.5, 6.45, 2.1, 2.23],
                    [0.38, 0.56, 0.56, 0.56, 0.4, 0.4, 0.4, 0.4], rows, size=12, header_size=13,
                    margins=(0.1, 0.03, 0.08, 0.03))
    merge_rows(tbl, 0, 1, 4, [{"runs": "先做", "bold": True, "size": 15, "color": RED, "align": "c"},
                              {"runs": "v0.3\n风险点", "size": 11, "color": MUTED, "align": "c"}], fill="FBE9E7")
    merge_rows(tbl, 0, 5, 7, [{"runs": "其次", "bold": True, "size": 15, "color": DARK, "align": "c"},
                              {"runs": "对齐质量", "size": 11, "color": MUTED, "align": "c"}])
    add_text(slide, CL, 5.84, CW, 0.5,
             [[("暂缓：", {"bold": True, "color": NAVY}),
               ("第 1 项要对照检验报告单；第 3、5、6 项依赖阶段 1 还没做的功能；第 9 项等站点可区分性审计脚本；"
                "第 7、10 项下周再做。", {})]],
             size=11.5, color=MUTED, margins=(0, 0, 0, 0))


def build_decide(slide):
    frame(slide, "决定", "两件这个周末定，五件下周组会前定",
          "A 档决定归你：下面只列选项和理由；标“建议”的是参考，不是结论。",
          ask_lead="顺序：", ask="周日先定上面两件，下周一 Claude 就能开工；下面五件不急，下周组会前定。",
          notes=NOTES["decide"])
    cw = (CW - 0.25) / 2
    y0, ch_ = 1.82, 2.76
    # 左：§8 第 11 项
    card(slide, CL, y0, cw, ch_, "① 再住院的事件进不进标签（§8 第 11 项，A 档）", header_h=0.46, header_size=14.5)
    add_text(slide, CL + 0.18, y0 + 0.5, cw - 0.36, 0.32, "现状：住院任务的标签只看本次住院；D73 只定了谁做索引。",
             size=12, color=MUTED, anchor="m", margins=(0, 0, 0, 0))
    opts = [
        ("A", "不算（维持现状）", "口径干净；但出院后 14 天内再住院才转重的会记成“未转重”", False),
        ("B", "算", "更贴近“病程”；但改变了标签定义，要出新版本", False),
        ("C", "主分析不算、敏感性分析算", "先数一数真实数据里 14 天内再住院有多少，再定", True),
    ]
    oy = y0 + 0.86
    for letter, name, why, rec in opts:
        oh = 0.56
        if rec:
            add_rect(slide, CL + 0.12, oy - 0.03, cw - 0.24, oh + 0.06, fill=LIGHT)
        badge = add_rect(slide, CL + 0.2, oy + 0.1, 0.36, 0.36, fill=DARK if rec else WHITE,
                         line=None if rec else MID, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        shape_text(badge, letter, 13, WHITE if rec else NAVY, bold=True)
        title_runs = [(name, {"bold": True, "color": NAVY})]
        if rec:
            title_runs.append(("　建议", {"bold": True, "color": RED}))
        add_text(slide, CL + 0.68, oy, cw - 0.88, 0.26, [title_runs], size=13, color=NAVY, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, CL + 0.68, oy + 0.26, cw - 0.88, 0.3, why, size=11.5, color=BODY, anchor="m",
                 margins=(0, 0, 0, 0))
        oy += oh + 0.05
    # 右：§8 第 14 项
    rx = CL + cw + 0.25
    card(slide, rx, y0, cw, ch_, "② 阶段 1 其余项开不开工（§8 第 14 项，非 A 档）", header_h=0.46, header_size=14.5)
    add_text(slide, rx + 0.18, y0 + 0.58, cw - 0.36, 1.62, bullets([
        "GD 也从出院记录认转院、转科；转院的结局记“不可观测”",
        "胸部 CT / X 线报告的结论栏认肺炎（两院同一条规则）",
        "共用药品对应表、流感检测对应表",
        "每条记录标来源类型；检验另存报告时间",
    ], gap=4), size=12.5, color=BODY, margins=(0, 0, 0, 0))
    add_rect(slide, rx + 0.12, y0 + ch_ - 0.5, cw - 0.24, 0.38, fill=LIGHT)
    add_text(slide, rx + 0.24, y0 + ch_ - 0.5, cw - 0.48, 0.38, "同意即可开工；全部在虚构数据上开发和测试。",
             size=12.5, color=NAVY, bold=True, anchor="m", margins=(0, 0, 0, 0))

    add_text(slide, CL, 4.68, CW, 0.32, "可以缓到下周组会前定（A 档或治理）", size=14, color=NAVY, bold=True,
             anchor="m", margins=(0, 0, 0, 0))
    later = [
        ("发布层兼容 MEDS？", "放在 flucase 里，还是单独做成工具"),
        ("支持哪几种切分？", "按病人、按流感季、按中心"),
        ("首次发布前留封存样本？", "发布后再收回就晚了（MIMIC-IV 的做法）"),
        ("院内发病进主人群？", "GAM 有 109/173 次住院首次阳性在入院 48 小时后"),
        ("跨院只用最小公共数据集？", "年龄、症状、体温、SpO2 等门急诊拿得到的"),
    ]
    lw_ = (CW - 4 * 0.15) / 5
    for i, (t, sub) in enumerate(later):
        x = CL + i * (lw_ + 0.15)
        add_rect(slide, x, 5.04, lw_, 1.2, fill=WHITE, line=BORDER, shadow=True)
        add_text(slide, x + 0.12, 5.1, lw_ - 0.24, 0.42, t, size=12.5, color=NAVY, bold=True, anchor="m",
                 margins=(0, 0, 0, 0))
        add_text(slide, x + 0.12, 5.52, lw_ - 0.24, 0.66, sub, size=11, color=MUTED, margins=(0, 0, 0, 0))


def build_next(slide):
    frame(slide, "下一步", "下周衔接：从计数到结论，赶在流感季前",
          "关键路径：计数交回 → 定稿口径与读取器 → 医生试用 → 流感高峰前上线前瞻收集。",
          ask_lead="11 月底前要有：", ask="v0.3 在两院的真实计数、定稿的口径、可复现的 std_v3 数据包，和六模型汇总表。",
          notes=NOTES["next"])
    ly = 2.42
    add_line(slide, 0.85, ly, 12.45, ly, DARK, 2.5, arrow_end=True)
    ms = [
        ("10-26 周一", "交回计数；改规则、\n开阶段 1 其余项", CYAN),
        ("10-27–28", "医生可用性测试\n1–2 位，虚构样例", CYAN),
        ("10-29–30", "v0.3 两院真实\n结局计数小结", CYAN),
        ("合作同学结果到后", "六模型汇总表\n与配对比较", DARK),
        ("11 月初", "GD、GAM 医生\n用真实数据试导", CYAN),
        ("11 月中", "前瞻录入表单\nv1 定稿", CYAN),
        ("12 月", "流感高峰", RED),
    ]
    x0, step = 1.35, (11.95 - 1.35) / 6
    for i, (date, label, col) in enumerate(ms):
        mx = x0 + i * step
        add_oval(slide, mx, ly, 0.26, fill=col, line=WHITE, line_w=2)
        add_text(slide, mx - 0.86, 1.76, 1.72, 0.48, date, size=13.5, color=RED if col == RED else NAVY,
                 bold=True, align="c", anchor="b")
        add_text(slide, mx - 0.86, 2.62, 1.72, 0.6, label, size=12, color=MUTED, align="c")
    for j, (text, col) in enumerate((("flucase / 医院线", CYAN), ("六模型实验", DARK), ("流感季", RED))):
        lx = 8.15 + j * 1.6
        add_oval(slide, lx, 3.3, 0.16, fill=col, line=WHITE, line_w=1.5)
        add_text(slide, lx + 0.13, 3.16, 1.4, 0.28, text, size=10.5, color=MUTED, anchor="m",
                 margins=(0.02, 0, 0, 0))

    item = lambda t: {"text": t, "bold": True, "color": NAVY}
    rows = [
        ["风险", "要盯的信号", "对策"],
        [item("氧疗代理在两院检出机会不同"), "GD 的 O2_NEW 远多于 GAM；给氧时多数没有低氧",
         "按中心分层报告；用 guideline 口径量出代理贡献了多少事件"],
        [item("ICU 科室名称误判"), "“呼吸与危重症医学科”这类普通病房被认成 ICU",
         "v0.3 已排除；看第 13 项的名称清单再补规则"],
        [item("GD 检验迟迟补不上编号"), "只有 JZH 的检验进不了库，检验覆盖率偏低",
         "改好给 GD 数据方的草稿，何时发由你定（D27）"],
        [item("两院做不出可信的跨院验证"), "留一中心只有 2 折，估不出中心间差异",
         "新中心或新流感季作外部验证，事先锁定、只评一次"],
        [item("流感季临近"), "医生测试、前瞻表单拖到 12 月", "下周约测试；11 月中定稿前瞻表单"],
    ]
    add_table(slide, CL, 3.56, [3.0, 4.35, 4.88], [0.38, 0.44, 0.44, 0.44, 0.44, 0.44], rows, size=12,
              header_size=12.5, margins=(0.1, 0.03, 0.08, 0.03))


def build_closing(slide):
    band = add_rect(slide, 0, 2.55, SW, 1.6, fill="4472C4", name="致谢横带")
    shape_text(band, "谢谢！欢迎讨论与指正", 40, WHITE)
    for r in band.text_frame.paragraphs[0].runs:
        rPr = r._r.get_or_add_rPr()
        eff = etree.Element(qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="50800", dist="38100", dir="5400000", algn="ctr",
                              rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="45000")
        rPr.find(qn("a:solidFill")).addnext(eff)
    add_text(slide, 1.0, 4.45, SW - 2.0, 0.5, "周末必做三件事见第 11 页 · 核对清单见第 13 页 · 待定决定见第 14 页",
             size=16, color=NAVY, align="c", anchor="m")
    slide.notes_slide.notes_text_frame.text = NOTES["closing"]


# ---------------------------------------------------------------------------
# 封面与删页
# ---------------------------------------------------------------------------
def set_shape_text(shape, text):
    """改写形状里第一段的第一个文字片段，删掉其余片段，保留原格式。"""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fix_cover(cover):
    for shp in cover.shapes:
        if not shp.has_text_frame:
            continue
        text = shp.text_frame.text.strip()
        if text.startswith("老年流感"):
            set_shape_text(shp, "老年流感转重研究：本周总结与周末计划")
        elif text.startswith("理想数据"):
            set_shape_text(shp, "本周成果 · 现状判断 · 周末必做 · 下周衔接")
        elif text.startswith("2026"):
            set_shape_text(shp, "2026年10月23日")
    cover.notes_slide.notes_text_frame.text = NOTES["cover"]


def delete_slides_after_first(prs):
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst)[1:]:
        rid = sld_id.get(qn("r:id"))
        sld_id_lst.remove(sld_id)
        prs.part.drop_rel(rid)


def main():
    prs = Presentation(str(SRC))
    fix_cover(prs.slides[0])
    delete_slides_after_first(prs)
    blank = prs.slide_layouts.get_by_name("Blank")
    builders = [build_summary, build_overview, build_week, build_gd, build_rules, build_export, build_sixmodel,
                build_numbers, build_status, build_must, build_schedule, build_checks, build_decide, build_next,
                build_closing]
    for build in builders:
        build(prs.slides.add_slide(blank))
    prs.core_properties.title = DECK_TITLE
    prs.core_properties.subject = "本周（10-17 至 10-23）进展、现状判断、周末安排与下周衔接"
    prs.save(str(OUT))
    print(f"已生成：{OUT}（共 {len(prs.slides)} 页）")


if __name__ == "__main__":
    main()
