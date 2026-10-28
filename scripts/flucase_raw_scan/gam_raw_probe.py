#!/usr/bin/env python3
"""GAM 原始导出（门诊、急诊、护理记录，以及以前的住院）的结构扫描：只输出结构与计数，不输出病例内容。

为 flucase 的 GAM 新格式（BaseModel plans/GAM三源格式与规则集v0.4_执行计划_20261027.md §3）准备，研究者在本机运行，报告交回：
1. flucase 现有的 GAM 格式（住院 gam_ip、门诊 gam_op）认不认得每个文件：状态、得分、认出的表、缺的表和列；
2. 每张表：表头、行数；编号列（不同值个数、每个编号的行数）；时间列（能解析的比例、精度、年月范围）；
   目录类列（项目、名称、方式、科室、级别等）的名称目录；数值列的分布；文字列的长度与关键词计数；
3. 长表（项目名称 + 值）的项目目录：每个项目的行数、编号数、数值占比、p1 / p50 / p99、常见单位；
4. 编号重合：每张表的编号列和其他文件的编号列有多少相同（加盐哈希，只给计数），看护理记录挂到住院还是门急诊、按就诊还是按病人；
5. 附：通用结构扫描（gd_raw_scan.py 的报告，名字遮蔽、列的写法形状）。

不输出：任何单元格原文——目录类名称除外，规则同 gd_raw_vocab.py：只列在 ≥5 个编号里出现、没有 ≥6 位连续数字、没有人称或联系方式类的词、
≤60 字的名称，其余只计数；编号、姓名、时间原值、病历和护理记录正文的任何片段都不输出。按编号分类的计数 1–4 写成 "<5"。
编号重合用本次随机加盐的哈希计算，盐不落盘。

用法（BaseModel 根目录；第 1 部分要用 flucase 的 venv，别的 Python 也能跑，只是跳过第 1 部分）：
  flucase/backend/.venv/Scripts/python.exe scripts/flucase_raw_scan/gam_raw_probe.py "data/GAM- 护理记录" --out 报告.md [--show-names]
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gd_raw_scan as S  # noqa: E402

K = 5
CAP = 300000                          # 每列最多记这么多个不同编号 / 名称
PERSON = re.compile(r"主任|医生|医师|教授|护士|护师|家属|患者|病人|先生|女士|老师|同志|签字|签名|电话|手机|身份证|住址|地址")
PERSON_HEAD = re.compile(r"姓名|名字|医生|医师|护士|签名|书写人|记录人|操作员|审核|报告人|申请人|执行人|录入|联系人|电话|地址|身份证|住址|籍贯|"
                         r"工作单位|职业|婚姻|民族|责任护士|签字")
ID_HEAD = re.compile(r"关联字段|住院号|病人ID|患者ID|病人编号|患者编号|就诊流水号|流水号|门诊号|急诊号|登记号|就诊号|病案号|病历号|ID$|编号$", re.I)
TIME_HEAD = re.compile(r"时间|日期")
CATALOG_HEAD = re.compile(r"项目|名称|类型|类别|方式|科室|病区|级别|等级|途径|单位|意识|神志|吸氧|给氧|氧疗|护理|部位|体位|标识|状态|来源")
VALUE_HEAD = re.compile(r"^值$|结果|数值|测量值|记录值|内容")
TEXT_HEAD = re.compile(r"病情|观察|措施|记录|内容|描述|备注|小结|主诉|现病史|既往史|体格|诊断|意见|建议")
NUM = re.compile(r"^\s*[<>＜＞≤≥]?=?\s*([-+]?\d+(?:\.\d+)?)\s*$")
LONG_DIGITS = re.compile(r"\d{6,}")
DT_RE = re.compile(r"^\s*(\d{4})\s*[-/.年]\s*(\d{1,2})\s*[-/.月]\s*(\d{1,2})(?:\s*日)?(?:[\sT]+(\d{1,2})\s*[:：时]\s*(\d{1,2})(?:\s*[:：分]\s*(\d{1,2}))?)?")
# 文字列的关键词（只数命中的行数）：吸氧方式与流量、呼吸支持、血氧、监护与 ICU、发热与退热、意识
KEYWORDS = {
    "吸氧": r"吸氧|给氧|氧疗", "鼻导管": r"鼻导管", "面罩": r"面罩", "高流量": r"高流量|HFNC", "无创": r"无创|BiPAP|BIPAP",
    "有创 / 插管": r"有创|气管插管|气管切开|呼吸机辅助", "流量 L/min": r"\d+(?:\.\d+)?\s*(?:L|l|升)\s*/\s*(?:min|分)",
    "SpO2 / 血氧": r"SpO2|SPO2|血氧|氧饱和", "监护": r"监护", "ICU / 重症": r"ICU|重症", "特级 / 一级护理": r"特级护理|一级护理",
    "体温": r"体温|T\s*[:：]?\s*3\d", "发热": r"发热|发烧", "无发热 / 退热": r"无发热|未再发热|热已退|体温正常", "意识 / 神志": r"意识|神志|嗜睡|昏迷|谵妄",
}
KW = {k: re.compile(v) for k, v in KEYWORDS.items()}
VITAL_HEAD = re.compile(r"体温|脉搏|心率|呼吸|血压|收缩压|舒张压|血氧|SpO2|氧饱和|饱和度|氧流量|流量|FiO2|吸氧浓度|氧浓度|体重|身高|入量|出量|尿量|疼痛|评分|GCS", re.I)


def c5(n: int) -> str:
    return "0" if n == 0 else (str(n) if n >= K else "<5")


def pct(a: int, b: int) -> str:
    return f"{100 * a / b:.0f}%" if b else "—"


def quant(xs: list[float]) -> str:
    if not xs:
        return "—"
    xs = sorted(xs)
    q = lambda p: xs[min(len(xs) - 1, int(p * (len(xs) - 1)))]  # noqa: E731
    return f"p1 {q(.01):g} / p50 {q(.5):g} / p99 {q(.99):g}"


def num(s: str | None) -> float | None:
    m = NUM.match(s or "")
    return float(m.group(1)) if m else None


def safe_name(s: str) -> bool:
    return bool(s) and len(s) <= 60 and "\n" not in s and not LONG_DIGITS.search(s) and not PERSON.search(s) and not S.DATE_ANY.search(s)


def dt_of(v) -> tuple[dt.datetime | None, str]:
    """时间 → (时间, 精度 day / minute / second)。"""
    if isinstance(v, dt.datetime):
        return v, ("second" if v.second else ("minute" if (v.hour or v.minute) else "day"))
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day), "day"
    m = DT_RE.match(S.as_text(v) or "")
    if not m:
        return None, ""
    y, mo, d, h, mi, sec = m.groups()
    try:
        t = dt.datetime(int(y), int(mo), int(d), int(h or 0), int(mi or 0), int(sec or 0))
    except ValueError:
        return None, ""
    return t, ("second" if sec else ("minute" if h else "day"))


class Col:
    def __init__(self, j: int, head: str):
        self.j, self.head = j, head
        self.is_id = bool(ID_HEAD.search(head)) and not PERSON_HEAD.search(head)
        self.is_time = bool(TIME_HEAD.search(head)) and not self.is_id
        self.is_person = bool(PERSON_HEAD.search(head))
        self.is_catalog = bool(CATALOG_HEAD.search(head)) and not (self.is_id or self.is_time or self.is_person)
        self.n = 0
        self.nums: list[float] = []
        self.n_num = 0
        self.lens: list[int] = []
        self.ids: set[int] = set()
        self.per_id: Counter = Counter()
        self.prec: Counter = Counter()
        self.tmin: dt.datetime | None = None
        self.tmax: dt.datetime | None = None
        self.t_bad = 0
        self.names: dict[str, set[int]] = defaultdict(set)
        self.name_rows: Counter = Counter()
        self.kw: Counter = Counter()
        self.shapes: Counter = Counter()          # 短取值的写法形状（如血压 999/99），不含原值


class Tab:
    def __init__(self, fid: str, sheet: str, header: list[str]):
        self.fid, self.sheet, self.header = fid, sheet, header
        self.cols = [Col(j, h) for j, h in enumerate(header) if h]
        self.n = 0
        self.key = next((c for c in self.cols if c.is_id), None)
        self.item = next((c for c in self.cols if re.search(r"项目|名称", c.head) and c.is_catalog), None)
        self.value = next((c for c in self.cols if VALUE_HEAD.search(c.head) and not c.is_catalog), None)
        self.unit = next((c for c in self.cols if "单位" in c.head), None)
        self.long = self.item is not None and self.value is not None
        self.items: dict[str, dict] = defaultdict(lambda: {"rows": 0, "ids": set(), "num": 0, "vals": [], "u": Counter()})
        self.span: dict[int, list] = {}             # 编号 → [最早, 最晚]（护理记录看一个编号覆盖多少天）
        self.tcol = next((c for c in self.cols if c.is_time), None)

    @property
    def label(self) -> str:
        return f"{self.fid}“{S.render_label(self.sheet, ascii_strict=True)}”"

    def add(self, row: list) -> None:
        self.n += 1
        kid = None
        if self.key is not None and self.key.j < len(row):
            s = S.as_text(row[self.key.j])
            if s:
                kid = S.h64(S.canon_id(s))
        for c in self.cols:
            v = row[c.j] if c.j < len(row) else None
            s = S.as_text(v)
            if s is None:
                continue
            c.n += 1
            if c.is_id:
                h = S.h64(S.canon_id(s))
                if len(c.ids) < CAP:
                    c.ids.add(h)
                if c is self.key:
                    c.per_id[h] += 1
                continue
            if c.is_person:
                continue
            if c.is_time:
                t, p = dt_of(v)
                if t is None:
                    c.t_bad += 1
                else:
                    c.prec[p] += 1
                    c.tmin = t if c.tmin is None or t < c.tmin else c.tmin
                    c.tmax = t if c.tmax is None or t > c.tmax else c.tmax
                    if c is self.tcol and kid is not None:
                        sp = self.span.setdefault(kid, [t, t])
                        sp[0], sp[1] = min(sp[0], t), max(sp[1], t)
                continue
            x = num(s) if not isinstance(v, (int, float)) else float(v)
            if x is not None:
                c.n_num += 1
                if len(c.nums) < 50000:
                    c.nums.append(x)
            if len(c.lens) < 50000:
                c.lens.append(len(s))
            if len(s) <= 12 and len(c.shapes) < 200:
                c.shapes[S.shape(s)] += 1
            if c.is_catalog and x is None and len(c.names) < CAP:
                c.names[s].add(kid if kid is not None else self.n)
                c.name_rows[s] += 1
            if len(s) > 12:
                for k, rx in KW.items():
                    if rx.search(s):
                        c.kw[k] += 1
        if self.long:
            name = S.as_text(row[self.item.j]) if self.item.j < len(row) else None
            if name:
                it = self.items[name] if len(self.items) < 5000 or name in self.items else None
                if it is None:
                    return
                it["rows"] += 1
                if kid is not None:
                    it["ids"].add(kid)
                val = S.as_text(row[self.value.j]) if self.value.j < len(row) else None
                x = num(val)
                if x is not None:
                    it["num"] += 1
                    if len(it["vals"]) < 20000:
                        it["vals"].append(x)
                if self.unit is not None and self.unit.j < len(row):
                    u = S.as_text(row[self.unit.j])
                    it["u"][u if u and safe_name(u) else ("（空）" if not u else "〔不列〕")] += 1


def read_tables(scan: S.Scan):
    from openpyxl import load_workbook
    for f in scan.files:
        fid = scan.fid[f]
        kind = S.sniff(f)
        sheets = []
        try:
            if kind == "xlsx":
                wb = load_workbook(f, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    ws.reset_dimensions()
                    sheets.append((ws.title, ws.iter_rows(values_only=True)))
            elif kind == "xls":
                sheets = list(S.iter_xls(f))
            elif kind == "html":
                sheets = list(S.iter_html_tables(f))
            elif kind == "文本":
                enc, got = S.iter_delimited(f)
                if got:
                    sheets = [("文本表", got[1])]
        except Exception as e:  # noqa: BLE001
            print(f"{fid} 读不了：{type(e).__name__}", file=sys.stderr)
            continue
        for name, rows in sheets:
            head: list = []
            it = iter(rows)
            for r in it:
                head.append(list(r))
                if len(head) >= 15:
                    break
            h = S.pick_header(head)
            if h is None:
                continue
            t = Tab(fid, name, [(S.as_text(v) or "").strip() for v in head[h]])
            print(f"…{t.label}", file=sys.stderr)
            for r in head[h + 1:]:
                if any(S.as_text(v) for v in r):
                    t.add(list(r))
            empty = 0
            for r in it:
                r = list(r)
                if not any(S.as_text(v) for v in r):
                    empty += 1
                    if empty >= S.EMPTY_STOP:
                        break
                    continue
                empty = 0
                t.add(r)
            yield f, t


def flucase_match(scan: S.Scan) -> list[str]:
    """用 flucase 现有的 GAM 格式认每个文件（只看表头）。"""
    try:
        from flucase.repo.readers.gam_ehr import GamIpReader, GamOpReader
    except Exception as e:  # noqa: BLE001
        return [f"（没有用 flucase 的 venv 运行，跳过：{type(e).__name__}）"]
    out = ["| 文件 | 格式 | 认定 | 得分 | 认出的工作表 / 差异 |", "|---|---|---|---|---|"]
    for f in scan.files:
        if f.suffix.lower() not in (".xlsx", ".xlsm", ".xls"):
            continue
        for r in (GamIpReader(), GamOpReader()):
            try:
                m = r.match(f)
            except Exception as e:  # noqa: BLE001
                out.append(f"| {scan.fid[f]} | {r.name} | 出错 {type(e).__name__} | — | — |")
                continue
            parts = [S.render_label(x, ascii_strict=True) for x in (m.evidence or [])]
            diffs = "；".join(S.render_label(d.get("detail", ""), ascii_strict=True) for d in (m.differences or [])[:12])
            parts += [f"差异：{diffs}"] if diffs else []
            out.append(f"| {scan.fid[f]} | {r.name} | {m.status} | {m.score:g} | {'；'.join(parts) or '—'} |")
    return out


def report(scan: S.Scan, tabs: list[Tab]) -> str:
    L: list[str] = []
    w = L.append
    w("# GAM 原始导出扫描（门诊、急诊、护理记录；只含结构与计数）\n")
    w(f"- 扫描时间：{dt.datetime.now():%Y-%m-%d %H:%M}；文件 {len(scan.files)} 个，表 {len(tabs)} 张；名字{'未遮蔽' if scan.show else '已遮蔽'}")
    w("- 名称目录只列在 ≥5 个编号里出现、无长数字、无人称词、≤60 字的名称；按编号分类的计数 1–4 写成 \"<5\"\n")
    w("## 1 flucase 现有 GAM 格式认不认得\n")
    L += flucase_match(scan)
    w("\n状态：exact 完全符合；similar 相似（差异里写缺什么）；none 认不出。\n")
    w("## 2 各表概况\n")
    w("| 表 | 行数 | 列数 | 编号列（不同值） | 时间列 | 长表（项目 + 值） |")
    w("|---|---|---|---|---|---|")
    for t in tabs:
        ids = "、".join(f"{S.render_label(c.head)}（{c5(len(c.ids))}）" for c in t.cols if c.is_id) or "—"
        tms = "、".join(S.render_label(c.head) for c in t.cols if c.is_time) or "—"
        lg = f"{S.render_label(t.item.head)} + {S.render_label(t.value.head)}" if t.long else "—"
        w(f"| {t.label} | {t.n} | {len(t.cols)} | {ids} | {tms} | {lg} |")
    for t in tabs:
        w(f"\n## 表 {t.label}\n")
        w("表头：" + " · ".join(S.render_label(h) or "（空）" for h in t.header) + "\n")
        w("| 列 | 填充 | 类别 | 数值占比 | 数值分布 | 长度 p50 / p90 | 短取值的写法形状 |")
        w("|---|---|---|---|---|---|---|")
        for c in t.cols:
            kind = "编号" if c.is_id else ("时间" if c.is_time else ("人员（不看）" if c.is_person else ("目录" if c.is_catalog else "取值")))
            dist = quant(c.nums) if c.nums and not c.is_id and (VITAL_HEAD.search(c.head) or c.n_num / max(c.n, 1) > .8) else "—"
            ln = f"{statistics.median(c.lens):g} / {sorted(c.lens)[int(.9 * (len(c.lens) - 1))]}" if c.lens else "—"
            shp = "、".join(f"{k}（{pct(v, c.n)}）" for k, v in c.shapes.most_common(3)) if c.shapes and not c.is_catalog else "—"
            w(f"| {S.render_label(c.head)} | {pct(c.n, t.n)} | {kind} | {pct(c.n_num, c.n) if not (c.is_id or c.is_time) else '—'} | {dist} | {ln} "
              f"| {shp} |")
        if t.key is not None and t.key.per_id:
            per = sorted(t.key.per_id.values())
            line = f"\n编号列 {S.render_label(t.key.head)}：{c5(len(per))} 个不同值；每个编号的行数 {quant([float(x) for x in per])}"
            if t.span:
                days = [round((b - a).total_seconds() / 86400, 1) for a, b in t.span.values()]
                line += f"；按 {S.render_label(t.tcol.head)}，每个编号覆盖的天数 {quant(days)}（按就诊编号的一般几天到几周，按病人编号的会跨多次就诊）"
            w(line)
        for c in t.cols:
            if c.is_time and (c.prec or c.t_bad):
                tot = sum(c.prec.values()) + c.t_bad
                w(f"\n时间列 {S.render_label(c.head)}：能解析 {pct(sum(c.prec.values()), tot)}；精度 " +
                  "、".join(f"{p} {pct(n, tot)}" for p, n in c.prec.most_common()) +
                  (f"；范围 {c.tmin:%Y-%m} 至 {c.tmax:%Y-%m}" if c.tmin and c.tmax else ""))
        cats = [c for c in t.cols if c.is_catalog and c.names]
        for c in cats:
            if c.lens and statistics.median(c.lens) > 20:
                w(f"\n目录列 {S.render_label(c.head)}：取值多为长文字（长度中位数 {statistics.median(c.lens):g}），不列写法")
                continue
            ok = [(n, len(ids), c.name_rows[n]) for n, ids in c.names.items() if len(ids) >= K and safe_name(n)]
            ok.sort(key=lambda x: -x[2])
            hidden = len(c.names) - len(ok)
            w(f"\n目录列 {S.render_label(c.head)}：{len(c.names)} 种写法；列出 {len(ok)} 种" + (f"，另 {hidden} 种只计数（<5 个编号或不合列出规则）" if hidden else ""))
            if ok:
                w("\n| 写法 | 编号数 | 行数 |")
                w("|---|---|---|")
                for n, k, r in ok[:80]:
                    w(f"| {n} | {k} | {r} |")
                if len(ok) > 80:
                    w(f"| …另 {len(ok) - 80} 种 | | |")
        kws = [c for c in t.cols if c.kw]
        if kws:
            w("\n文字列关键词（命中的行数，只数不摘）：\n")
            w("| 列 | " + " | ".join(KEYWORDS) + " |")
            w("|---|" + "---|" * len(KEYWORDS))
            for c in kws:
                w(f"| {S.render_label(c.head)} | " + " | ".join(str(c.kw.get(k, 0)) for k in KEYWORDS) + " |")
        if t.long and t.items:
            ok = [(n, d) for n, d in t.items.items() if len(d["ids"]) >= K and safe_name(n)]
            ok.sort(key=lambda x: -x[1]["rows"])
            w(f"\n长表项目目录（{len(t.items)} 种项目，列出 {len(ok)} 种）：\n")
            w("| 项目 | 行数 | 编号数 | 数值占比 | 数值分布 | 常见单位 |")
            w("|---|---|---|---|---|---|")
            for n, d in ok[:150]:
                u = "、".join(f"{k}（{v}）" for k, v in d["u"].most_common(3)) or "—"
                w(f"| {n} | {d['rows']} | {c5(len(d['ids']))} | {pct(d['num'], d['rows'])} | {quant(d['vals'])} | {u} |")
    # 编号重合
    w("\n## 3 编号重合（加盐哈希，只给计数）\n")
    w("每行：这张表的编号列有多少个不同值也出现在另一个文件的编号列里（同一文件里的表之间不列）。看护理记录按什么编号挂、挂到住院还是门急诊。"
      "重合 <5 时不给占比。\n")
    w("| 表.列 | 不同值 | 另一表.列 | 重合 | 占比 |")
    w("|---|---|---|---|---|")
    idcols = [(t, c) for t in tabs for c in t.cols if c.is_id and c.ids]
    for t, c in idcols:
        for u, d in idcols:
            if u.fid == t.fid:
                continue
            k = len(c.ids & d.ids)
            if k:
                w(f"| {t.label}.{S.render_label(c.head)} | {c5(len(c.ids))} | {u.label}.{S.render_label(d.head)} | {c5(k)} "
                  f"| {pct(k, len(c.ids)) if k >= K else '—'} |")
    return "\n".join(L)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--show-names", action="store_true", help="文件名、目录名不遮蔽（研究者确认名字里没有病人信息时才用）")
    ap.add_argument("--no-structure", action="store_true", help="不附通用结构扫描（gd_raw_scan.py）")
    a = ap.parse_args()
    if not a.root.is_dir():
        sys.exit(f"不是目录：{a.root}")
    scan = S.Scan(a.root.resolve(), a.show_names)
    if a.no_structure:
        for dp, dns, fns in __import__("os").walk(scan.root):
            dns[:] = sorted(d for d in dns if not d.startswith("."))
            scan.files += [Path(dp) / fn for fn in sorted(fns) if not (fn.startswith(".") or fn.startswith("~$") or fn == "Thumbs.db")]
        scan.build_keep()
        for i, f in enumerate(scan.files):
            scan.fid[f] = f"F{i + 1:03d}"
        tail = ""
    else:
        scan.run()
        tail = "\n\n---\n\n" + "\n".join(("#" + x if x.startswith("#") else x) for x in scan.report().splitlines()).replace(
            "# 原始导出结构扫描", "# 附：通用结构扫描", 1)
    tabs = [t for _, t in read_tables(scan)]
    rep = report(scan, tabs) + tail
    if a.out:
        a.out.write_text(rep, encoding="utf-8")
        print(f"报告已写到 {a.out}", file=sys.stderr)
    print(rep)


if __name__ == "__main__":
    main()
